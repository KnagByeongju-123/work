# -*- coding: utf-8 -*-
"""
태진다이텍 MES 웹 서버 - 작업지시 관리 (사내망 전용, 인터넷 필요 없음)

- 이 PC 1대에만 오라클(Instant Client + oracledb)을 설치하고 이 서버를 켜 둡니다.
- 다른 PC·폰은 브라우저로  http://이PC주소:8000  에 접속합니다. (설치 없음)
- 기존 프로그램(mes_7.x.py)과 같은 폴더에 두면 설정(mold_config.json), SPM(spm.xlsx),
  작업기록(mold_register_log.csv), 삭제 백업(mold_deleted_backup.jsonl)을 같이 씁니다.
필요: pip install oracledb openpyxl   (그 외는 파이썬 기본 기능만 사용)
"""
import os
import re
import io
import sys
import csv
import json
import copy
import socket
import threading
import traceback
import datetime as dt
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse, unquote, quote

try:
    import oracledb
except Exception as _e:          # 서버 화면에 알려주고 계속 (화면은 열림)
    oracledb = None
    _ORA_ERR = repr(_e)
try:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except Exception:
    openpyxl = None

VERSION = "web 1.0 (MES v7.7 작업지시)"
BASE = os.path.dirname(os.path.abspath(__file__))
CFG_PATH = os.path.join(BASE, "mold_config.json")
LOG_PATH = os.path.join(BASE, "mold_register_log.csv")
BACKUP_PATH = os.path.join(BASE, "mold_deleted_backup.jsonl")
SPM_XLSX = os.path.join(BASE, "spm.xlsx")
LOG_COLS = ["일시", "등록자", "구분", "테이블", "키", "입력내용"]

DB_USER = "infinity21_pimmes"          # MES 접속 계정 (기존 프로그램과 같음)
DB_PASS = "infinity21_pimmes"

DEFAULT_CFG = {
    "oracle": {"user": "", "dsn": "192.168.1.250:1521/XE", "lib_dir": r"C:\instantclient_19_30"},
    "schema": "INFINITY21_PIMMES",
    "kor_bytes": 3,
    "limit": 5000,
    "item_master_table": "ICOM_ITEM_MASTER",
    "wo_table": "IPLN_WORK_ORDER_MASTER",
    "bom_table": "ICOM_ITEM_CHILD",
    "item": {"table": "ICOM_MOLD_ITEM"},
    "web_port": 8000,
    "web_pin": "",                 # 비워두면 누구나 접속. 숫자를 넣으면 화면에서 PIN을 물어봄
    "web_reg_user": "PYWEB",       # MES 등록자/수정자 칸에 들어갈 값
}

ID_RE = re.compile(r"^[A-Z_][A-Z0-9_$#]*(\.[A-Z_][A-Z0-9_$#]*){0,2}$")
CTRL_RE = re.compile(r"[\x00-\x1f\x7f]")
RE_AUTO_DATE = re.compile(r"(ENTER|CREATE|CREATED|CREATION|REG|REGIST|INSERT|INS|LAST_UPDATE|"
                          r"UPDATE|UPDATED|UPD|MODIFY|MOD)_?(DATE|DT|TIME|DTM)$")
RE_AUTO_USER = re.compile(r"(ENTER|CREATE|CREATED|REG|REGIST|INSERT|INS|LAST_UPDATE|"
                          r"UPDATE|UPDATED|UPD|MODIFY|MOD)_?(BY|USER|USER_ID|EMP|EMP_NO)$")
RE_UPD_DATE = re.compile(r"(LAST_UPDATE|UPDATE|UPDATED|UPD|MODIFY|MODIFIED|MOD)_?(DATE|DT|TIME|DTM)$")
RE_UPD_USER = re.compile(r"(LAST_UPDATE|UPDATE|UPDATED|UPD|MODIFY|MODIFIED|MOD)_?(BY|USER|USER_ID|EMP|EMP_NO)$")
WO_RE = re.compile(r"^WO([0-9A-Z]+?)P(\d+)$")
ORA_HINT = [
    ("ORA-01400", "필수 컬럼이 비어 있습니다 (아래 괄호 안 마지막 이름이 컬럼)"),
    ("ORA-00001", "같은 값이 이미 있습니다 (중복)"),
    ("ORA-02291", "연결된 기준 데이터(부모 키)가 없습니다"),
    ("ORA-02292", "다른 테이블에서 쓰고 있는 값이라 바꿀 수 없습니다"),
    ("ORA-12899", "값이 너무 깁니다"),
    ("ORA-01722", "숫자 칸에 글자가 들어갔습니다"),
    ("ORA-01031", "이 계정에 입력/수정 권한이 없습니다 (DB 관리자에게 권한 요청)"),
    ("ORA-00942", "테이블이 없거나 볼 권한이 없습니다"),
    ("ORA-01861", "날짜 형식이 맞지 않습니다"),
    ("ORA-00904", "없는 컬럼 이름입니다"),
    ("ORA-01438", "숫자가 허용 자릿수보다 큽니다"),
]

# ---------- 작업지시 표 설정 (기존 프로그램 GRID_SPECS["wo"] 와 같음) ----------
WO_FIRST = ["WORK_ORDER_DATE", "MACHINE_CODE", "MACHINE_NAME", "WORK_ORDER_NO", "ITEM_CODE", "ITEM_NAME",
            "MOLD_CODE", "LINK_MOLD__", "PLAN_QTY", "ACTUAL_QTY", "PASS_QTY", "BAD_QTY",
            "WORK_ORDER_STATUS", "PLAN_STATUS", "WORK_SHIFT", "START_DATE", "END_DATE", "ST_VALUE"]
WO_NAMES = {"WORK_ORDER_DATE": "지시일", "MACHINE_CODE": "설비", "MACHINE_NAME": "설비명",
            "WORK_ORDER_NO": "작업지시번호", "ITEM_CODE": "품번", "ITEM_NAME": "품명",
            "MOLD_CODE": "금형", "LINK_MOLD__": "현재 연결 금형(참고)", "PLAN_QTY": "계획수량",
            "ACTUAL_QTY": "실적수량", "PASS_QTY": "양품", "BAD_QTY": "불량",
            "WORK_ORDER_STATUS": "상태", "PLAN_STATUS": "계획상태", "WORK_SHIFT": "근무조",
            "START_DATE": "시작", "END_DATE": "종료", "ST_VALUE": "ST(초)"}
WO_LABEL = ["WORK_ORDER_NO", "WORK_ORDER_DATE", "MACHINE_CODE"]
WO_LOCKED = {"ORGANIZATION_ID": "회사코드는 바꿀 수 없습니다.",
             "WORK_ORDER_NO": "작업지시번호는 실적·자재출고와 연결되어 있어 바꿀 수 없습니다.",
             "WIP_ENTITY_ID": "ERP 작업번호라 바꿀 수 없습니다.",
             "MOLD_CODE": "금형은 품번-금형 연결로 자동으로 정해집니다 (MES 트리거).\n"
                          "금형을 바꾸려면 [금형·품번 관리]에서 품번 연결을 고치세요.",
             "MACHINE_NAME": "설비명은 설비를 바꾸면 자동으로 바뀝니다.",
             "ENTER_DATE": "등록일은 바꾸지 않습니다.", "ENTER_BY": "등록자는 바꾸지 않습니다."}
WO_SENSITIVE = {
    "ACTUAL_QTY": "실적수량은 현장 실적(생산현황판)과 연결된 값입니다.",
    "PASS_QTY": "양품수량은 현장 실적과 연결된 값입니다.",
    "BAD_QTY": "불량수량은 현장 실적과 연결된 값입니다.",
    "INPUT_QTY": "투입수량은 현장 실적과 연결된 값입니다.",
    "WORK_ORDER_STATUS": "상태를 바꾸면 현장 화면에서 작업지시가 보이거나 사라질 수 있습니다.",
    "PLAN_STATUS": "상태를 바꾸면 현장 화면에서 작업지시가 보이거나 사라질 수 있습니다.",
    "MACHINE_CODE": "설비를 바꾸면 다른 설비로 작업지시가 옮겨갑니다 (설비명은 자동).",
    "ITEM_CODE": "품번을 바꾸면 금형이 새 품번의 연결 금형으로 자동으로 바뀝니다.",
    "WORK_ORDER_DATE": "지시일을 바꾸면 일자별 계획·실적 집계가 달라집니다.",
}
DONE_COLS = ("ACTUAL_QTY", "PASS_QTY", "BAD_QTY", "INPUT_QTY")
WO_RESET = ("ACTUAL_QTY", "PASS_QTY", "BAD_QTY", "INPUT_QTY", "OUT_RST_QTY", "IN_RST_QTY", "OUT_PREVIOUS_QTY",
            "IN_PREVIOUS_QTY", "AVG_TACT_TIME", "LASTEST_TACT_TIME", "EXTRA_IN_QTY", "EXTRA_OUT_QTY", "INTERFACE_DATE")
T_UP = "IPLN_WORK_ORDER_UPLOADER"
PROC = "P_PLN_WORKORDER_BATCH_CREATE"

# 작업지시번호 앞부분(WO 다음 ~ P 앞) 규칙 후보
WO_RULES = (
    ("년2+월(1~12)+일2", lambda d: f"{d:%y}{d.month}{d:%d}"),
    ("년2+월2+일2", lambda d: f"{d:%y}{d:%m}{d:%d}"),
    ("년2+월(1~9,A~C)+일2", lambda d: f"{d:%y}{'123456789ABC'[d.month - 1]}{d:%d}"),
    ("년1+월2+일2", lambda d: f"{d.year % 10}{d:%m}{d:%d}"),
    ("년1+월(1~9,A~C)+일2", lambda d: f"{d.year % 10}{'123456789ABC'[d.month - 1]}{d:%d}"),
)

SPM_DEFAULT = {
    '1515-196': 45, '1547-102-1': 35, '1547-102-2': 45, '1549-156': 40, '1549-157': 35, '1551-189': 35,
    '1553-101-1': 35, '1553-101-2': 35, '1642-118-1': 30, '1654-119-1': 36, '1654-119-2': 32,
    '1657-127-1': 35, '1657-130': 35, '1660-131': 35, '1660-132': 35, '1664-128-1': 35, '1664-128-2': 35,
    '1752-297': 35, '1755-296': 35, '1781-0001-00': 35, '1807-653': 25, '1816-0010-00': 35, '1839-629': 28,
    '1837-0033-01': 35, '1853-515-1': 35, '1853-515-2': 35, '1853-586': 35, '1856-627-2': 60,
    '1856-628-1': 35, '1858-618': 35, '1858-618-2': 35, '1858-619': 35, '1861-596': 35, '1861-597': 35,
    '1862-0041-01': 30, '1862-0041-21': 25, '1862-0042-21': 25, '1862-604': 35, '1862-624': 35,
    '1862-651': 35, '1866-599': 30, '1866-617-1': 35, '1866-617-2': 35, '1869-621': 35, '1870-602-1': 35,
    '1870-602-2': 35, '1871-0043-01': 25, '1871-0043-21': 25, '1873-0044-21': 25, '1873-0045-01': 25,
    '1874-623': 35, '1888-0005': 20, '1945-349-1': 35, '1945-349-2': 40, '1946-0053-01': 35, '1947-353': 35,
    '1950-354-1': 35, '1950-354-2': 35, '1950-374-1': 35, '1950-374-2': 35, '1951-386': 40,
    '1952-0056-01': 30, '1952-0056-21': 25, '1952-346-1': 35, '1952-346-2': 35, '1952-367-2': 30,
    '1958-365-1': 40, '1958-365-2(2호금형)': 30, '1958-373-1': 35, '1958-373-2': 35, '1959-0024-01': 30,
    '1959-0024-21': 25, '1959-360-1': 35, '1959-360-2': 35, '1960-0031-01': 28, '1960-0031-21': 25,
    '1960-361': 35, '1962-0030': 35, '1963-0055-01': 30, '1963-0055-21': 25, '1965-326-1': 35,
    '1965-326-2': 35, '1967-0052-01': 25, '1967-0052-21': 25, '1967-391': 35, '1967-392': 35,
    '1968-0016-01': 30, '1968-0016-02': 40, '1968-362': 35, '1968-364-1': 40, '1968-364-2': 24,
    '1969-0030-01': 35, '1970-363': 35, '1970-372-1': 35, '1970-372-2': 35, '1972-383': 40,
    '1972-383(2호금형)': 40, '1973-325': 30, '1973-370(375)': 35, '1974-0023-01': 25, '1974-0023-21': 25,
    '1976-0004-01': 35, '1979-0051-01': 20, '1982-371': 25, '54610-C1000-(TR)': 22, '54610-C1001-(BK)': 25,
    '54610-D3500-(TR)': 22, '54610-D3501-(BK)': 25, '54610-D3000-(TR)': 23, '54610-D3001-(BK)': 22,
    'AB3038-SC': 23, 'AB3248-SC-01': 22, 'AT3037-GC-00': 23, 'IJ223101-SC-00': 25, 'C657-127': 35,
    'C873-622-2': 35, '1GE0-0009-01': 30, '116090-ST2': 28, '116080-ST0': 25, 'IW/OW-BDA-1106-1': 38,
    'IW/OW-BDA-1106-2': 38, 'IW/OW-BDA-1106-3': 38, 'IW/OW-BDA-1113-1': 38, 'IW/OW-BDA-1113-2': 38,
    'IW/OW-BDA-1113D': 38, 'IW/OW-BDA-1124 BA': 35, 'IW/OW-BDA-1156': 38, 'IW/OW-BDA-1159B': 38,
    'IW/OW-BDA-1165A': 38, 'IW/OW-BDA-1179': 38, 'IW/OW-BDA-1179-2': 38, 'MAG624828': 30, 'MAG629430': 10,
    'MJH637370': 25, '48338-2H010': 36, '45631-26010 A': 25, '45631-26010 B/C': 25, '45631-26010-C': 25,
    '45624-2F010-U': 25, '45624-2F010-L': 25, '45524-26210': 25, '45464-4G600': 25, '45424-4G100L-A': 25,
    '45424-4G100L-B': 15, '45424-4G100L-C': 15, 'D000136433': 22, 'B17101000-0013-00': 22,
}


class UserError(Exception):
    """화면에 그대로 보여줄 안내 (경고)"""


# ==============================================================
# 공통 도구
# ==============================================================
def load_cfg():
    cfg = copy.deepcopy(DEFAULT_CFG)
    if os.path.exists(CFG_PATH):
        with open(CFG_PATH, encoding="utf-8") as f:
            saved = json.load(f)
        for k, v in saved.items():
            if k in ("oracle", "item") and isinstance(v, dict):
                cfg[k].update(v)
            elif k in cfg:
                cfg[k] = v
    return cfg


CFG = load_cfg()


def ora_hint(e):
    s = str(e)
    for code, msg in ORA_HINT:
        if code in s:
            return f"▶ {msg}\n\n{s}"
    return s


def ident(s, what="이름"):
    s = (s or "").strip().upper()
    if not ID_RE.match(s):
        raise ValueError(f"{what} 이름이 올바르지 않습니다: '{s}'")
    return s


def qualify(name):
    tb = ident(name, "테이블")
    s = (CFG.get("schema") or "").strip()
    return tb if (not s or "." in tb) else ident(s, "스키마") + "." + tb


def wo_tbl():
    return qualify(CFG.get("wo_table") or "IPLN_WORK_ORDER_MASTER")


def link_tbl():
    return qualify((CFG.get("item") or {}).get("table") or "ICOM_MOLD_ITEM")


def mst_tbl():
    return qualify(CFG.get("item_master_table") or "ICOM_ITEM_MASTER")


def blen(s, kb=None):
    kb = int(kb or CFG.get("kor_bytes") or 3)
    return sum(1 if ord(ch) < 128 else kb for ch in s)


def show(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def cell_text(v, limit=300):
    try:
        if v is None:
            return ""
        if hasattr(v, "read"):
            v = v.read(1, limit)
        if isinstance(v, (bytes, bytearray)):
            v = v[:limit // 2].hex().upper()
        s = show(v)
    except Exception as e:
        s = f"<읽기오류 {type(e).__name__}>"
    s = CTRL_RE.sub(" ", s)
    return s if len(s) <= limit else s[:limit] + "…"


def parse_day(text):
    t = (text or "").strip()
    if re.fullmatch(r"\d{8}", t):
        return dt.date(int(t[:4]), int(t[4:6]), int(t[6:]))
    m = re.fullmatch(r"(\d{4})\D(\d{1,2})\D(\d{1,2})", t)
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    raise UserError(f"날짜 형식이 틀립니다: {text}  (예: 2026-09-30)")


def day_dt(text):
    return dt.datetime.combine(parse_day(text), dt.time())


def conv(c, v):
    """화면 글자 -> DB 값. 빈칸은 None, 형식 오류는 ValueError"""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if v == "":
            return None
    t = c["type"]
    if t == "NUM":
        try:
            f = float(str(v).replace(",", ""))
        except ValueError:
            raise ValueError("숫자 아님")
        return int(f) if f.is_integer() else f
    if t == "DATE":
        s = str(v).strip().replace(".", "-").replace("/", "-").rstrip("-")
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y%m%d"):
            try:
                return dt.datetime.strptime(s, fmt)
            except ValueError:
                pass
        raise ValueError("날짜 아님")
    return str(v).strip()


def make_defs(desc):
    """조회 결과 컬럼 정보 -> {컬럼: {type, editable, req, maxlen, dbtype}}"""
    defs = {}
    for d in desc:
        name, t, isize, null_ok = d[0], d[1], d[3], d[6]
        tn = str(getattr(t, "name", t)).upper()
        if any(w in tn for w in ("LOB", "RAW", "LONG", "ROWID")):
            typ = None
        elif "CHAR" in tn:
            typ = "TEXT"
        elif any(w in tn for w in ("NUMBER", "BINARY_", "FLOAT", "INTEGER", "DECIMAL")):
            typ = "NUM"
        elif "DATE" in tn or "TIMESTAMP" in tn:
            typ = "DATE"
        else:
            typ = None
        defs[name] = {"type": typ or "TEXT", "editable": typ is not None and not name.endswith("__"),
                      "req": null_ok is False, "maxlen": int(isize or 0) if typ == "TEXT" else 0,
                      "dbtype": tn.replace("DB_TYPE_", "")}
    return defs


def table_meta(conn, tbl):
    owner, t = tbl.split(".") if "." in tbl else (None, tbl)
    cur = conn.cursor()
    cur.execute("SELECT COLUMN_NAME, DATA_TYPE, NULLABLE, DATA_DEFAULT, DATA_LENGTH FROM ALL_TAB_COLUMNS "
                "WHERE OWNER = NVL(:o, USER) AND TABLE_NAME = :t ORDER BY COLUMN_ID", o=owner, t=t)
    return {n: {"type": ty, "notnull": nl == "N", "hasdef": str(d or "").strip().upper() not in ("", "NULL"),
                "len": ln} for n, ty, nl, d, ln in cur.fetchall()}


def fit_type(meta, col, v):
    if v is None:
        return None
    t = meta.get(col, {}).get("type", "")
    if t == "NUMBER" and isinstance(v, str):
        f = float(v)
        return int(f) if f.is_integer() else f
    if t.startswith(("VARCHAR", "CHAR", "NVARCHAR")) and not isinstance(v, str):
        return show(v)
    return v


def insert_dict(cur, tbl, meta, vals):
    vals = {c: fit_type(meta, c, v) for c, v in vals.items() if c in meta and v is not None}
    miss = [c for c, m in meta.items() if m["notnull"] and not m["hasdef"] and c not in vals]
    if miss:
        raise RuntimeError(f"{tbl}: 반드시 필요한 칸이 비었습니다 - {', '.join(miss)}")
    cols = list(vals)
    cur.execute(f"INSERT INTO {tbl} ({', '.join(ident(c, '컬럼') for c in cols)}) "
                f"VALUES ({', '.join(':' + str(i) for i in range(1, len(cols) + 1))})", [vals[c] for c in cols])
    if cur.rowcount != 1:
        raise RuntimeError(f"{tbl}: 등록되지 않았습니다.")
    return vals


def jsonable(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return show(v)
    if isinstance(v, dict):
        return {k: jsonable(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [jsonable(x) for x in v]
    if isinstance(v, (str, int, float, bool)) or v is None:
        return v
    return cell_text(v)


_log_lock = threading.Lock()


def append_log(rows):
    with _log_lock:
        if os.path.exists(LOG_PATH):
            with open(LOG_PATH, encoding="utf-8-sig") as f:
                head = next(csv.reader(f), [])
            if head != LOG_COLS:
                os.replace(LOG_PATH, LOG_PATH.replace(".csv", "_v1.csv"))
        new = not os.path.exists(LOG_PATH)
        with open(LOG_PATH, "a", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            if new:
                w.writerow(LOG_COLS)
            w.writerows(rows)


def now_s():
    return f"{dt.datetime.now():%Y-%m-%d %H:%M:%S}"


def reg_user():
    return (CFG.get("web_reg_user") or "").strip() or "PYWEB"


# ==============================================================
# 오라클 접속
# ==============================================================
_client_lock = threading.Lock()
_client_ready = False


def client_dirs(lib):
    import glob
    cands = [lib] if lib else []
    for pat in (os.path.join(BASE, "instantclient*"), os.path.join(BASE, "*", "instantclient*"),
                os.path.join(os.path.dirname(BASE), "instantclient*"),
                r"C:\oracle\instantclient*", r"C:\instantclient*", r"D:\instantclient*"):
        cands += sorted(glob.glob(pat), reverse=True)
    for d in os.environ.get("PATH", "").split(os.pathsep):
        if d and os.path.isfile(os.path.join(d, "oci.dll")):
            cands.append(d)
    out = []
    for d in cands:
        if d and os.path.isdir(d) and d not in out:
            out.append(d)
    return out


def init_client():
    global _client_ready
    with _client_lock:
        if _client_ready or not oracledb.is_thin_mode():
            _client_ready = True
            return
        lib = (CFG["oracle"].get("lib_dir") or "").strip()
        tried = []
        for d in client_dirs(lib):
            try:
                oracledb.init_oracle_client(lib_dir=d)
                _client_ready = True
                return
            except Exception as e:
                tried.append(f"  {d}\n    → {str(e).splitlines()[0][:150]}")
        raise RuntimeError(
            "Oracle Instant Client를 찾지 못했습니다 (11g 접속에 꼭 필요).\n"
            "해결: instantclient_19_xx 폴더를 server.py 와 같은 폴더나 C:\\ 에 두세요.\n"
            "그래도 안 되면 Visual C++ 재배포 패키지(2017 이상, x64)를 설치하세요.\n"
            + ("찾아본 곳:\n" + "\n".join(tried) if tried else f"찾아본 곳: {lib or '설정 없음'}, {BASE}, C:\\"))


def connect():
    if oracledb is None:
        raise RuntimeError("서버 PC에 oracledb가 없습니다.  명령 프롬프트에서: python -m pip install oracledb")
    init_client()
    try:
        return oracledb.connect(user=DB_USER, password=DB_PASS, dsn=CFG["oracle"]["dsn"])
    except Exception as e:
        raise RuntimeError(f"MES 접속 실패\n{e}")


def run_db(fn, *a, **kw):
    conn = connect()
    try:
        return fn(conn, *a, **kw)
    finally:
        try:
            conn.close()
        except Exception:
            pass


# ==============================================================
# 작업지시 조회 / 표에서 고치기 / 삭제
# ==============================================================
_comments = {}


def wo_where(f):
    conds, b = [], {}
    d1, d2 = day_dt(f.get("d1")), day_dt(f.get("d2"))
    conds.append("T.WORK_ORDER_DATE >= :d1 AND T.WORK_ORDER_DATE < :d2")
    b["d1"], b["d2"] = d1, d2 + dt.timedelta(days=1)
    mc = str(f.get("mc") or "").strip().upper()
    if mc:
        conds.append("UPPER(T.MACHINE_CODE) LIKE :mc")
        b["mc"] = f"%{mc}%"
    it = str(f.get("item") or "").strip().upper()
    if it:
        conds.append("(UPPER(T.ITEM_CODE) LIKE :it OR UPPER(T.ITEM_NAME) LIKE :it)")
        b["it"] = f"%{it}%"
    if f.get("nomold"):
        conds.append("NVL(TRIM(T.MOLD_CODE), '*') = '*'")
    return conds, b


def head(col):
    nm = WO_NAMES.get(col) or _comments.get(col) or ""
    return f"{nm} ({col})" if nm and not col.endswith("__") else (nm or col)


def wo_query(conn, f):
    conds, binds = wo_where(f)
    tbl = wo_tbl()
    extra = (f"(SELECT MAX(X.MOLD_CODE) FROM {link_tbl()} X WHERE X.ORGANIZATION_ID = T.ORGANIZATION_ID "
             f"AND X.ITEM_CODE = T.ITEM_CODE) AS LINK_MOLD__")
    sql = f"SELECT ROWIDTOCHAR(T.ROWID) AS RID__, T.*, {extra} FROM {tbl} T"
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += " ORDER BY T.WORK_ORDER_DATE, T.MACHINE_CODE, T.WORK_ORDER_NO"
    lim = int(CFG.get("limit") or 5000)
    sql = f"SELECT * FROM ({sql}) WHERE ROWNUM <= {lim}"
    cur = conn.cursor()
    cur.execute(sql, binds)
    rows, desc = cur.fetchall(), cur.description
    if not _comments:
        owner, tname = tbl.split(".") if "." in tbl else (None, tbl)
        try:
            c2 = conn.cursor()
            c2.execute("SELECT COLUMN_NAME, COMMENTS FROM ALL_COL_COMMENTS WHERE OWNER = NVL(:o, USER) "
                       "AND TABLE_NAME = :t", o=owner, t=tname)
            _comments.update({a: (b or "").strip() for a, b in c2.fetchall()})
        except Exception:
            pass
    names = [d[0] for d in desc]
    defs = make_defs(desc[1:])
    order = [names.index(c) for c in WO_FIRST if c in names[1:]]
    order += [i for i in range(1, len(names)) if i not in order]
    cols = [names[i] for i in order]
    data = [[r[i] for i in order] for r in rows]
    return cols, data, [r[0] for r in rows], defs, lim


def wo_tags(cols, data):
    need = ("WORK_ORDER_DATE", "MACHINE_CODE", "ITEM_CODE")
    dups = {}
    has = all(c in cols for c in need)
    if has:
        idx = [cols.index(c) for c in need]
        for r in data:
            k = tuple(show(r[i]) for i in idx)
            dups[k] = dups.get(k, 0) + 1
    tags = []
    for r in data:
        mold = show(r[cols.index("MOLD_CODE")]).strip() if "MOLD_CODE" in cols else ""
        link = show(r[cols.index("LINK_MOLD__")]).strip() if "LINK_MOLD__" in cols else ""
        if mold in ("", "*"):
            tags.append("err")
        elif link and mold != link:
            tags.append("warn")
        elif has and dups.get(tuple(show(r[cols.index(x)]) for x in need), 0) > 1:
            tags.append("dup")
        else:
            tags.append("")
    return tags


def api_wo_search(body, user):
    cols, data, rids, defs, lim = run_db(wo_query, body)
    tags = wo_tags(cols, data)
    parts = []
    for t, txt in (("err", "금형 없음(*)"), ("warn", "금형이 현재 연결과 다름"), ("dup", "같은 날·설비·품번 여러 건")):
        if tags.count(t):
            parts.append(f"{txt} {tags.count(t)}건")
    return {"cols": cols, "heads": [head(c) for c in cols],
            "names": {c: (WO_NAMES.get(c) or _comments.get(c, "")) for c in cols},
            "defs": {c: defs.get(c, {"type": "TEXT", "editable": False, "req": False, "maxlen": 0, "dbtype": ""})
                     for c in cols},
            "rows": [[cell_text(v) for v in r] for r in data], "rids": rids, "tags": tags,
            "summary": " / ".join(parts), "limit": lim, "hit_limit": len(data) >= lim,
            "locked": WO_LOCKED, "sensitive": WO_SENSITIVE, "time": f"{dt.datetime.now():%H:%M:%S}"}


def row_by_rid(conn, tbl, rid):
    cur = conn.cursor()
    cur.execute(f"SELECT * FROM {tbl} WHERE ROWID = CHARTOROWID(:r)", r=rid)
    r = cur.fetchone()
    return (dict(zip([d[0] for d in cur.description], r)), make_defs(cur.description)) if r else (None, {})


def label_of(row):
    return " / ".join(show(row.get(c)).strip() for c in WO_LABEL if c in row)


def api_wo_save_grid(body, user):
    """표에서 고친 칸 저장. edits: [{rid, changes: {col: {old, new}}}]
    old(화면에 보였던 값)가 지금 MES 값과 다르면 전체 취소 (그 사이 누가 바꿈)"""
    edits = body.get("edits") or []
    if not edits:
        raise UserError("수정한 내용이 없습니다.")
    tbl = wo_tbl()
    reg = reg_user()

    def work(conn):
        cur = conn.cursor()
        done = []
        try:
            for e in edits:
                rid = e["rid"]
                row, defs = row_by_rid(conn, tbl, rid)
                if row is None:
                    raise UserError("이미 없어진 줄이 있습니다. 다시 조회하세요.")
                lab = label_of(row)
                sets, where, b = [], ["ROWID = CHARTOROWID(:rid)"], {"rid": rid}
                log = []
                for k, (col, ch) in enumerate((e.get("changes") or {}).items()):
                    col = ident(col, "컬럼")
                    if col in WO_LOCKED:
                        raise UserError(f"{lab} · {head(col)}: {WO_LOCKED[col]}")
                    d = defs.get(col)
                    if not d or not d["editable"]:
                        raise UserError(f"{lab} · {head(col)}: 이 칸은 고칠 수 없습니다.")
                    cur_v = row.get(col)
                    if cell_text(cur_v) != str(ch.get("old") or ""):
                        raise UserError(f"{lab} · {head(col)}: 그 사이 다른 곳에서 바뀌었습니다 "
                                        f"(지금 값 '{cell_text(cur_v)}'). 다시 조회하세요.")
                    try:
                        new = conv(d, ch.get("new"))
                        if new is None and d["req"]:
                            raise ValueError("필수 칸이라 비울 수 없음")
                        if d["type"] == "TEXT" and d["maxlen"] and new is not None and blen(new) > d["maxlen"]:
                            raise ValueError(f"너무 김 ({blen(new)}/{d['maxlen']}byte)")
                        if d["type"] == "NUM" and new is not None and new < 0:
                            raise ValueError("음수는 넣을 수 없습니다")
                    except ValueError as ex:
                        raise UserError(f"{lab} · {head(col)}: {ex}")
                    sets.append(f"{col} = :n{k}")
                    b[f"n{k}"] = new
                    if cur_v is None:
                        where.append(f"{col} IS NULL")
                    else:
                        where.append(f"{col} = :o{k}")
                        b[f"o{k}"] = cur_v
                    log.append(f"{col}: {cell_text(cur_v)} -> {cell_text(new)}")
                if not sets:
                    continue
                changed = set(e.get("changes") or {})
                for c, d in defs.items():
                    if c in changed or not d["editable"]:
                        continue
                    if d["type"] == "DATE" and RE_UPD_DATE.search(c):
                        sets.append(f"{ident(c)} = SYSDATE")
                    elif d["type"] == "TEXT" and RE_UPD_USER.search(c):
                        sets.append(f"{ident(c)} = :u_{c.lower()}")
                        b[f"u_{c.lower()}"] = reg
                cur.execute(f"UPDATE {tbl} SET {', '.join(sets)} WHERE {' AND '.join(where)}", b)
                if cur.rowcount != 1:
                    raise UserError(f"{lab}: 그 사이 다른 곳에서 바뀌었거나 없어진 줄입니다. 다시 조회하세요.")
                done.append((lab, log))
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        return done
    try:
        done = run_db(work)
    except UserError as e:
        raise UserError(f"{e}\n\n(전체 취소 - MES는 그대로)")
    except Exception as e:
        raise RuntimeError(f"수정 중 오류 -> 전체 취소 (MES는 그대로)\n\n{ora_hint(e)}")
    t = now_s()
    append_log([[t, user, "작업지시수정(웹)", tbl, lab, "; ".join(log)] for lab, log in done])
    return {"ok": True, "n": len(done)}


def delete_block(row):
    for c in DONE_COLS:
        v = row.get(c)
        if v not in (None, 0, 0.0):
            return f"실적이 있습니다 ({c}={show(v)}). 실적 있는 작업지시는 지우지 않습니다."
    return None


def usage_of_value(conn, col, value, skip):
    """col 값을 쓰는 다른 테이블 (표마다 1건 찾으면 멈춤, 오래 걸리는 표는 15초 뒤 넘어감)"""
    owner = (CFG.get("schema") or "").strip().upper() or None
    cur = conn.cursor()
    cur.execute("SELECT c.TABLE_NAME, c.COLUMN_NAME, NVL(t.NUM_ROWS, 0) FROM ALL_TAB_COLUMNS c JOIN ALL_TABLES t "
                "ON t.OWNER = c.OWNER AND t.TABLE_NAME = c.TABLE_NAME "
                "WHERE c.OWNER = NVL(:o, USER) AND c.COLUMN_NAME = :c ORDER BY 3, 1", o=owner, c=col)
    tables = [(t, c) for t, c, n in cur.fetchall() if t not in skip]
    used = []
    for tname, cname in tables:
        try:
            conn.call_timeout = 15000
            c2 = conn.cursor()
            c2.execute(f"SELECT COUNT(*) FROM {qualify(tname)} WHERE {ident(cname)} = :v AND ROWNUM = 1", v=value)
            if c2.fetchone()[0]:
                used.append(tname)
        except Exception:
            try:
                conn.ping()
            except Exception:
                break
        finally:
            try:
                conn.call_timeout = 0
            except Exception:
                pass
    return used


def api_wo_delete_check(body, user):
    rids = body.get("rids") or []
    if not rids:
        raise UserError("삭제할 줄을 선택하세요.")
    tbl = wo_tbl()

    def work(conn):
        lines, blocked, used = [], [], []
        for rid in rids:
            row, _ = row_by_rid(conn, tbl, rid)
            if row is None:
                blocked.append("이미 없어진 줄이 있습니다. 다시 조회하세요.")
                continue
            lab = label_of(row)
            lines.append(lab)
            why = delete_block(row)
            if why:
                blocked.append(f"{lab}: {why}")
        if not blocked:
            for rid in rids:
                row, _ = row_by_rid(conn, tbl, rid)
                no = show(row.get("WORK_ORDER_NO")).strip()
                if no:
                    for t in usage_of_value(conn, "WORK_ORDER_NO", no, {tbl.split(".")[-1]}):
                        used.append(f"{no} → {t}: 기록 있음")
        return lines, blocked, used
    lines, blocked, used = run_db(work)
    return {"lines": lines, "blocked": blocked, "used": used}


def api_wo_delete(body, user):
    rids = body.get("rids") or []
    tbl = wo_tbl()
    if not rids:
        raise UserError("삭제할 줄을 선택하세요.")
    if body.get("used") and (body.get("typed") or "").strip() != "삭제":
        raise UserError("다른 테이블에 기록이 있어 '삭제'라고 입력해야 지울 수 있습니다.")

    def work(conn):
        cur = conn.cursor()
        rows = []
        try:
            for rid in rids:
                row, _ = row_by_rid(conn, tbl, rid)
                if row is None:
                    raise UserError("이미 없어진 줄입니다. 다시 조회하세요.")
                why = delete_block(row)
                if why:
                    raise UserError(f"{label_of(row)}: {why}")
                rows.append(row)
            with open(BACKUP_PATH, "a", encoding="utf-8") as f:       # 지우기 전에 백업
                for row in rows:
                    f.write(json.dumps({"time": now_s(), "user": user, "table": tbl,
                                        "row": {k: show(v) for k, v in row.items()}}, ensure_ascii=False) + "\n")
            for rid, row in zip(rids, rows):
                cur.execute(f"DELETE FROM {tbl} WHERE ROWID = CHARTOROWID(:1)", [rid])
                if cur.rowcount != 1:
                    raise UserError(f"{label_of(row)}: 이미 없어진 줄입니다. 다시 조회하세요.")
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        return rows
    try:
        rows = run_db(work)
    except UserError:
        raise
    except Exception as e:
        raise RuntimeError(f"삭제 중 오류 -> 전체 취소 (MES는 그대로)\n\n{ora_hint(e)}")
    append_log([[now_s(), user, "작업지시삭제(웹)", tbl, label_of(r),
                 "; ".join(f"{k}={cell_text(v)}" for k, v in r.items() if v is not None)] for r in rows])
    return {"ok": True, "n": len(rows)}


def api_wo_export(body, user):
    if openpyxl is None:
        raise RuntimeError("서버 PC에 openpyxl이 없습니다.  python -m pip install openpyxl")
    cols, data, _r, _d, _l = run_db(wo_query, body)
    keep = [i for i, c in enumerate(cols) if not c.endswith("__")]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "작업지시"
    ws.append([cols[i] for i in keep])
    ws.append([WO_NAMES.get(cols[i]) or _comments.get(cols[i], "") for i in keep])
    for c in ws[1]:
        c.font = Font(bold=True)
    for c in ws[2]:
        c.font = Font(size=9, color="808080")
    for r in data:
        ws.append([r[i] if isinstance(r[i], (int, float, dt.datetime)) else cell_text(r[i], 2000) for i in keep])
    for j, i in enumerate(keep, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(30, len(cols[i]) + 4))
    ws.freeze_panes = "B3"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue(), f"작업지시_{dt.datetime.now():%Y%m%d_%H%M}.xlsx"


# ==============================================================
# 작업지시 고치기 창 (기존 WoForm)
# ==============================================================
def api_form(body, user):
    rid = body.get("rid")
    tbl = wo_tbl()

    def work(conn):
        meta = table_meta(conn, tbl)
        cur = conn.cursor()
        cur.execute(f"SELECT MACHINE_CODE, MAX(MACHINE_NAME) FROM {tbl} WHERE WORK_ORDER_DATE >= SYSDATE - 365 "
                    f"AND MACHINE_CODE IS NOT NULL GROUP BY MACHINE_CODE ORDER BY MACHINE_CODE")
        machines = [[show(a), show(b)] for a, b in cur.fetchall()]
        shifts, st = [], {}
        if "WORK_SHIFT" in meta:
            cur.execute(f"SELECT WORK_SHIFT, COUNT(*) FROM {tbl} WHERE WORK_ORDER_DATE >= SYSDATE - 180 "
                        f"GROUP BY WORK_SHIFT ORDER BY 2 DESC")
            shifts = [show(a) for a, _ in cur.fetchall() if a is not None]
        for c in ("WORK_ORDER_STATUS", "PLAN_STATUS"):
            if c in meta:
                cur.execute(f"SELECT {c}, COUNT(*) FROM {tbl} WHERE WORK_ORDER_DATE >= SYSDATE - 180 "
                            f"GROUP BY {c} ORDER BY 2 DESC")
                st[c] = [show(a) for a, _ in cur.fetchall() if a is not None]
        row, _ = row_by_rid(conn, tbl, rid)
        return meta, machines, shifts, st, row
    meta, machines, shifts, st, row = run_db(work)
    if row is None:
        raise UserError("선택한 작업지시를 MES에서 찾지 못했습니다. 다시 조회하세요.")
    done = {c: show(row.get(c)) for c in DONE_COLS if row.get(c) not in (None, 0)}
    return {"machines": machines, "shifts": shifts, "status": st,
            "has_shift": "WORK_SHIFT" in meta,
            "row": {k: cell_text(v) for k, v in row.items()},
            "done": done}


def api_item_lookup(body, user):
    it = str(body.get("item") or "").strip()
    if not it:
        return {"n": 0}
    tbl = wo_tbl()

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"SELECT MAX(ITEM_NAME), COUNT(*) FROM {mst_tbl()} WHERE ITEM_CODE = :i", i=it)
        name, n = cur.fetchone()
        cur.execute(f"SELECT MAX(MOLD_CODE) FROM {link_tbl()} WHERE ITEM_CODE = :i", i=it)
        mold = cur.fetchone()[0]
        cur.execute(f"SELECT MACHINE_CODE, PLAN_QTY FROM (SELECT MACHINE_CODE, PLAN_QTY FROM {tbl} "
                    f"WHERE ITEM_CODE = :i ORDER BY WORK_ORDER_DATE DESC, WORK_ORDER_NO DESC) WHERE ROWNUM = 1", i=it)
        last = cur.fetchone()
        return name, n, mold, last
    name, n, mold, last = run_db(work)
    return {"n": n or 0, "name": show(name), "mold": show(mold),
            "last_mc": show(last[0]) if last else "", "last_qty": show(last[1]) if last else ""}


def api_item_find(body, user):
    q = str(body.get("q") or "").strip().upper()

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"SELECT * FROM (SELECT M.ITEM_CODE, MAX(M.ITEM_NAME), "
                    f"(SELECT MAX(X.MOLD_CODE) FROM {link_tbl()} X WHERE X.ITEM_CODE = M.ITEM_CODE) "
                    f"FROM {mst_tbl()} M WHERE UPPER(M.ITEM_CODE) LIKE :q OR UPPER(M.ITEM_NAME) LIKE :q "
                    f"GROUP BY M.ITEM_CODE ORDER BY M.ITEM_CODE) WHERE ROWNUM <= 300", q=f"%{q}%")
        return cur.fetchall()
    return {"rows": [[show(a), show(b), show(m)] for a, b, m in run_db(work)]}


def wo_derive(conn, vals):
    """설비명·품명·규격은 MES 함수로 다시 채움"""
    out = {}
    cur = conn.cursor()
    org, mc, it = vals.get("ORGANIZATION_ID"), vals.get("MACHINE_CODE"), vals.get("ITEM_CODE")
    if mc is not None:
        try:
            cur.execute("SELECT F_GET_MACHINE_NAME(:1, :2) FROM DUAL", [mc, org])
            r = cur.fetchone()[0]
            if r:
                out["MACHINE_NAME"] = r
        except Exception:
            pass
    if it is not None:
        try:
            cur.execute("SELECT F_GET_ITEM_NAME(:o, :i), F_GET_ITEM_SPEC(:o, :i) FROM DUAL", o=org, i=it)
            a, b = cur.fetchone()
            if a:
                out["ITEM_NAME"] = a
            if b:
                out["ITEM_SPEC"] = b
        except Exception:
            pass
    return out


def api_form_save(body, user):
    """고치기 창 저장. confirm=false 면 바뀔 내용·경고만 돌려줌, true 면 저장.
    stamp(창을 열 때 값)가 지금 MES 값과 다르면 저장 안 함."""
    rid, vin, stamp, confirm = body.get("rid"), body.get("values") or {}, body.get("stamp") or {}, body.get("confirm")
    tbl = wo_tbl()

    def work(conn):
        meta = table_meta(conn, tbl)
        o, _ = row_by_rid(conn, tbl, rid)
        if o is None:
            raise UserError("작업지시를 찾지 못했습니다. 다시 조회하세요.")
        changed_out = [c for c, s in stamp.items() if c in o and cell_text(o[c]) != s]
        if changed_out:
            raise UserError("창을 연 뒤에 다른 곳(MES 화면·현장)에서 이 작업지시가 바뀌었습니다.\n"
                            f"바뀐 칸: {', '.join(changed_out[:8])}\n창을 닫고 다시 조회한 뒤 고치세요. (저장 안 함)")

        def cast(c, text):
            m = meta.get(c, {})
            t = str(text or "").strip()
            if t == "":
                return None
            if m.get("type") == "NUMBER":
                try:
                    f = float(t.replace(",", ""))
                except ValueError:
                    raise UserError(f"{WO_NAMES.get(c, c)}: 숫자를 넣으세요 ({t})")
                return int(f) if f.is_integer() else f
            if m.get("type") == "DATE":
                return day_dt(t)
            if m.get("len") and blen(t) > int(m["len"]):
                raise UserError(f"{WO_NAMES.get(c, c)}: 너무 깁니다 (최대 {m['len']}byte)")
            return t

        day = parse_day(vin.get("WORK_ORDER_DATE"))
        od = o.get("WORK_ORDER_DATE")
        v = {"WORK_ORDER_DATE": od if od and od.date() == day else dt.datetime.combine(day, dt.time()),
             "ITEM_CODE": cast("ITEM_CODE", vin.get("ITEM_CODE")),
             "MACHINE_CODE": cast("MACHINE_CODE", vin.get("MACHINE_CODE")),
             "PLAN_QTY": cast("PLAN_QTY", vin.get("PLAN_QTY"))}
        if "WORK_SHIFT" in meta and "WORK_SHIFT" in vin:
            v["WORK_SHIFT"] = cast("WORK_SHIFT", vin.get("WORK_SHIFT"))
        for c in ("WORK_ORDER_STATUS", "PLAN_STATUS"):
            if c in meta and c in vin:
                v[c] = cast(c, vin.get(c))
        for c, lab in (("ITEM_CODE", "품번"), ("MACHINE_CODE", "설비"), ("PLAN_QTY", "계획수량")):
            if v[c] is None:
                raise UserError(f"{lab}을(를) 넣으세요.")
        if v["PLAN_QTY"] <= 0:
            raise UserError("계획수량은 1 이상이어야 합니다.")

        ch = {c: (o.get(c), x) for c, x in v.items() if cell_text(o.get(c)) != cell_text(x)}
        if not ch:
            raise UserError("바뀐 내용이 없습니다.")
        if "WORK_ORDER_DATE" in ch and od:
            delta = v["WORK_ORDER_DATE"] - dt.datetime.combine(od.date(), dt.time())
            for c in ("START_DATE", "END_DATE"):
                if isinstance(o.get(c), dt.datetime):
                    ch[c] = (o[c], o[c] + delta)
        if "ITEM_CODE" in ch and "MODEL_NAME" in o and o.get("MODEL_NAME") == o.get("ITEM_CODE"):
            ch["MODEL_NAME"] = (o["MODEL_NAME"], v["ITEM_CODE"])
        warn = []
        done = {c: o.get(c) for c in DONE_COLS if o.get(c) not in (None, 0)}
        key3 = set(ch) & {"ITEM_CODE", "MACHINE_CODE", "WORK_ORDER_DATE"}
        if done and key3:
            warn.append("이 작업지시는 실적이 있습니다 (" + ", ".join(f"{c}={show(x)}" for c, x in done.items()) +
                        "). 품번·설비·날짜를 바꾸면 실적 집계가 달라집니다.")
        cur = conn.cursor()
        if key3:
            cur.execute(f"SELECT COUNT(*) FROM {mst_tbl()} WHERE ITEM_CODE = :i", i=v["ITEM_CODE"])
            if not cur.fetchone()[0]:
                raise UserError(f"품목 마스터에 없는 품번입니다: {v['ITEM_CODE']}")
            cur.execute(f"SELECT MAX(MOLD_CODE) FROM {link_tbl()} WHERE ITEM_CODE = :i", i=v["ITEM_CODE"])
            if not cur.fetchone()[0]:
                warn.append(f"품번 {v['ITEM_CODE']}에 연결된 금형이 없어 금형이 '*'로 들어갑니다.")
            d = v["WORK_ORDER_DATE"]
            cur.execute(f"SELECT WORK_ORDER_NO FROM {tbl} WHERE WORK_ORDER_DATE >= :a AND WORK_ORDER_DATE < :b "
                        f"AND MACHINE_CODE = :m AND ITEM_CODE = :i AND ROWID <> CHARTOROWID(:r)",
                        a=dt.datetime.combine(d.date(), dt.time()),
                        b=dt.datetime.combine(d.date(), dt.time()) + dt.timedelta(days=1),
                        m=v["MACHINE_CODE"], i=v["ITEM_CODE"], r=rid)
            same = [show(x[0]) for x in cur.fetchall()]
            if same:
                warn.append(f"같은 날·설비·품번 작업지시가 이미 있습니다: {', '.join(same[:5])}")
        if "ITEM_CODE" in ch:
            der = wo_derive(conn, {**o, **v})
            for c in ("ITEM_NAME", "ITEM_SPEC"):
                if c in der and c in o and cell_text(der[c]) != cell_text(o.get(c)):
                    ch[c] = (o.get(c), der[c])
        lines = [{"col": c, "name": WO_NAMES.get(c, c), "old": cell_text(a), "new": cell_text(b)}
                 for c, (a, b) in ch.items()]
        if not confirm:
            return {"preview": True, "no": show(o.get("WORK_ORDER_NO")), "lines": lines, "warn": warn}

        reg = reg_user()
        auto = {}
        for c, m in meta.items():
            if m["type"] == "DATE" and RE_UPD_DATE.search(c):
                auto[c] = "SYSDATE"
            elif m["type"].startswith("VARCHAR") and RE_UPD_USER.search(c):
                auto[c] = reg
        sets, where, b = [], ["ROWID = CHARTOROWID(:rid)"], {"rid": rid}
        for i, (c, (a, x)) in enumerate(ch.items()):
            sets.append(f"{ident(c)} = :n{i}")
            b[f"n{i}"] = x
            if a is None:
                where.append(f"{ident(c)} IS NULL")
            else:
                where.append(f"{ident(c)} = :o{i}")
                b[f"o{i}"] = a
        for j, (c, how) in enumerate(auto.items()):
            if c in ch:
                continue
            if how == "SYSDATE":
                sets.append(f"{ident(c)} = SYSDATE")
            else:
                sets.append(f"{ident(c)} = :u{j}")
                b[f"u{j}"] = how
        try:
            cur.execute(f"UPDATE {tbl} SET {', '.join(sets)} WHERE {' AND '.join(where)}", b)
            if cur.rowcount != 1:
                raise UserError("그 사이 다른 곳(MES 화면·현장)에서 바뀌었거나 없어진 작업지시입니다.\n"
                                "다시 조회한 뒤 고치세요. (저장 안 함)")
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        new, _ = row_by_rid(conn, tbl, rid)
        return {"ok": True, "no": show(o.get("WORK_ORDER_NO")), "lines": lines,
                "mold": show((new or {}).get("MOLD_CODE")), "mname": show((new or {}).get("MACHINE_NAME"))}
    try:
        res = run_db(work)
    except UserError:
        raise
    except Exception as e:
        raise RuntimeError(ora_hint(e))
    if res.get("ok"):
        append_log([[now_s(), user, "작업지시수정(웹)", tbl, res["no"],
                     "; ".join(f"{x['col']}: {x['old']} -> {x['new']}" for x in res["lines"])]])
    return res


# ==============================================================
# 여러 건 일괄등록 / 최근 작업지시 불러오기 (기존 WoBatch)
# ==============================================================
def api_batch_lists(body, user):
    tb = wo_tbl()

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"SELECT MACHINE_CODE, MAX(MACHINE_NAME) FROM {tb} WHERE WORK_ORDER_DATE >= SYSDATE - 730 "
                    f"AND MACHINE_CODE IS NOT NULL GROUP BY MACHINE_CODE ORDER BY MACHINE_CODE")
        machines = {show(a).strip(): show(b).strip() for a, b in cur.fetchall()}
        cur.execute(f"SELECT ITEM_CODE, MAX(ITEM_NAME), MAX(MACHINE_CODE) KEEP (DENSE_RANK LAST ORDER BY WORK_ORDER_DATE), "
                    f"COUNT(*) FROM {tb} WHERE WORK_ORDER_DATE >= SYSDATE - 730 AND ITEM_CODE IS NOT NULL "
                    f"GROUP BY ITEM_CODE ORDER BY COUNT(*) DESC")
        seen = {}
        for it, nm, mc, n in cur.fetchall():
            seen[show(it).strip()] = (show(nm).strip(), show(mc).strip())
        cur.execute(f"SELECT MACHINE_CODE, ITEM_CODE, COUNT(*), MAX(WORK_ORDER_DATE) FROM {tb} "
                    f"WHERE WORK_ORDER_DATE >= SYSDATE - 730 AND ITEM_CODE IS NOT NULL AND MACHINE_CODE IS NOT NULL "
                    f"GROUP BY MACHINE_CODE, ITEM_CODE ORDER BY MAX(WORK_ORDER_DATE) DESC")
        mc_items, item_mcs = {}, {}
        for mc, it, n, last in cur.fetchall():
            mc, it = show(mc).strip(), show(it).strip()
            mc_items.setdefault(mc, []).append([it, n, show(last)])
            item_mcs.setdefault(it.upper(), []).append([mc, n, show(last)])
        try:
            cur.execute(f"SELECT L.ITEM_CODE, (SELECT MAX(M.ITEM_NAME) FROM {mst_tbl()} M WHERE M.ITEM_CODE = L.ITEM_CODE) "
                        f"FROM {link_tbl()} L")
            for it, nm in cur.fetchall():
                it = show(it).strip()
                if it and it not in seen:
                    seen[it] = (show(nm).strip(), "")
        except Exception:
            pass
        return machines, seen, mc_items, item_mcs
    machines, seen, mc_items, item_mcs = run_db(work)
    return {"machines": machines, "items": [[k, nm, mc] for k, (nm, mc) in seen.items()],
            "mc_items": mc_items, "item_mcs": item_mcs}


_spm_cache = {"src": None, "src_done": False, "ref": None, "mtime": None}


def spm_table():
    """spm.xlsx (기존 프로그램 [⑤ SPM 관리]와 같은 파일) -> {품번 대문자: SPM}"""
    if openpyxl is None:
        return {}, "SPM 관리(openpyxl 없음)"
    try:
        if not os.path.exists(SPM_XLSX):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "SPM"
            ws.append(["품번", "SPM", "수정일", "수정자", "비고"])
            stamp = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
            for k in sorted(SPM_DEFAULT, key=str.upper):
                ws.append([k, SPM_DEFAULT[k], stamp, "초기값", ""])
            wb.save(SPM_XLSX)
        mt = os.path.getmtime(SPM_XLSX)
        if _spm_cache["ref"] is None or _spm_cache["mtime"] != mt:
            wb = openpyxl.load_workbook(SPM_XLSX, read_only=True, data_only=True)
            out = {}
            for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
                r = list(r or []) + [None] * 2
                it = str(r[0] or "").strip().upper()
                try:
                    v = float(str(r[1]).replace(",", ""))
                except (TypeError, ValueError):
                    continue
                if it and it not in out:
                    out[it] = v
            wb.close()
            _spm_cache.update(ref=out, mtime=mt)
        return _spm_cache["ref"], "SPM 관리"
    except Exception:
        return {}, "SPM 관리(읽기 실패)"


def spm_of(conn, mc, it, days=90):
    """SPM 찾기: SPM 관리표 → 생산현황판 SPM 평균 → 작업지시 ST. return (spm, 캐비티, 설명)"""
    days = max(1, min(int(days or 90), 730))
    tb = wo_tbl()
    cur = conn.cursor()
    try:
        conn.call_timeout = 10000
    except Exception:
        pass
    try:
        if not _spm_cache["src_done"]:
            owner = (CFG.get("schema") or "").strip().upper() or None
            cur.execute("SELECT TABLE_NAME, MAX(CASE WHEN COLUMN_NAME LIKE '%SPM%' THEN COLUMN_NAME END), "
                        "MAX(CASE WHEN COLUMN_NAME = 'WORK_ORDER_NO' THEN 1 END) FROM ALL_TAB_COLUMNS "
                        "WHERE OWNER = NVL(:o, USER) AND TABLE_NAME IN ('ICOM_PRODUCTION_BOARD_WO_HIST', "
                        "'ICOM_PRODUCTION_BOARD_HIST', 'ICOM_PRODUCTION_BOARD') GROUP BY TABLE_NAME", o=owner)
            for t, col, has_wo in cur.fetchall():
                if col and has_wo:
                    _spm_cache["src"] = (qualify(t), col, t)
                    break
            _spm_cache["src_done"] = True
        cav = 1
        try:
            cur.execute(f"SELECT MAX(CAVITY_QTY) FROM {link_tbl()} WHERE ITEM_CODE = :i", i=it)
            cav = int(cur.fetchone()[0] or 1) or 1
        except Exception:
            cav = 1
        ref, ref_src = spm_table()
        v = ref.get(str(it or "").strip().upper())
        if v:
            return float(v), cav, f"{ref_src} {v:g} × 캐비티 {cav}"
        if _spm_cache["src"]:
            t, col, tname = _spm_cache["src"]
            for cond, label in (("AND W.MACHINE_CODE = :m", "이 설비"), ("", "다른 설비 포함")):
                b = {"i": it}
                if cond:
                    b["m"] = mc
                cur.execute(f"SELECT AVG(V), COUNT(*) FROM (SELECT P.{col} V FROM {t} P WHERE P.WORK_ORDER_NO IN "
                            f"(SELECT W.WORK_ORDER_NO FROM {tb} W WHERE W.ITEM_CODE = :i {cond} "
                            f"AND W.WORK_ORDER_DATE >= SYSDATE - {days}) AND P.{col} > 0)", b)
                v, n = cur.fetchone()
                if v:
                    return float(v), cav, f"{tname} 최근 {days}일 평균 ({label}, {n}건) × 캐비티 {cav}"
        cur.execute(f"SELECT MAX(ST_VALUE) KEEP (DENSE_RANK LAST ORDER BY WORK_ORDER_DATE) FROM {tb} "
                    f"WHERE ITEM_CODE = :i AND MACHINE_CODE = :m AND WORK_ORDER_DATE >= SYSDATE - {days} "
                    f"AND NVL(ST_VALUE, 0) > 0", i=it, m=mc)
        st = cur.fetchone()[0]
        if not st:
            cur.execute(f"SELECT MAX(ST_VALUE) KEEP (DENSE_RANK LAST ORDER BY WORK_ORDER_DATE) FROM {tb} "
                        f"WHERE ITEM_CODE = :i AND WORK_ORDER_DATE >= SYSDATE - {days} AND NVL(ST_VALUE, 0) > 0", i=it)
            st = cur.fetchone()[0]
        if st:
            return 60.0 / float(st), 1, f"최근 {days}일 작업지시 ST {show(st)}초/개 → 60÷ST"
        return None, cav, f"최근 {days}일 SPM 자료 없음 - 계획수량 직접 입력"
    finally:
        try:
            conn.call_timeout = 0
        except Exception:
            pass


def api_batch_spm(body, user):
    mc, it = str(body.get("mc") or "").strip(), str(body.get("item") or "").strip().upper()
    if not it:
        raise UserError("품번을 넣으세요.")
    spm, cav, src = run_db(spm_of, mc, it)
    return {"spm": spm, "cav": cav, "src": src}


def api_batch_recent(body, user):
    """최근 30일에서 호기별 마지막 작업지시 1건 + SPM × 60 × 근무시간 × 캐비티로 계획수량"""
    try:
        hours = float(body.get("hours"))
        if not 0 < hours <= 24:
            raise ValueError
    except (TypeError, ValueError):
        raise UserError("근무시간은 0보다 크고 24 이하인 숫자로 넣으세요.")
    tb = wo_tbl()

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"""
            SELECT MACHINE_CODE, MACHINE_NAME, ITEM_CODE, ITEM_NAME, WORK_ORDER_DATE, WORK_ORDER_NO
              FROM (SELECT T.MACHINE_CODE, T.MACHINE_NAME, T.ITEM_CODE, T.ITEM_NAME, T.WORK_ORDER_DATE, T.WORK_ORDER_NO,
                           ROW_NUMBER() OVER (PARTITION BY T.MACHINE_CODE
                                              ORDER BY T.WORK_ORDER_DATE DESC, T.WORK_ORDER_NO DESC, T.ROWID DESC) AS RN
                      FROM {tb} T
                     WHERE T.WORK_ORDER_DATE >= TRUNC(SYSDATE) - 30
                       AND T.MACHINE_CODE IS NOT NULL AND T.ITEM_CODE IS NOT NULL)
             WHERE RN = 1
             ORDER BY WORK_ORDER_DATE DESC, MACHINE_CODE""")
        base = cur.fetchall()
        cur.execute(f"SELECT MACHINE_CODE, MAX(MACHINE_NAME) FROM {tb} WHERE WORK_ORDER_DATE >= SYSDATE - 730 "
                    f"AND MACHINE_CODE IS NOT NULL GROUP BY MACHINE_CODE")
        machines = {show(a).strip(): show(b).strip() for a, b in cur.fetchall()}
        out = []
        for mc, mname, it, iname, last_day, last_wo in base:
            mc, it = show(mc).strip(), show(it).strip().upper()
            if not mc or not it:
                continue
            try:
                spm, cav, src = spm_of(conn, mc, it, days=30)
            except Exception as e:
                spm, cav, src = None, 1, f"SPM 조회 실패: {str(e)[:50]}"
            qty = ""
            if spm:
                qty = str(int(round(float(spm) * 60 * hours * int(cav or 1) / 10.0) * 10))
            out.append({"line": machines.get(mc) or show(mname).strip() or mc, "date_t": "", "shift": "",
                        "item": it, "qty_t": qty, "prio": "", "lv": "" if qty else "warn",
                        "st": (f"최근 {last_day:%m/%d} · {src}" if last_day else src),
                        "recent_date": show(last_day)[:10], "recent_wo": show(last_wo),
                        "spm": round(float(spm), 2) if spm else None})
        return out
    return {"rows": run_db(work)}


def parse_bdate(t):
    t = str(t or "").strip().split(" ")[0]
    if not t:
        return None
    if re.fullmatch(r"\d{1,2}[/.-]\d{1,2}", t):
        m, d = re.split(r"[/.-]", t)
        return dt.datetime(dt.date.today().year, int(m), int(d))
    if re.fullmatch(r"\d{5}(\.0)?", t):
        return dt.datetime(1899, 12, 30) + dt.timedelta(days=int(float(t)))
    return day_dt(t)


def batch_check(conn, rows, d_def_text, shift_def):
    """기존 WoBatch.check 와 같은 검사. rows(dict 목록)에 결과를 채워서 돌려줌. return org"""
    tbl = wo_tbl()
    bom = qualify(CFG.get("bom_table") or "ICOM_ITEM_CHILD")
    d_def = day_dt(d_def_text)
    for r in rows:
        r["lv"], r["st"] = "", ""
        for k in ("line", "date_t", "shift", "item", "qty_t", "prio"):
            r[k] = str(r.get(k) or "").strip()
        try:
            dd = parse_bdate(r["date_t"]) if r["date_t"] else d_def
            r["date"] = dd
            r["shift_v"] = str(r["shift"] or shift_def or "1").strip()
        except Exception:
            r["date"], r["lv"], r["st"] = None, "err", "계획일자 형식 오류"
        try:
            q = float(r["qty_t"].replace(",", ""))
            if q <= 0:
                raise ValueError
            r["qty"] = int(q) if q.is_integer() else q
        except Exception:
            r["qty"], r["lv"] = None, "err"
            r["st"] = (r["st"] + " / " if r["st"] else "") + "계획수량 오류"
        r["item"] = r["item"].upper()
    cur = conn.cursor()
    cur.execute(f"SELECT ORGANIZATION_ID FROM (SELECT ORGANIZATION_ID, COUNT(*) N FROM {tbl} "
                f"WHERE WORK_ORDER_DATE >= SYSDATE - 90 GROUP BY ORGANIZATION_ID ORDER BY N DESC) WHERE ROWNUM = 1")
    org = (cur.fetchone() or [1])[0]
    cur.execute(f"SELECT MACHINE_CODE, MAX(MACHINE_NAME) FROM {tbl} WHERE WORK_ORDER_DATE >= SYSDATE - 730 "
                f"AND MACHINE_CODE IS NOT NULL GROUP BY MACHINE_CODE")
    machines = {show(a).strip(): show(b).strip() for a, b in cur.fetchall()}
    info = {}
    for it in {r["item"] for r in rows if r["item"]}:
        cur.execute(f"SELECT MAX(ITEM_NAME), COUNT(*) FROM {mst_tbl()} WHERE ITEM_CODE = :i AND ORGANIZATION_ID = :o",
                    i=it, o=org)
        name, n = cur.fetchone()
        cur.execute(f"SELECT MAX(MOLD_CODE) FROM {link_tbl()} WHERE ITEM_CODE = :i", i=it)
        mold = cur.fetchone()[0]
        cur.execute(f"SELECT COUNT(*) FROM {bom} WHERE ITEM_CODE = :i AND NVL(UNIT_PER_QTY, 0) > 0", i=it)
        b = cur.fetchone()[0]
        info[it] = (n, show(name), show(mold), b)
    dups = {}
    for r in rows:
        if r.get("date") and r["item"]:
            cur.execute(f"SELECT MAX(WORK_ORDER_NO) FROM {tbl} WHERE WORK_ORDER_DATE >= :a AND WORK_ORDER_DATE < :b "
                        f"AND ITEM_CODE = :i", a=r["date"], b=r["date"] + dt.timedelta(days=1), i=r["item"])
            dups[(r["date"], r["item"])] = show(cur.fetchone()[0])
    norm = lambda x: re.sub(r"\s+", "", str(x or "")).upper()  # noqa: E731
    by_name, by_num = {}, {}
    for code, name in machines.items():
        by_name.setdefault(norm(name), []).append(code)
        m = re.search(r"(\d+)\s*호기", name)
        if m:
            by_num.setdefault(int(m.group(1)), []).append(code)
    seen = set()
    for r in rows:
        if r["lv"] == "err":
            continue
        msgs, lv = [], "ok"
        ln = r["line"]
        num_codes = [c for c in machines if c.isdigit() and ln.isdigit() and int(c) == int(ln)]
        if ln in machines:
            r["mc"] = ln
        elif len(by_name.get(norm(ln), [])) == 1:
            r["mc"] = by_name[norm(ln)][0]
        elif ln.isdigit() and len(by_num.get(int(ln), [])) == 1:
            r["mc"] = by_num[int(ln)][0]
        elif num_codes:
            r["mc"] = num_codes[0]
        else:
            r["mc"] = ""
            msgs.append("설비를 찾을 수 없음 (설비명·설비코드·호기 번호)")
            lv = "err"
        r["mname"] = machines.get(r["mc"], "")
        n, name, mold, b = info.get(r["item"], (0, "", "", 0))
        r["iname"], r["mold"] = name, mold or "*"
        if not n:
            msgs.append("품목 마스터에 없음 → MES 일괄처리에서 빠짐")
            lv = "err"
        if not mold:
            msgs.append("금형 연결 없음 (금형 * 로 들어감)")
            lv = "warn" if lv == "ok" else lv
        if n and not b:
            msgs.append("피치당 소재중량 없음 (소재 사용량 0)")
            lv = "warn" if lv == "ok" else lv
        d = dups.get((r.get("date"), r["item"]))
        if d:
            msgs.append(f"같은 날 이 품번 작업지시 있음 ({d})")
            lv = "dup" if lv == "ok" else lv
        k = (r.get("date"), r["mc"], r["item"])
        if k in seen:
            msgs.append("표 안에 같은 날·라인·품번 줄이 또 있음")
            lv = "warn" if lv == "ok" else lv
        seen.add(k)
        r["lv"], r["st"] = lv, " / ".join(msgs) or "정상"
    cnt = {}
    for r in rows:
        k = (r.get("date"), r.get("mc"))
        cnt[k] = cnt.get(k, 0) + 1
        r["prio_auto"] = str(cnt[k])
    return org


def wo_learn_rule(conn, tbl):
    cur = conn.cursor()
    cur.execute(f"SELECT WORK_ORDER_NO, WORK_ORDER_DATE FROM (SELECT WORK_ORDER_NO, WORK_ORDER_DATE FROM {tbl} "
                f"WHERE WORK_ORDER_DATE >= ADD_MONTHS(SYSDATE, -15) AND WORK_ORDER_NO LIKE 'WO%' "
                f"ORDER BY DBMS_RANDOM.VALUE) WHERE ROWNUM <= 4000")
    score = [[0, 0] for _ in WO_RULES]
    n_late = 0
    for no, d in cur.fetchall():
        m = WO_RE.match(str(no or "").strip())
        if not m or not d:
            continue
        late = d.month >= 10
        n_late += late
        for i, (_, f) in enumerate(WO_RULES):
            if f(d) == m.group(1):
                score[i][1] += 1
                score[i][0] += late
    best = max(range(len(WO_RULES)), key=lambda i: (score[i][0], score[i][1]))
    if score[best][1] == 0:
        best = 0
    note = (f"{WO_RULES[best][0]} (과거 번호 {score[best][1]}건 일치"
            + (f", 10~12월 {score[best][0]}/{n_late}건" if n_late else ", 10~12월 자료 없음") + ")")
    return WO_RULES[best][1], note


def wo_numbers(conn, tbl, dates):
    cur = conn.cursor()
    cur.execute(f"SELECT WORK_ORDER_NO, WORK_ORDER_DATE FROM {tbl} WHERE WORK_ORDER_DATE >= :d",
                d=dt.datetime.now() - dt.timedelta(days=90))
    seq, same = 0, {}
    for no, d in cur.fetchall():
        m = WO_RE.match(str(no or "").strip())
        if not m:
            continue
        seq = max(seq, int(m.group(2)))
        if d:
            same.setdefault(d.date(), m.group(1))
    rule, note, out = None, "", []
    for d in dates:
        pre = same.get(d.date())
        if pre is None:
            if rule is None:
                rule, note = wo_learn_rule(conn, tbl)
            pre = rule(d)
        seq += 1
        out.append(f"WO{pre}P{seq}")
    return out, note


def batch_out(rows):
    keep = ("line", "date_t", "shift", "item", "qty_t", "prio", "mc", "mname", "iname", "mold", "no", "lv", "st",
            "prio_auto", "qty", "date", "shift_v", "recent_date", "recent_wo", "spm")
    return [{k: jsonable(r.get(k)) for k in keep if k in r} for r in rows]


def api_batch_check(body, user):
    rows = body.get("rows") or []
    if not rows:
        raise UserError("등록할 줄이 없습니다.")
    org = run_db(batch_check, rows, body.get("d_def"), body.get("shift"))
    return {"rows": batch_out(rows), "org": org}


def order_rows(rows):
    pr = lambda r: int(r["prio"]) if str(r.get("prio") or "").isdigit() else int(r.get("prio_auto") or 999)  # noqa: E731
    return sorted(range(len(rows)), key=lambda i: (rows[i]["date"], rows[i]["mc"], pr(rows[i]), i))


def api_batch_make_no(body, user):
    rows = body.get("rows") or []
    if not rows:
        raise UserError("등록할 줄이 없습니다.")
    tbl = wo_tbl()

    def work(conn):
        batch_check(conn, rows, body.get("d_def"), body.get("shift"))
        if any(r["lv"] == "err" for r in rows):
            return None
        order = order_rows(rows)
        nos, note = wo_numbers(conn, tbl, [rows[i]["date"] for i in order])
        for i, no in zip(order, nos):
            rows[i]["no"] = no
        return note
    note = run_db(work)
    if note is None:
        return {"rows": batch_out(rows), "error": "빨간 줄이 있습니다. 고치거나 뺀 뒤 하세요."}
    return {"rows": batch_out(rows), "note": note}


def api_batch_apply(body, user):
    rows = body.get("rows") or []
    if not rows:
        raise UserError("등록할 줄이 없습니다.")
    nos = [str(r.get("no") or "").strip() for r in rows]
    if not all(nos) or not all(WO_RE.match(n) for n in nos):
        raise UserError("작업지시번호가 없는 줄이 있습니다. [W/O 번호생성]을 먼저 하세요.")
    if len(set(nos)) != len(nos):
        raise UserError("표 안에 같은 작업지시번호가 있습니다. [W/O 번호생성]을 다시 하세요.")
    try:
        hours = float(body.get("hours") or 0)
    except (TypeError, ValueError):
        hours = 0
    tbl = wo_tbl()
    up = qualify(T_UP)
    who = reg_user()
    now = dt.datetime.now()

    def work(conn):
        org = batch_check(conn, rows, body.get("d_def"), body.get("shift"))    # 반영 직전에 다시 검사
        bad = [r for r in rows if r["lv"] == "err"]
        if bad:
            raise UserError("검사에서 오류(빨강) 줄이 나왔습니다. 고친 뒤 다시 하세요:\n  "
                            + "\n  ".join(f"{r['line']} {r['item']}: {r['st']}" for r in bad[:10]))
        cur = conn.cursor()
        for no in nos:
            cur.execute(f"SELECT COUNT(*) FROM {tbl} WHERE WORK_ORDER_NO = :n", n=no)
            if cur.fetchone()[0]:
                raise UserError(f"작업지시번호 {no} 가 그 사이 생겼습니다. [W/O 번호생성]을 다시 하세요.")
        meta = table_meta(conn, up)
        cur.execute(f"SELECT NVL(MAX(SESSION_ID), 0) + 1 FROM {up}")
        sid = int(cur.fetchone()[0])
        try:
            for k, (r, no) in enumerate(zip(rows, nos), 1):
                insert_dict(cur, up, meta, {
                    "SESSION_ID": sid, "UPLOAD_SEQ": k, "ORGANIZATION_ID": org,
                    "MACHINE_CODE": r["mc"], "MACHINE_NAME": r.get("mname") or None,
                    "WORK_ORDER_DATE": r["date"], "WORK_ORDER_NO": no, "ITEM_CODE": r["item"],
                    "PLAN_QTY": r["qty"],
                    "PLAN_PRIORITY": int(r["prio"]) if str(r["prio"]).isdigit() else int(r.get("prio_auto") or k),
                    "WORK_SHIFT": r.get("shift_v") or "1", "STATUS_FLAG": "Y",
                    "COMMENTS": (f"web 최근30일 {hours:g}h 계획" if hours else "web 일괄등록"),
                    "ENTER_DATE": now, "ENTER_BY": who, "LAST_MODIFY_DATE": now, "LAST_MODIFY_BY": who})
            out, msg = cur.var(str), cur.var(str)
            cur.callproc(PROC, [org, sid, out, msg])        # MES 프로그램: 작업지시 생성 + COMMIT
            res = (out.getvalue(), msg.getvalue())
        except Exception:
            conn.rollback()
            raise
        if res[0] != "OK":
            conn.rollback()
            raise RuntimeError(f"MES 일괄처리 실패: {res[1]}")
        got = []
        for no in nos:
            cur.execute(f"SELECT MOLD_CODE, MACHINE_NAME FROM {tbl} WHERE WORK_ORDER_NO = :n", n=no)
            got.append(cur.fetchone())
        return sid, got
    try:
        sid, got = run_db(work)
    except UserError:
        raise
    except Exception as e:
        raise RuntimeError(ora_hint(e))
    ok = 0
    for r, no, gt in zip(rows, nos, got):
        r["no"] = no
        if gt:
            ok += 1
            r["lv"], r["st"] = "ok", f"반영됨 (금형 {show(gt[0])})"
        else:
            r["lv"], r["st"] = "err", "반영 안 됨 (품목 마스터 조직 확인)"
    append_log([[now_s(), user, "작업지시일괄(웹)", tbl, f"session {sid}",
                 "; ".join(f"{r['no']} {r['mc']} {r['item']} {r['qty']}" for r in rows)]])
    return {"rows": batch_out(rows), "ok": ok, "n": len(rows), "sid": sid}


def api_info(body, user):
    return {"version": VERSION, "pin": bool(str(CFG.get("web_pin") or "").strip()),
            "dsn": CFG["oracle"]["dsn"], "table": wo_tbl(), "today": f"{dt.date.today():%Y-%m-%d}"}


def api_ping(body, user):
    def work(conn):
        cur = conn.cursor()
        cur.execute("SELECT SYSDATE FROM DUAL")
        return cur.fetchone()[0]
    return {"ok": True, "db_time": show(run_db(work))}


ROUTES = {
    "/api/info": api_info, "/api/ping": api_ping,
    "/api/wo/search": api_wo_search, "/api/wo/save_grid": api_wo_save_grid,
    "/api/wo/delete_check": api_wo_delete_check, "/api/wo/delete": api_wo_delete,
    "/api/wo/form": api_form, "/api/wo/form_save": api_form_save,
    "/api/item/lookup": api_item_lookup, "/api/item/find": api_item_find,
    "/api/batch/lists": api_batch_lists, "/api/batch/spm": api_batch_spm,
    "/api/batch/recent": api_batch_recent, "/api/batch/check": api_batch_check,
    "/api/batch/make_no": api_batch_make_no, "/api/batch/apply": api_batch_apply,
}
NO_PIN = {"/api/info"}


# ==============================================================
# 웹 서버
# ==============================================================
class Handler(BaseHTTPRequestHandler):
    server_version = "TJD-MES-Web"

    def log_message(self, fmt, *args):
        if "/api/" in (self.path or ""):
            sys.stdout.write(f"[{dt.datetime.now():%H:%M:%S}] {self.client_address[0]} {fmt % args}\n")

    def send(self, code, data, ctype="application/json; charset=utf-8", extra=None):
        if not isinstance(data, (bytes, bytearray)):
            data = json.dumps(jsonable(data), ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        p = urlparse(self.path).path
        if p in ("/", "/index.html"):
            path = os.path.join(BASE, "index.html")
            if not os.path.exists(path):
                return self.send(404, "index.html 이 server.py 와 같은 폴더에 없습니다.".encode("utf-8"),
                                 "text/plain; charset=utf-8")
            with open(path, "rb") as f:
                return self.send(200, f.read(), "text/html; charset=utf-8")
        if p == "/api/info":
            return self.send(200, api_info({}, ""))
        self.send(404, {"error": "없는 주소입니다."})

    def do_POST(self):
        p = urlparse(self.path).path
        try:
            n = int(self.headers.get("Content-Length") or 0)
            body = json.loads(self.rfile.read(n).decode("utf-8") or "{}") if n else {}
        except Exception:
            return self.send(400, {"error": "요청 형식 오류"})
        pin = str(CFG.get("web_pin") or "").strip()
        if pin and p not in NO_PIN and self.headers.get("X-Pin", "") != pin:
            return self.send(401, {"error": "PIN이 맞지 않습니다.", "pin": True})
        name = unquote(self.headers.get("X-User", "") or "").strip()[:20] or "이름없음"
        user = f"{name}(웹 {self.client_address[0]})"
        try:
            if p == "/api/wo/export":
                data, fname = api_wo_export(body, user)
                return self.send(200, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 {"Content-Disposition": "attachment; filename*=UTF-8''" +
                                  quote(fname)})
            fn = ROUTES.get(p)
            if not fn:
                return self.send(404, {"error": "없는 기능입니다."})
            self.send(200, fn(body, user))
        except UserError as e:
            self.send(400, {"error": str(e), "kind": "warn"})
        except Exception as e:
            traceback.print_exc()
            self.send(500, {"error": ora_hint(e)})


def local_ips():
    ips = set()
    try:
        for info in socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET):
            ips.add(info[4][0])
    except Exception:
        pass
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("192.168.1.250", 1521))       # 실제로 보내지는 않음 (내 IP 알아내기)
        ips.add(s.getsockname()[0])
        s.close()
    except Exception:
        pass
    return sorted(i for i in ips if not i.startswith("127."))


def main():
    port = int(sys.argv[1]) if len(sys.argv) > 1 else int(CFG.get("web_port") or 8000)
    print("=" * 60)
    print(f" 태진다이텍 MES 웹 서버  {VERSION}")
    print(f" 설정: {CFG_PATH if os.path.exists(CFG_PATH) else '(설정 파일 없음 - 기본값)'}")
    print(f" DB  : {CFG['oracle']['dsn']}   표: {wo_tbl()}")
    if oracledb is None:
        print(f" ※ oracledb 없음: python -m pip install oracledb   ({_ORA_ERR})")
    else:
        try:
            print(f" DB 접속 확인: OK (DB 시각 {api_ping({}, '')['db_time']})")
        except Exception as e:
            print(" ※ DB 접속 실패 (서버는 켜 둡니다. 화면에서 다시 시도하세요)\n   " + str(e).replace("\n", "\n   "))
    if openpyxl is None:
        print(" ※ openpyxl 없음 (엑셀 저장·SPM 표 안 됨): python -m pip install openpyxl")
    print("-" * 60)
    print(" 다른 PC·폰 브라우저에서 접속:")
    for ip in local_ips() or ["이PC주소"]:
        print(f"   http://{ip}:{port}")
    print(f"   (이 PC에서는 http://localhost:{port})")
    print(" 끄려면 이 창을 닫거나 Ctrl+C")
    print("=" * 60)
    srv = ThreadingHTTPServer(("0.0.0.0", port), Handler)
    srv.daemon_threads = True
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
