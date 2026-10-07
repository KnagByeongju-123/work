# -*- coding: utf-8 -*-
"""
태진다이텍 MES 웹 서버 (사내망 전용, 인터넷 필요 없음)

- 기존 프로그램(태진다이텍_MES관리_v7_x.py)과 같은 폴더에 둡니다.
  그 파일을 불러와서 DB 규칙(검사·등록·수정·삭제·중복검사)을 그대로 씁니다. → PC 프로그램과 웹이 똑같이 동작
- 설정(mold_config.json), SPM(spm.xlsx), 작업기록(mold_register_log.csv), 삭제 백업도 같은 파일을 씁니다.
- 이 PC 1대에만 오라클(Instant Client + oracledb)을 설치하고 켜 둡니다. 다른 PC·폰은 브라우저로 접속.
필요: pip install oracledb openpyxl
"""
import os
import re
import io
import sys
import csv
import json
import glob
import uuid
import types
import base64
import socket
import threading
import time
import traceback
import importlib.util
import datetime as dt
from collections import OrderedDict
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse, unquote, quote

VERSION = "web 2.0"
HERE = os.path.dirname(os.path.abspath(__file__))
TL = threading.local()                      # 지금 요청한 사람 이름 (작업기록용)


class UserError(Exception):
    """화면에 그대로 보여줄 안내 (경고)"""


# ==============================================================
# 기존 프로그램 불러오기 (화면 부품 tkinter 는 빈 껍데기로 바꿔서 DB 로직만 씀)
# ==============================================================
class _Dummy:
    def __init__(self, *a, **k):
        pass

    def __call__(self, *a, **k):
        return _Dummy()

    def __getattr__(self, name):
        return _Dummy()

    def get(self, *a, **k):
        return ""


def _stub_tk():
    def mod(name):
        m = types.ModuleType(name)
        m.__getattr__ = lambda attr: _Dummy
        return m
    tk = mod("tkinter")
    tk.TclError = Exception
    subs = {}
    for s in ("ttk", "messagebox", "filedialog", "simpledialog", "font"):
        subs[s] = mod(f"tkinter.{s}")
        setattr(tk, s, subs[s])
    mb = subs["messagebox"]
    mb.askyesno = lambda *a, **k: False
    mb.showinfo = mb.showwarning = mb.showerror = lambda *a, **k: None
    sys.modules["tkinter"] = tk
    for s, m in subs.items():
        sys.modules[f"tkinter.{s}"] = m


def find_program():
    """같은 폴더(또는 한 칸 위)에서 기존 프로그램 파일 찾기 (가장 최근 것)"""
    cands = []
    for folder in (HERE, os.path.dirname(HERE)):
        for p in glob.glob(os.path.join(folder, "*.py")):
            if os.path.abspath(p) == os.path.abspath(__file__):
                continue
            try:
                with open(p, encoding="utf-8", errors="ignore") as f:
                    head = f.read()
            except Exception:
                continue
            if "class Mes" in head and "GRID_SPECS" in head and "class App" in head:
                cands.append(p)
    if not cands:
        raise SystemExit("기존 프로그램(태진다이텍_MES관리_v7_x.py)을 server.py 와 같은 폴더에 넣으세요.")
    return max(cands, key=os.path.getmtime)


_stub_tk()
PROGRAM = find_program()
_spec = importlib.util.spec_from_file_location("mesapp", PROGRAM)
M = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(M)

show, cell_text, conv, ident, blen, ora_hint, parse_day = M.show, M.cell_text, M.conv, M.ident, M.blen, M.ora_hint, M.parse_day
openpyxl = M.openpyxl
oracledb = M.oracledb
UPLOAD_DIR = os.path.join(M.BASE, "web_uploads")


class WebApp:
    """기존 프로그램의 App 대신 쓰는 껍데기 (cfg / mes / run_db / 작업기록)"""
    LOG_COLS = M.App.LOG_COLS

    def __init__(self):
        self.cfg = M.load_cfg()
        self.mes = M.Mes(self)
        self.xl = M.ExcelIO(self.mes)
        self.root = _Dummy()
        self.lock = threading.Lock()

    def connect(self):
        if not getattr(self.mes, "client_ready", False):
            with self.lock:                 # Instant Client 켜기는 한 번만
                return self.mes.connect()
        return self.mes.connect()

    def run_db(self, fn):
        conn = self.connect()
        try:
            return fn(conn)
        finally:
            try:
                conn.close()
            except Exception:
                pass

    def append_log(self, rows):
        u = getattr(TL, "user", "")
        if u:
            rows = [[r[0], u] + list(r[2:]) for r in rows]
        with self.lock:
            M.App.append_log(self, rows)

    def load_log(self):
        pass

    def busy(self, on):
        pass

    def err(self, e):
        raise e

    def save_cfg(self):
        M.save_cfg(self.cfg)


APP = WebApp()
CFG = APP.cfg
MES = APP.mes
LOCK = threading.Lock()


def reg_user(default=""):
    return (CFG.get("reg_user") or "").strip() or default


def now_s():
    return f"{dt.datetime.now():%Y-%m-%d %H:%M:%S}"


def parse_day_u(t):
    try:
        return parse_day(t)
    except ValueError as e:
        raise UserError(str(e))


def day_dt(text):
    return dt.datetime.combine(parse_day_u(text), dt.time())


def jsonable(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return show(v)
    if isinstance(v, dict):
        return {str(k): jsonable(x) for k, x in v.items()}
    if isinstance(v, (list, tuple, set)):
        return [jsonable(x) for x in v]
    if isinstance(v, (str, int, float, bool)) or v is None:
        return v
    return cell_text(v)


def rows_text(rows):
    return [[cell_text(v) for v in r] for r in rows]


def table_meta(conn, tbl):
    return M.table_meta(conn, tbl)


# ---------- 조회 결과 잠깐 보관 (수정·삭제할 때 원래 값 그대로 쓰기 위해) ----------
_cache = OrderedDict()


def cache_put(data):
    tok = uuid.uuid4().hex[:16]
    with LOCK:
        _cache[tok] = data
        while len(_cache) > 300:
            _cache.popitem(last=False)
    return tok


def cache_get(tok):
    d = _cache.get(tok or "")
    if d is None:
        raise UserError("조회한 지 오래되었거나 서버가 다시 켜졌습니다. [MES에서 조회]를 다시 하세요.")
    return d


# ---------- 파일 내려주기 / 올리기 ----------
class FileOut:
    def __init__(self, data, name, ctype="application/octet-stream"):
        self.data, self.name, self.ctype = data, name, ctype


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def need_xl():
    if openpyxl is None:
        raise RuntimeError("서버 PC에 openpyxl이 없습니다.  python -m pip install openpyxl")


def xlsx_of(sheets):
    """sheets: [(이름, 머리글[], 줄[][])] -> 엑셀 파일 bytes"""
    need_xl()
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    used = set()
    for k, (name, head, rows) in enumerate(sheets):
        ws = wb.active if k == 0 else wb.create_sheet()
        nm = re.sub(r"[\\/*?:\[\]]", "_", str(name))[:28] or f"Sheet{k + 1}"
        while nm in used:
            nm = nm[:25] + f"_{k}"
        used.add(nm)
        ws.title = nm
        ws.append(list(head))
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([v if isinstance(v, (int, float, dt.datetime)) else cell_text(v, 2000) for v in r])
        for j, h in enumerate(head, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(9, min(36, len(str(h)) * 2 + 4))
        ws.freeze_panes = "A2"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def save_upload(body, ext=".xlsx"):
    data = body.get("b64") or ""
    if not data:
        raise UserError("파일이 없습니다.")
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    tok = uuid.uuid4().hex[:16]
    path = os.path.join(UPLOAD_DIR, tok + ext)
    with open(path, "wb") as f:
        f.write(base64.b64decode(data.split(",")[-1]))
    return tok, path


def upload_path(tok):
    if not re.fullmatch(r"[0-9a-f]{16}", tok or ""):
        raise UserError("파일 정보가 올바르지 않습니다. 다시 불러오세요.")
    path = os.path.join(UPLOAD_DIR, tok + ".xlsx")
    if not os.path.exists(path):
        raise UserError("올린 파일이 서버에 없습니다. 다시 불러오세요.")
    return path


def make_defs(desc, names=None, comments=None):
    """조회 결과 컬럼 정보 -> {컬럼: {type, editable, req, maxlen, dbtype, name}}"""
    defs = {}
    for d in desc:
        name, t, isize, null_ok = d[0], d[1], d[3], d[6]
        tn = str(getattr(t, "name", t)).upper()
        if any(w in tn for w in ("LOB", "RAW", "LONG", "ROWID")):
            typ = None
        elif "CHAR" in tn:
            typ = "TEXT"
        elif any(w in tn for w in ("NUMBER", "BINARY_", "FLOAT", "INTEGER", "DECIMAL")):
            typ = "NUM"
        elif "DATE" in tn or "TIMESTAMP" in tn:
            typ = "DATE"
        else:
            typ = None
        nm = (names or {}).get(name) or (comments or {}).get(name) or ""
        defs[name] = {"type": typ or "TEXT", "editable": typ is not None and not name.endswith("__"),
                      "req": null_ok is False, "maxlen": int(isize or 0) if typ == "TEXT" else 0,
                      "dbtype": tn.replace("DB_TYPE_", ""), "name": nm}
    return defs


def check_value(d, text, col_label):
    """화면 글자 -> DB 값 (형식·길이·필수·음수 검사)"""
    try:
        new = conv(d, text)
        if new is None and d.get("req"):
            raise ValueError("필수 칸이라 비울 수 없음")
        if d["type"] == "TEXT" and d.get("maxlen") and new is not None:
            n = blen(new, int(CFG.get("kor_bytes") or 3))
            if n > int(d["maxlen"]):
                raise ValueError(f"너무 김 ({n}/{d['maxlen']}byte)")
        if d["type"] == "NUM" and isinstance(new, (int, float)) and new < 0:
            raise ValueError("음수는 넣을 수 없습니다")
    except ValueError as e:
        raise UserError(f"{col_label}: {e}")
    return new


# ==============================================================
# 공용 표: ③ 작업지시 / ② 단위중량 / 품목 (기존 GRID_SPECS 그대로)
# ==============================================================
_comments = {}


class FV:
    def __init__(self, v):
        self.v = v

    def get(self):
        return self.v

    def set(self, v):
        self.v = v


class GShim:
    """기존 GridTab 대신 (조건 함수들이 쓰는 값만)"""

    def __init__(self, key, f=None):
        if key not in M.GRID_SPECS:
            raise UserError("없는 표입니다.")
        self.key, self.spec, self.app = key, M.GRID_SPECS[key], APP
        self.fvars = {}
        for fk, _label, kind, default, _w in self.spec["filters"]:
            v = (f or {}).get(fk, default)
            self.fvars[fk] = FV(bool(v) if kind in ("check", "choice") else ("" if v is None else str(v)))
        self.cols, self.rows, self.dups = [], [], {}
        self.comments = _comments.setdefault(key, {})
        self.edits = {}

    def table(self):
        t = (CFG.get(self.spec["table_cfg"]) or "").strip()
        if not t:
            raise UserError(f"[환경설정]에 {self.spec['title']} 테이블명을 넣으세요.")
        return MES.qualify(t)

    def head(self, col):
        nm = self.spec["names"].get(col) or self.comments.get(col) or ""
        return f"{nm} ({col})" if nm and not col.endswith("__") else (nm or col)


def grid_query(conn, g):
    try:
        conds, binds = g.spec["where"](g)
    except ValueError as e:
        raise UserError(str(e))
    tbl = g.table()
    extra = g.spec.get("extra", lambda _g: [])(g)
    sel = ", ".join(["ROWIDTOCHAR(T.ROWID) AS RID__", "T.*"] + [f"{e} AS {a}" for e, a in extra])
    sql = f"SELECT {sel} FROM {tbl} T"
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += f" ORDER BY {g.spec['order']}"
    lim = int(CFG.get("limit") or 5000)
    sql = f"SELECT * FROM ({sql}) WHERE ROWNUM <= {lim}"
    cur = conn.cursor()
    cur.execute(sql, binds)
    rows, desc = cur.fetchall(), cur.description
    if not g.comments:
        owner, tname = tbl.split(".") if "." in tbl else (None, tbl)
        try:
            c2 = conn.cursor()
            c2.execute("SELECT COLUMN_NAME, COMMENTS FROM ALL_COL_COMMENTS WHERE OWNER = NVL(:o, USER) "
                       "AND TABLE_NAME = :t", o=owner, t=tname)
            g.comments.update({a: (b or "").strip() for a, b in c2.fetchall()})
        except Exception:
            pass
    names = [d[0] for d in desc]
    defs = make_defs(desc[1:], g.spec["names"], g.comments)
    cols, data = M.Mes.reorder(names[1:], [list(r[1:]) for r in rows], g.spec["first"])
    g.cols, g.rows = cols, data
    if g.spec.get("prepare"):
        g.spec["prepare"](g)
    return cols, data, [r[0] for r in rows], defs, lim


def grid_empty_msg(g):
    """단위중량 조회 0건: 금형 연결·품목 마스터·최근 생산 여부 (기존 bom_empty 와 같은 내용)"""
    if g.key != "bom":
        return ""
    q = g.fvars["q"].get().strip().upper()
    if not q:
        return ""
    link = MES.qualify(CFG["item"].get("table") or "ICOM_MOLD_ITEM")
    mst = MES.qualify(CFG.get("item_master_table") or "ICOM_ITEM_MASTER")
    dtl = MES.qualify("IMTL_RAW_SF_USE_CASE_DTL")

    def work(conn):
        cur = conn.cursor()
        out = {}
        for k, sql in (("mst", f"SELECT COUNT(*) FROM {mst} WHERE UPPER(ITEM_CODE) LIKE :q"),
                       ("link", f"SELECT COUNT(*) FROM {link} WHERE UPPER(ITEM_CODE) LIKE :q"),
                       ("prod", f"SELECT COUNT(*) FROM {dtl} WHERE UPPER(ITEM_CODE) LIKE :q AND ACTUCAL_DATE >= TRUNC(SYSDATE) - 60")):
            try:
                cur.execute(sql, q=f"%{q}%")
                out[k] = cur.fetchone()[0]
            except Exception:
                out[k] = None
        return out
    try:
        r = APP.run_db(work)
    except Exception:
        r = {}
    yn = lambda v, a, b: "?" if v is None else (a if v else b)  # noqa: E731
    return (f"'{q}' 가 들어간 제품의 단위중량(피치당 소재중량) 줄이 MES에 없습니다.\n\n"
            f"  · 금형-품번 연결: {yn(r.get('link'), '있음', '없음')}\n"
            f"  · 품목 마스터: {yn(r.get('mst'), '있음', '없음')}\n"
            f"  · 최근 60일 생산 기록: {yn(r.get('prod'), str(r.get('prod')) + '건 (소재 사용량이 0으로 쌓이는 중)', '없음')}\n\n"
            f"넣으려면: [수정 모드] 체크 → [새로 등록] → 제품 품번, 원소재(코일) 코드, 피치당 소재중량(g) 입력.\n"
            f"(비슷한 제품을 먼저 조회해서 그 줄을 선택하고 [새로 등록]하면 나머지 칸이 복사됩니다)")


def api_grid_search(body, user):
    g = GShim(body.get("key"), body.get("f"))
    cols, data, rids, defs, lim = APP.run_db(lambda c: grid_query(c, g))
    tags = [g.spec["tag"](g, r) or "" for r in data]
    return {"cols": cols, "heads": [g.head(c) for c in cols],
            "names": {c: (g.spec["names"].get(c) or g.comments.get(c, "")) for c in cols},
            "defs": {c: defs.get(c, {"type": "TEXT", "editable": False, "req": False, "maxlen": 0, "dbtype": "", "name": ""})
                     for c in cols},
            "rows": rows_text(data), "rids": rids, "tags": tags,
            "summary": g.spec["summary"](g).replace("※", "").strip(), "limit": lim, "hit_limit": len(data) >= lim,
            "time": f"{dt.datetime.now():%H:%M:%S}", "locked": g.spec["locked"], "sensitive": g.spec["sensitive"],
            "empty_msg": grid_empty_msg(g) if not data else ""}


def grid_row(conn, g, tbl, rid):
    """ROWID로 한 줄 다시 읽기 (조회 때와 같은 참고 컬럼 포함)"""
    extra = g.spec.get("extra", lambda _g: [])(g)
    sel = ", ".join(["T.*"] + [f"{e} AS {a}" for e, a in extra])
    cur = conn.cursor()
    cur.execute(f"SELECT {sel} FROM {tbl} T WHERE T.ROWID = CHARTOROWID(:r)", r=rid)
    r = cur.fetchone()
    names = [d[0] for d in cur.description]
    return (dict(zip(names, r)) if r else None), names, cur.description


def grid_label(g, row):
    return " / ".join(show(row.get(c)).strip() for c in g.spec["label_cols"] if c in row)


def api_grid_save(body, user):
    """표에서 고친 칸 저장. edits: [{rid, changes: {col: {old, new}}}]
    old(화면에 보였던 값)가 지금 MES 값과 다르면 전체 취소. 저장은 기존 Mes.update_by_rowid"""
    g = GShim(body.get("key"))
    edits = body.get("edits") or []
    if not edits:
        raise UserError("수정한 내용이 없습니다.")
    tbl = g.table()

    def work(conn):
        todo, alldefs = [], {}
        for e in edits:
            row, _names, desc = grid_row(conn, g, tbl, e["rid"])
            if row is None:
                raise UserError("이미 없어진 줄이 있습니다. 다시 조회하세요.")
            defs = make_defs(desc, g.spec["names"], g.comments)
            alldefs.update(defs)
            lab = grid_label(g, row)
            ch = {}
            for col, x in (e.get("changes") or {}).items():
                col = ident(col, "컬럼")
                if col in g.spec["locked"]:
                    raise UserError(f"{lab} · {g.head(col)}: {g.spec['locked'][col]}")
                d = defs.get(col)
                if not d or not d["editable"]:
                    raise UserError(f"{lab} · {g.head(col)}: 이 칸은 고칠 수 없습니다.")
                if cell_text(row.get(col)) != str(x.get("old") or ""):
                    raise UserError(f"{lab} · {g.head(col)}: 그 사이 다른 곳에서 바뀌었습니다 "
                                    f"(지금 값 '{cell_text(row.get(col))}'). 다시 조회하세요.")
                ch[col] = (row.get(col), check_value(d, x.get("new"), f"{lab} · {g.head(col)}"))
            if ch:
                todo.append((e["rid"], {"changes": ch, "label": lab}))
        if not todo:
            raise UserError("바뀐 내용이 없습니다.")
        n = MES.update_by_rowid(conn, tbl, todo, {c: d for c, d in alldefs.items() if not c.endswith("__")})
        return n, todo
    n, todo = APP.run_db(work)
    APP.append_log([[now_s(), "", f"{g.spec['title']}수정", tbl, e["label"],
                     "; ".join(f"{c}: {cell_text(o)} -> {cell_text(nv)}" for c, (o, nv) in e["changes"].items())]
                    for _, e in todo])
    return {"ok": True, "n": n}


def grid_new_prep(conn, g, tbl, rid):
    """새로 등록 준비: 복사할 줄, 컬럼 정보, 자동 칸 (기존 GridTab.new_row 와 같은 규칙)"""
    if rid:
        base, _names, desc = grid_row(conn, g, tbl, rid)
        if base is None:
            raise UserError("선택한 줄을 MES에서 찾지 못했습니다. 다시 조회하세요.")
    else:
        cur = conn.cursor()
        cur.execute(f"SELECT T.* FROM {tbl} T WHERE 1 = 0")
        desc, base = cur.description, {}
    defs = make_defs(desc, g.spec["names"], g.comments)
    g.cols = [d[0] for d in desc]   # 조건 함수(new_check 등)가 g.cols 를 씀
    meta = MES.col_meta(conn, tbl)
    vals = g.spec["new_defaults"](g, dict(base), meta) if g.spec.get("new_defaults") else dict(base)
    cols = [d[0] for d in desc if not d[0].endswith("__")]
    reg = reg_user()
    auto = {}
    for c in cols:
        d = defs.get(c, {})
        if d.get("type") == "DATE" and (M.RE_AUTO_DATE.search(c) or M.RE_UPD_DATE.search(c)):
            auto[c] = "SYSDATE"
        elif d.get("type") == "TEXT" and (M.RE_AUTO_USER.search(c) or M.RE_UPD_USER.search(c)) and reg:
            auto[c] = reg
    for c in auto:
        vals.pop(c, None)
    locked_new = set(g.spec.get("new_locked", {}))
    need = {c for c in cols if meta.get(c, {}).get("notnull") and not meta.get(c, {}).get("hasdef") and c not in auto}
    return cols, defs, vals, auto, locked_new, need, base


def api_grid_new_form(body, user):
    g = GShim(body.get("key"))
    if g.spec.get("hide_new"):
        raise UserError("이 표는 [새로 등록]을 쓰지 않습니다.")
    tbl = g.table()
    cols, defs, vals, auto, locked_new, need, base = APP.run_db(lambda c: grid_new_prep(c, g, tbl, body.get("rid")))
    out = []
    for c in cols:
        d = defs.get(c, {})
        if c in auto:
            kind, v = "auto", ("현재시각" if auto[c] == "SYSDATE" else auto[c]) + " (자동)"
        elif c in locked_new:
            kind, v = "auto", f"(자동: {g.spec['new_locked'][c]})"
        elif not d.get("editable"):
            kind, v = "off", cell_text(vals.get(c))
        else:
            kind, v = "edit", cell_text(vals.get(c))
        out.append({"col": c, "name": g.spec["names"].get(c) or g.comments.get(c, ""), "value": v, "kind": kind,
                    "need": c in need, "dbtype": d.get("dbtype", "")})
    return {"cols": out, "copied": bool(base), "hint": g.spec.get("new_hint", "")}


def api_grid_insert(body, user):
    g = GShim(body.get("key"))
    if g.spec.get("hide_new"):
        raise UserError("이 표는 [새로 등록]을 쓰지 않습니다.")
    tbl = g.table()
    texts = body.get("values") or {}

    def work(conn):
        cols, defs, vals, auto, locked_new, need, _base = grid_new_prep(conn, g, tbl, body.get("rid"))
        for c, t in texts.items():
            if c in cols and c not in auto and c not in locked_new and defs.get(c, {}).get("editable"):
                vals[c] = check_value(dict(defs[c], req=False), t, c)
        miss = [c for c in cols if c in need and c not in locked_new and vals.get(c) is None]
        if miss:
            raise UserError("반드시 입력할 칸이 비었습니다:\n  " + ", ".join(miss))
        chk = g.spec.get("new_check")
        if chk:
            msg = chk(g, conn, tbl, vals)
            if msg:
                raise UserError(msg)
        label = " / ".join(show(vals.get(c)) for c in g.spec["label_cols"])
        if not body.get("confirm"):
            return {"preview": True, "label": label}
        ins = {c: v for c, v in vals.items() if v is not None and c in cols and c not in locked_new}
        MES.insert_row(conn, tbl, ins, auto, g.spec.get("derive"))
        return {"ok": True, "label": label, "ins": ins}
    res = APP.run_db(work)
    if res.get("ok"):
        APP.append_log([[now_s(), "", f"{g.spec['title']}등록", tbl, res["label"],
                         "; ".join(f"{c}={cell_text(v)}" for c, v in res.pop("ins").items())]])
    return res


def job_grid_delete_check(body, user, progress, cancelled):
    g = GShim(body.get("key"))
    rids = body.get("rids") or []
    if not rids:
        raise UserError("삭제할 줄을 선택하세요.")
    tbl = g.table()

    def work(conn):
        lines, blocked, used, rows = [], [], [], []
        for rid in rids:
            row, names, _d = grid_row(conn, g, tbl, rid)
            if row is None:
                blocked.append("이미 없어진 줄이 있습니다. 다시 조회하세요.")
                continue
            g.cols = names
            lab = grid_label(g, row)
            lines.append(lab)
            why = g.spec.get("delete_block", lambda _g, _r: None)(g, [row[c] for c in names])
            if why:
                blocked.append(f"{lab}: {why}")
            rows.append(row)
        uc = g.spec.get("usage_col")
        if not blocked and uc:
            for v in sorted({show(r.get(uc)).strip() for r in rows if r.get(uc) is not None}):
                if cancelled():
                    raise UserError("취소했습니다.")
                u, _ = MES.usage_of_value(conn, uc, v, {tbl.split(".")[-1]}, progress, cancelled)
                used += [f"{v} → {t}: 기록 있음" for t, _c, _n in u]
            if used and g.spec.get("usage_block"):
                blocked.append("아래처럼 다른 곳에서 쓰이고 있어 삭제할 수 없습니다:\n" + "\n".join(used[:15]))
        return {"lines": lines, "blocked": blocked, "used": used}
    return APP.run_db(work)


def api_grid_delete(body, user):
    g = GShim(body.get("key"))
    rids = body.get("rids") or []
    tbl = g.table()
    if not rids:
        raise UserError("삭제할 줄을 선택하세요.")
    if body.get("used") and (body.get("typed") or "").strip() != "삭제":
        raise UserError("다른 테이블에 기록이 있어 '삭제'라고 입력해야 지울 수 있습니다.")

    def work(conn):
        targets = []
        for rid in rids:
            row, names, _d = grid_row(conn, g, tbl, rid)
            if row is None:
                raise UserError("이미 없어진 줄입니다. 다시 조회하세요.")
            g.cols = names
            why = g.spec.get("delete_block", lambda _g, _r: None)(g, [row[c] for c in names])
            if why:
                raise UserError(f"{grid_label(g, row)}: {why}")
            targets.append((rid, grid_label(g, row), row))
        with open(M.BACKUP_PATH, "a", encoding="utf-8") as f:
            for rid, label, row in targets:
                f.write(json.dumps({"time": now_s(), "user": user, "table": tbl,
                                    "row": {k: show(v) for k, v in row.items() if not k.endswith("__")}},
                                   ensure_ascii=False) + "\n")
        n = MES.delete_by_rowid(conn, tbl, targets)
        return n, targets
    n, targets = APP.run_db(work)
    APP.append_log([[now_s(), "", f"{g.spec['title']}삭제", tbl, label,
                     "; ".join(f"{k}={cell_text(v)}" for k, v in row.items() if v is not None and not k.endswith("__"))]
                    for _rid, label, row in targets])
    return {"ok": True, "n": n}


def api_grid_export(body, user):
    g = GShim(body.get("key"), body.get("f"))
    need_xl()
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    cols, data, _r, _d, _l = APP.run_db(lambda c: grid_query(c, g))
    keep = [i for i, c in enumerate(cols) if not c.endswith("__")]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = g.spec["title"]
    ws.append([cols[i] for i in keep])
    ws.append([g.spec["names"].get(cols[i]) or g.comments.get(cols[i], "") for i in keep])
    for c in ws[1]:
        c.font = Font(bold=True)
    for c in ws[2]:
        c.font = Font(size=9, color="808080")
    for r in data:
        ws.append([r[i] if isinstance(r[i], (int, float, dt.datetime)) else cell_text(r[i], 2000) for i in keep])
    for j, i in enumerate(keep, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(30, len(cols[i]) + 4))
    ws.freeze_panes = "B3"
    bio = io.BytesIO()
    wb.save(bio)
    return FileOut(bio.getvalue(), f"{g.spec['title']}_{dt.datetime.now():%Y%m%d_%H%M}.xlsx", XLSX)


def api_grid_import(body, user):
    """엑셀로 일괄 수정: 바뀐 칸만 골라서 돌려줌 (화면에서 주황으로 표시 → [수정 저장])"""
    g = GShim(body.get("key"), body.get("f"))
    mcs = g.spec.get("match_cols") or ([g.spec["match_col"]] if g.spec.get("match_col") else [])
    if not mcs:
        raise UserError("이 표는 엑셀 일괄 수정을 쓰지 않습니다.")
    need_xl()
    _tok, path = save_upload(body)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        xrows = list(wb.active.iter_rows(values_only=True))
    finally:
        try:
            os.remove(path)
        except Exception:
            pass
    cols, data, rids, defs, _l = APP.run_db(lambda c: grid_query(c, g))
    if any(m not in cols for m in mcs):
        raise UserError("먼저 [MES에서 조회]로 고칠 줄들을 불러오세요.")
    if not xrows:
        raise UserError("엑셀에 내용이 없습니다.")
    hdr = [str(h or "").strip().upper() for h in xrows[0]]
    miss = [m for m in mcs if m not in hdr]
    if miss:
        raise UserError(f"1행에 {', '.join(miss)} 컬럼이 있어야 합니다.")
    kis = [hdr.index(m) for m in mcs]
    mcis = [cols.index(m) for m in mcs]
    idx = {}
    for i, r in enumerate(data):
        idx.setdefault(tuple(show(r[c]).strip().upper() for c in mcis), []).append(i)
    targets = [(j, h) for j, h in enumerate(hdr) if h in cols and h not in mcs]
    start = 2 if len(xrows) > 1 and all((str(v or "").strip() in ("", g.spec["names"].get(hdr[j], ""), g.comments.get(hdr[j], "")))
                                        for j, v in enumerate(xrows[1]) if j < len(hdr)) else 1
    changes, notfound, errs, skipped = [], [], [], set()
    for r in xrows[start:]:
        if any(k >= len(r) or r[k] is None or str(r[k]).strip() == "" for k in kis):
            continue
        keyt = tuple(show(r[k]).strip().upper() for k in kis)
        hits = idx.get(keyt)
        if not hits:
            notfound.append("/".join(keyt))
            continue
        for i in hits:
            for j, col in targets:
                if j >= len(r):
                    continue
                d = defs.get(col)
                v = r[j]
                old = data[i][cols.index(col)]
                if col in g.spec["locked"] or not d or not d["editable"]:
                    if ("" if v is None else show(v)).strip() != cell_text(old).strip():
                        skipped.add(col)   # 실제로 바꾼 칸만 알려줌
                    continue
                newtxt = "" if v is None else (show(v) if not isinstance(v, str) else v)
                try:
                    newv = check_value(d, newtxt, col)
                except UserError as e:
                    errs.append(f"{'/'.join(keyt)} {e}")
                    continue
                if cell_text(newv) == cell_text(old):
                    continue
                changes.append({"rid": rids[i], "col": col, "old": cell_text(old), "new": cell_text(newv)})
    msg = f"엑셀에서 {len(changes)}칸이 바뀐 것으로 표시했습니다 (주황).\n확인 후 [수정 저장]을 누르세요."
    if notfound:
        msg += (f"\n\n조회 목록에 없는 {' + '.join(mcs)} {len(notfound)}개: {', '.join(notfound[:10])}"
                + (" ..." if len(notfound) > 10 else "") + "\n(조회 조건을 넓혀서 다시 조회 후 불러오세요)")
    if skipped:
        msg += f"\n\n고칠 수 없는 칸이라 건너뜀: {', '.join(sorted(skipped))}"
    if errs:
        msg += f"\n\n형식 오류 {len(errs)}건:\n" + "\n".join(errs[:10])
    return {"changes": changes, "msg": msg}


def api_grid_tool(body, user):
    """표별 점검 도구 (단위중량: 최근 60일 단위중량 0 생산 점검)"""
    if body.get("key") != "bom":
        raise UserError("없는 기능입니다.")
    dtl = MES.qualify("IMTL_RAW_SF_USE_CASE_DTL")

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"SELECT ITEM_CODE, RAW_ITEM_CODE, COUNT(*), SUM(SHOT_QTY), MAX(ACTUCAL_DATE) FROM {dtl} "
                    f"WHERE ACTUCAL_DATE >= TRUNC(SYSDATE) - 60 AND NVL(UNIT_PER_QTY, 0) = 0 "
                    f"GROUP BY ITEM_CODE, RAW_ITEM_CODE ORDER BY MAX(ACTUCAL_DATE) DESC")
        return cur.fetchall()
    rows = APP.run_db(work)
    if not rows:
        return {"msg": "최근 60일 동안 단위중량 0으로 생산된 기록이 없습니다."}
    return {"title": "최근 60일 - 단위중량 0으로 생산된 조합",
            "note": "이 조합들은 단위중량이 없어 소재 사용량이 0으로 쌓였습니다. [새로 등록]으로 단위중량을 넣으세요.\n"
                    "(넣은 뒤부터 계산됩니다. 이미 지난 기록은 바뀌지 않습니다)",
            "cols": ["제품 품번", "원소재 코드", "기록 수", "샷 합계", "마지막 생산일"], "rows": rows_text(rows)}


# ==============================================================
# ③ 작업지시: 고치기 창 / 일괄등록 / 최근 불러오기 (기존 WoForm, WoBatch 와 같은 규칙)
# ==============================================================
WO_NAMES = M.GRID_SPECS["wo"]["names"]
DONE_COLS = M.WoForm.DONE_COLS


def wo_tbl():
    return MES.qualify(CFG.get("wo_table") or "IPLN_WORK_ORDER_MASTER")


def link_tbl():
    return MES.qualify(CFG["item"].get("table") or "ICOM_MOLD_ITEM")


def mst_tbl():
    return MES.qualify(CFG.get("item_master_table") or "ICOM_ITEM_MASTER")


def row_by_rid(conn, tbl, rid):
    cur = conn.cursor()
    cur.execute(f"SELECT * FROM {tbl} WHERE ROWID = CHARTOROWID(:r)", r=rid)
    r = cur.fetchone()
    return dict(zip([d[0] for d in cur.description], r)) if r else None


def api_wo_form(body, user):
    rid = body.get("rid")
    tbl = wo_tbl()

    def work(conn):
        meta = table_meta(conn, tbl)
        cur = conn.cursor()
        cur.execute(f"SELECT MACHINE_CODE, MAX(MACHINE_NAME) FROM {tbl} WHERE WORK_ORDER_DATE >= SYSDATE - 365 "
                    f"AND MACHINE_CODE IS NOT NULL GROUP BY MACHINE_CODE ORDER BY MACHINE_CODE")
        machines = [[show(a), show(b)] for a, b in cur.fetchall()]
        shifts, st = [], {}
        if "WORK_SHIFT" in meta:
            cur.execute(f"SELECT WORK_SHIFT, COUNT(*) FROM {tbl} WHERE WORK_ORDER_DATE >= SYSDATE - 180 "
                        f"GROUP BY WORK_SHIFT ORDER BY 2 DESC")
            shifts = [show(a) for a, _ in cur.fetchall() if a is not None]
        for c in ("WORK_ORDER_STATUS", "PLAN_STATUS"):
            if c in meta:
                cur.execute(f"SELECT {c}, COUNT(*) FROM {tbl} WHERE WORK_ORDER_DATE >= SYSDATE - 180 "
                            f"GROUP BY {c} ORDER BY 2 DESC")
                st[c] = [show(a) for a, _ in cur.fetchall() if a is not None]
        return meta, machines, shifts, st, row_by_rid(conn, tbl, rid)
    meta, machines, shifts, st, row = APP.run_db(work)
    if row is None:
        raise UserError("선택한 작업지시를 MES에서 찾지 못했습니다. 다시 조회하세요.")
    done = {c: show(row.get(c)) for c in DONE_COLS if row.get(c) not in (None, 0)}
    return {"machines": machines, "shifts": shifts, "status": st, "has_shift": "WORK_SHIFT" in meta,
            "row": {k: cell_text(v) for k, v in row.items()}, "done": done}


def api_wo_item_lookup(body, user):
    it = str(body.get("item") or "").strip()
    if not it:
        return {"n": 0}
    tbl = wo_tbl()

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"SELECT MAX(ITEM_NAME), COUNT(*) FROM {mst_tbl()} WHERE ITEM_CODE = :i", i=it)
        name, n = cur.fetchone()
        cur.execute(f"SELECT MAX(MOLD_CODE) FROM {link_tbl()} WHERE ITEM_CODE = :i", i=it)
        mold = cur.fetchone()[0]
        cur.execute(f"SELECT MACHINE_CODE, PLAN_QTY FROM (SELECT MACHINE_CODE, PLAN_QTY FROM {tbl} "
                    f"WHERE ITEM_CODE = :i ORDER BY WORK_ORDER_DATE DESC, WORK_ORDER_NO DESC) WHERE ROWNUM = 1", i=it)
        return name, n, mold, cur.fetchone()
    name, n, mold, last = APP.run_db(work)
    return {"n": n or 0, "name": show(name), "mold": show(mold),
            "last_mc": show(last[0]) if last else "", "last_qty": show(last[1]) if last else ""}


def api_wo_item_find(body, user):
    q = str(body.get("q") or "").strip().upper()

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"SELECT * FROM (SELECT M.ITEM_CODE, MAX(M.ITEM_NAME), "
                    f"(SELECT MAX(X.MOLD_CODE) FROM {link_tbl()} X WHERE X.ITEM_CODE = M.ITEM_CODE) "
                    f"FROM {mst_tbl()} M WHERE UPPER(M.ITEM_CODE) LIKE :q OR UPPER(M.ITEM_NAME) LIKE :q "
                    f"GROUP BY M.ITEM_CODE ORDER BY M.ITEM_CODE) WHERE ROWNUM <= 300", q=f"%{q}%")
        return cur.fetchall()
    return {"rows": rows_text(APP.run_db(work))}


def api_wo_form_save(body, user):
    """고치기 창 저장 (기존 WoForm.save_edit). confirm=false 면 바뀔 내용·경고만, true 면 저장.
    stamp(창을 열 때 값)가 지금 MES 값과 다르면 저장 안 함."""
    rid, vin, stamp, confirm = body.get("rid"), body.get("values") or {}, body.get("stamp") or {}, body.get("confirm")
    tbl = wo_tbl()

    def work(conn):
        meta = table_meta(conn, tbl)
        o = row_by_rid(conn, tbl, rid)
        if o is None:
            raise UserError("작업지시를 찾지 못했습니다. 다시 조회하세요.")
        changed_out = [c for c, s in stamp.items() if c in o and cell_text(o[c]) != s]
        if changed_out:
            raise UserError("창을 연 뒤에 다른 곳(MES 화면·현장)에서 이 작업지시가 바뀌었습니다.\n"
                            f"바뀐 칸: {', '.join(changed_out[:8])}\n창을 닫고 다시 조회한 뒤 고치세요. (저장 안 함)")

        def cast(c, text):
            m = meta.get(c, {})
            t = str(text or "").strip()
            if t == "":
                return None
            if m.get("type") == "NUMBER":
                try:
                    f = float(t.replace(",", ""))
                except ValueError:
                    raise UserError(f"{WO_NAMES.get(c, c)}: 숫자를 넣으세요 ({t})")
                return int(f) if f.is_integer() else f
            if m.get("type") == "DATE":
                return day_dt(t)
            if m.get("len") and blen(t, int(CFG.get("kor_bytes") or 3)) > int(m["len"]):
                raise UserError(f"{WO_NAMES.get(c, c)}: 너무 깁니다 (최대 {m['len']}byte)")
            return t

        day = parse_day_u(vin.get("WORK_ORDER_DATE"))
        od = o.get("WORK_ORDER_DATE")
        v = {"WORK_ORDER_DATE": od if od and od.date() == day else dt.datetime.combine(day, dt.time()),
             "ITEM_CODE": cast("ITEM_CODE", vin.get("ITEM_CODE")),
             "MACHINE_CODE": cast("MACHINE_CODE", vin.get("MACHINE_CODE")),
             "PLAN_QTY": cast("PLAN_QTY", vin.get("PLAN_QTY"))}
        if "WORK_SHIFT" in meta and "WORK_SHIFT" in vin:
            v["WORK_SHIFT"] = cast("WORK_SHIFT", vin.get("WORK_SHIFT"))
        for c in ("WORK_ORDER_STATUS", "PLAN_STATUS"):
            if c in meta and c in vin:
                v[c] = cast(c, vin.get(c))
        for c, lab in (("ITEM_CODE", "품번"), ("MACHINE_CODE", "설비"), ("PLAN_QTY", "계획수량")):
            if v[c] is None:
                raise UserError(f"{lab}을(를) 넣으세요.")
        if v["PLAN_QTY"] <= 0:
            raise UserError("계획수량은 1 이상이어야 합니다.")
        ch = {c: (o.get(c), x) for c, x in v.items() if cell_text(o.get(c)) != cell_text(x)}
        if not ch:
            raise UserError("바뀐 내용이 없습니다.")
        if "WORK_ORDER_DATE" in ch and od:
            delta = v["WORK_ORDER_DATE"] - dt.datetime.combine(od.date(), dt.time())
            for c in ("START_DATE", "END_DATE"):
                if isinstance(o.get(c), dt.datetime):
                    ch[c] = (o[c], o[c] + delta)
        if "ITEM_CODE" in ch and "MODEL_NAME" in o and o.get("MODEL_NAME") == o.get("ITEM_CODE"):
            ch["MODEL_NAME"] = (o["MODEL_NAME"], v["ITEM_CODE"])
        warn = []
        done = {c: o.get(c) for c in DONE_COLS if o.get(c) not in (None, 0)}
        key3 = set(ch) & {"ITEM_CODE", "MACHINE_CODE", "WORK_ORDER_DATE"}
        if done and key3:
            warn.append("이 작업지시는 실적이 있습니다 (" + ", ".join(f"{c}={show(x)}" for c, x in done.items()) +
                        "). 품번·설비·날짜를 바꾸면 실적 집계가 달라집니다.")
        cur = conn.cursor()
        if key3:
            cur.execute(f"SELECT COUNT(*) FROM {mst_tbl()} WHERE ITEM_CODE = :i", i=v["ITEM_CODE"])
            if not cur.fetchone()[0]:
                raise UserError(f"품목 마스터에 없는 품번입니다: {v['ITEM_CODE']}")
            cur.execute(f"SELECT MAX(MOLD_CODE) FROM {link_tbl()} WHERE ITEM_CODE = :i", i=v["ITEM_CODE"])
            if not cur.fetchone()[0]:
                warn.append(f"품번 {v['ITEM_CODE']}에 연결된 금형이 없어 금형이 '*'로 들어갑니다.")
            d0 = dt.datetime.combine(v["WORK_ORDER_DATE"].date(), dt.time())
            cur.execute(f"SELECT WORK_ORDER_NO FROM {tbl} WHERE WORK_ORDER_DATE >= :a AND WORK_ORDER_DATE < :b "
                        f"AND MACHINE_CODE = :m AND ITEM_CODE = :i AND ROWID <> CHARTOROWID(:r)",
                        a=d0, b=d0 + dt.timedelta(days=1), m=v["MACHINE_CODE"], i=v["ITEM_CODE"], r=rid)
            same = [show(x[0]) for x in cur.fetchall()]
            if same:
                warn.append(f"같은 날·설비·품번 작업지시가 이미 있습니다: {', '.join(same[:5])}")
        if "ITEM_CODE" in ch:
            der = M.wo_derive(conn, {**o, **v})
            for c in ("ITEM_NAME", "ITEM_SPEC"):
                if c in der and c in o and cell_text(der[c]) != cell_text(o.get(c)):
                    ch[c] = (o.get(c), der[c])
        lines = [{"col": c, "name": WO_NAMES.get(c, c), "old": cell_text(a), "new": cell_text(b)} for c, (a, b) in ch.items()]
        if not confirm:
            return {"preview": True, "no": show(o.get("WORK_ORDER_NO")), "lines": lines, "warn": warn}
        reg = reg_user("PYWO")
        auto = {}
        for c, m in meta.items():
            if m["type"] == "DATE" and M.RE_UPD_DATE.search(c):
                auto[c] = "SYSDATE"
            elif m["type"].startswith("VARCHAR") and M.RE_UPD_USER.search(c):
                auto[c] = reg
        sets, where, b = [], ["ROWID = CHARTOROWID(:rid)"], {"rid": rid}
        for i, (c, (a, x)) in enumerate(ch.items()):
            sets.append(f"{ident(c, '컬럼')} = :n{i}")
            b[f"n{i}"] = x
            if a is None:
                where.append(f"{ident(c, '컬럼')} IS NULL")
            else:
                where.append(f"{ident(c, '컬럼')} = :o{i}")
                b[f"o{i}"] = a
        for j, (c, how) in enumerate(auto.items()):
            if c in ch:
                continue
            if how == "SYSDATE":
                sets.append(f"{ident(c, '컬럼')} = SYSDATE")
            else:
                sets.append(f"{ident(c, '컬럼')} = :u{j}")
                b[f"u{j}"] = how
        try:
            cur.execute(f"UPDATE {tbl} SET {', '.join(sets)} WHERE {' AND '.join(where)}", b)
            if cur.rowcount != 1:
                raise UserError("그 사이 다른 곳(MES 화면·현장)에서 바뀌었거나 없어진 작업지시입니다.\n"
                                "다시 조회한 뒤 고치세요. (저장 안 함)")
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        new = row_by_rid(conn, tbl, rid) or {}
        return {"ok": True, "no": show(o.get("WORK_ORDER_NO")), "lines": lines,
                "mold": show(new.get("MOLD_CODE")), "mname": show(new.get("MACHINE_NAME"))}
    try:
        res = APP.run_db(work)
    except UserError:
        raise
    except Exception as e:
        raise RuntimeError(ora_hint(e))
    if res.get("ok"):
        APP.append_log([[now_s(), "", "작업지시수정", tbl, res["no"],
                         "; ".join(f"{x['col']}: {x['old']} -> {x['new']}" for x in res["lines"])]])
    return res


# ---------- 일괄등록 / 최근 불러오기 ----------
def batch_obj():
    """기존 WoBatch 의 계산 함수(목록 읽기·SPM 찾기·최근 작업지시)만 쓰기 위한 껍데기"""
    b = object.__new__(M.WoBatch)
    b.app, b.g = APP, GShim("wo")
    b.machines, b.item_list, b.item_last, b.mc_items, b.item_names, b.item_mcs = {}, [], {}, {}, {}, {}
    return b


def api_batch_lists(body, user):
    b = batch_obj()
    APP.run_db(b.load_lists)
    return {"machines": b.machines, "items": [[k, b.item_names.get(k, ""), b.item_last.get(k, "")] for k, _ in b.item_list],
            "mc_items": {k: [[it, n, show(last)] for it, n, last in v] for k, v in b.mc_items.items()},
            "item_mcs": {k: [[mc, n, show(last)] for mc, n, last in v] for k, v in b.item_mcs.items()}}


def api_batch_spm(body, user):
    mc, it = str(body.get("mc") or "").strip(), str(body.get("item") or "").strip().upper()
    if not it:
        raise UserError("품번을 넣으세요.")
    b = batch_obj()
    spm, cav, src = APP.run_db(lambda c: b.spm_of(c, mc, it))
    return {"spm": spm, "cav": cav, "src": src}


def api_batch_recent(body, user):
    try:
        hours = float(body.get("hours"))
        if not 0 < hours <= 24:
            raise ValueError
    except (TypeError, ValueError):
        raise UserError("근무시간은 0보다 크고 24 이하인 숫자로 넣으세요.")
    b = batch_obj()

    def work(conn):
        b.load_lists(conn)
        return b.recent_rows(conn, hours)
    rows = APP.run_db(work)
    return {"rows": [{"line": r["line"], "date_t": "", "shift": "", "item": r["item"], "qty_t": r["qty_t"], "prio": "",
                      "lv": r["lv"], "st": r["st"], "recent_date": show(r.get("recent_date"))[:10],
                      "recent_wo": r.get("recent_wo", ""), "spm": round(r["spm"], 2) if r.get("spm") else None}
                     for r in rows]}


def batch_check(conn, b, rows, d_def_text, shift_def):
    """기존 WoBatch.check 와 같은 검사 (화면 대신 rows 에 결과를 채움). return org"""
    tbl = wo_tbl()
    bom = MES.qualify(CFG.get("bom_table") or "ICOM_ITEM_CHILD")
    d_def = day_dt(d_def_text)
    for r in rows:
        r["lv"], r["st"] = "", ""
        for k in ("line", "date_t", "shift", "item", "qty_t", "prio"):
            r[k] = str(r.get(k) or "").strip()
        try:
            r["date"] = M.WoBatch.parse_date(r["date_t"]) if r["date_t"] else d_def
            r["shift_v"] = str(r["shift"] or shift_def or "1").strip()
        except Exception:
            r["date"], r["lv"], r["st"] = None, "err", "계획일자 형식 오류"
        try:
            q = float(r["qty_t"].replace(",", ""))
            if q <= 0:
                raise ValueError
            r["qty"] = int(q) if q.is_integer() else q
        except Exception:
            r["qty"], r["lv"] = None, "err"
            r["st"] = (r["st"] + " / " if r["st"] else "") + "계획수량 오류"
        r["item"] = r["item"].upper()
    cur = conn.cursor()
    cur.execute(f"SELECT ORGANIZATION_ID FROM (SELECT ORGANIZATION_ID, COUNT(*) N FROM {tbl} "
                f"WHERE WORK_ORDER_DATE >= SYSDATE - 90 GROUP BY ORGANIZATION_ID ORDER BY N DESC) WHERE ROWNUM = 1")
    org = (cur.fetchone() or [1])[0]
    cur.execute(f"SELECT MACHINE_CODE, MAX(MACHINE_NAME) FROM {tbl} WHERE WORK_ORDER_DATE >= SYSDATE - 400 "
                f"AND MACHINE_CODE IS NOT NULL GROUP BY MACHINE_CODE")
    machines = {**b.machines, **{show(a).strip(): show(c).strip() for a, c in cur.fetchall()}}
    info = {}
    for it in {r["item"] for r in rows if r["item"]}:
        cur.execute(f"SELECT MAX(ITEM_NAME), COUNT(*) FROM {mst_tbl()} WHERE ITEM_CODE = :i AND ORGANIZATION_ID = :o",
                    i=it, o=org)
        name, n = cur.fetchone()
        cur.execute(f"SELECT MAX(MOLD_CODE) FROM {link_tbl()} WHERE ITEM_CODE = :i", i=it)
        mold = cur.fetchone()[0]
        cur.execute(f"SELECT COUNT(*) FROM {bom} WHERE ITEM_CODE = :i AND NVL(UNIT_PER_QTY, 0) > 0", i=it)
        info[it] = (n, show(name), show(mold), cur.fetchone()[0])
    dups = {}
    for r in rows:
        if r.get("date") and r["item"]:
            cur.execute(f"SELECT MAX(WORK_ORDER_NO) FROM {tbl} WHERE WORK_ORDER_DATE >= :a AND WORK_ORDER_DATE < :b "
                        f"AND ITEM_CODE = :i", a=r["date"], b=r["date"] + dt.timedelta(days=1), i=r["item"])
            dups[(r["date"], r["item"])] = show(cur.fetchone()[0])
    norm = lambda x: re.sub(r"\s+", "", str(x or "")).upper()  # noqa: E731
    by_name, by_num = {}, {}
    for code, name in machines.items():
        by_name.setdefault(norm(name), []).append(code)
        m = re.search(r"(\d+)\s*호기", name)
        if m:
            by_num.setdefault(int(m.group(1)), []).append(code)
    seen = set()
    for r in rows:
        if r["lv"] == "err":
            continue
        msgs, lv = [], "ok"
        ln = r["line"]
        nums = [c for c in machines if c.isdigit() and ln.isdigit() and int(c) == int(ln)]
        if ln in machines:
            r["mc"] = ln
        elif len(by_name.get(norm(ln), [])) == 1:
            r["mc"] = by_name[norm(ln)][0]
        elif ln.isdigit() and len(by_num.get(int(ln), [])) == 1:
            r["mc"] = by_num[int(ln)][0]
        elif nums:
            r["mc"] = nums[0]
        else:
            r["mc"] = ""
            msgs.append("설비를 찾을 수 없음 (설비명·설비코드·호기 번호)")
            lv = "err"
        r["mname"] = machines.get(r["mc"], "")
        n, name, mold, bb = info.get(r["item"], (0, "", "", 0))
        r["iname"], r["mold"] = name, mold or "*"
        if not n:
            msgs.append("품목 마스터에 없음 → MES 일괄처리에서 빠짐")
            lv = "err"
        if not mold:
            msgs.append("금형 연결 없음 (금형 * 로 들어감)")
            lv = "warn" if lv == "ok" else lv
        if n and not bb:
            msgs.append("피치당 소재중량 없음 (소재 사용량 0)")
            lv = "warn" if lv == "ok" else lv
        d = dups.get((r.get("date"), r["item"]))
        if d:
            msgs.append(f"같은 날 이 품번 작업지시 있음 ({d})")
            lv = "dup" if lv == "ok" else lv
        k = (r.get("date"), r["mc"], r["item"])
        if k in seen:
            msgs.append("표 안에 같은 날·라인·품번 줄이 또 있음")
            lv = "warn" if lv == "ok" else lv
        seen.add(k)
        r["lv"], r["st"] = lv, " / ".join(msgs) or "정상"
    cnt = {}
    for r in rows:
        k = (r.get("date"), r.get("mc"))
        cnt[k] = cnt.get(k, 0) + 1
        r["prio_auto"] = str(cnt[k])
    return org


def batch_out(rows):
    keep = ("line", "date_t", "shift", "item", "qty_t", "prio", "mc", "mname", "iname", "mold", "no", "lv", "st",
            "prio_auto", "qty", "date", "shift_v", "recent_date", "recent_wo", "spm")
    return [{k: jsonable(r.get(k)) for k in keep if k in r} for r in rows]


def api_batch_check(body, user):
    rows = body.get("rows") or []
    if not rows:
        raise UserError("등록할 줄이 없습니다.")
    b = batch_obj()
    org = APP.run_db(lambda c: batch_check(c, b, rows, body.get("d_def"), body.get("shift")))
    return {"rows": batch_out(rows), "org": org}


def order_rows(rows):
    pr = lambda r: int(r["prio"]) if str(r.get("prio") or "").isdigit() else int(r.get("prio_auto") or 999)  # noqa: E731
    return sorted(range(len(rows)), key=lambda i: (rows[i]["date"], rows[i]["mc"], pr(rows[i]), i))


def api_batch_make_no(body, user):
    rows = body.get("rows") or []
    if not rows:
        raise UserError("등록할 줄이 없습니다.")
    tbl = wo_tbl()
    b = batch_obj()

    def work(conn):
        batch_check(conn, b, rows, body.get("d_def"), body.get("shift"))
        if any(r["lv"] == "err" for r in rows):
            return None
        order = order_rows(rows)
        nos, note = M.wo_numbers(conn, tbl, [rows[i]["date"] for i in order])
        for i, no in zip(order, nos):
            rows[i]["no"] = no
        return note
    note = APP.run_db(work)
    if note is None:
        return {"rows": batch_out(rows), "error": "빨간 줄이 있습니다. 고치거나 뺀 뒤 하세요."}
    return {"rows": batch_out(rows), "note": note}


def api_batch_apply(body, user):
    rows = body.get("rows") or []
    if not rows:
        raise UserError("등록할 줄이 없습니다.")
    nos = [str(r.get("no") or "").strip() for r in rows]
    if not all(nos) or not all(M.WO_RE.match(n) for n in nos):
        raise UserError("작업지시번호가 없는 줄이 있습니다. [W/O 번호생성]을 먼저 하세요.")
    if len(set(nos)) != len(nos):
        raise UserError("표 안에 같은 작업지시번호가 있습니다. [W/O 번호생성]을 다시 하세요.")
    try:
        hours = float(body.get("hours") or 0)
    except (TypeError, ValueError):
        hours = 0
    tbl = wo_tbl()
    up = MES.qualify(M.WoBatch.T_UP)
    who = reg_user("PYBATCH")
    now = dt.datetime.now()
    b = batch_obj()

    def work(conn):
        org = batch_check(conn, b, rows, body.get("d_def"), body.get("shift"))    # 반영 직전에 다시 검사
        bad = [r for r in rows if r["lv"] == "err"]
        if bad:
            raise UserError("검사에서 오류(빨강) 줄이 나왔습니다. 고친 뒤 다시 하세요:\n  "
                            + "\n  ".join(f"{r['line']} {r['item']}: {r['st']}" for r in bad[:10]))
        cur = conn.cursor()
        for no in nos:
            cur.execute(f"SELECT COUNT(*) FROM {tbl} WHERE WORK_ORDER_NO = :n", n=no)
            if cur.fetchone()[0]:
                raise UserError(f"작업지시번호 {no} 가 그 사이 생겼습니다. [W/O 번호생성]을 다시 하세요.")
        meta = table_meta(conn, up)
        cur.execute(f"SELECT NVL(MAX(SESSION_ID), 0) + 1 FROM {up}")
        sid = int(cur.fetchone()[0])
        try:
            for k, (r, no) in enumerate(zip(rows, nos), 1):
                M.insert_dict(cur, up, meta, {
                    "SESSION_ID": sid, "UPLOAD_SEQ": k, "ORGANIZATION_ID": org,
                    "MACHINE_CODE": r["mc"], "MACHINE_NAME": r.get("mname") or None,
                    "WORK_ORDER_DATE": r["date"], "WORK_ORDER_NO": no, "ITEM_CODE": r["item"],
                    "PLAN_QTY": r["qty"],
                    "PLAN_PRIORITY": int(r["prio"]) if str(r["prio"]).isdigit() else int(r.get("prio_auto") or k),
                    "WORK_SHIFT": r.get("shift_v") or "1", "STATUS_FLAG": "Y",
                    "COMMENTS": (f"python 최근30일 {hours:g}h 계획" if hours else "python 일괄등록"),
                    "ENTER_DATE": now, "ENTER_BY": who, "LAST_MODIFY_DATE": now, "LAST_MODIFY_BY": who})
            out, msg = cur.var(str), cur.var(str)
            cur.callproc(M.WoBatch.PROC, [org, sid, out, msg])      # MES 프로그램: 작업지시 생성 + COMMIT
            res = (out.getvalue(), msg.getvalue())
        except Exception:
            conn.rollback()
            raise
        if res[0] != "OK":
            conn.rollback()
            raise RuntimeError(f"MES 일괄처리 실패: {res[1]}")
        got = []
        for no in nos:
            cur.execute(f"SELECT MOLD_CODE, MACHINE_NAME FROM {tbl} WHERE WORK_ORDER_NO = :n", n=no)
            got.append(cur.fetchone())
        return sid, got
    try:
        sid, got = APP.run_db(work)
    except UserError:
        raise
    except Exception as e:
        raise RuntimeError(ora_hint(e))
    ok = 0
    for r, no, gt in zip(rows, nos, got):
        r["no"] = no
        if gt:
            ok += 1
            r["lv"], r["st"] = "ok", f"반영됨 (금형 {show(gt[0])})"
        else:
            r["lv"], r["st"] = "err", "반영 안 됨 (품목 마스터 조직 확인)"
    APP.append_log([[now_s(), "", "작업지시일괄", tbl, f"session {sid}",
                     "; ".join(f"{r['no']} {r['mc']} {r['item']} {r['qty']}" for r in rows)]])
    return {"rows": batch_out(rows), "ok": ok, "n": len(rows), "sid": sid}


# ==============================================================
# ① 금형·품번 관리 (기존 App.build_old 쪽)
# ==============================================================
def defs_out(defs, cols):
    return {c: {"type": defs.get(c, {}).get("type", "TEXT"), "editable": bool(defs.get(c, {}).get("editable", True)),
                "req": bool(defs.get(c, {}).get("req")), "maxlen": defs.get(c, {}).get("maxlen", 0),
                "key": bool(defs.get(c, {}).get("key")), "name": defs.get(c, {}).get("name", c),
                "dbtype": defs.get(c, {}).get("dbtype", "")} for c in cols}


def key_cols(p, cols):
    k = M.App.row_key(APP, p, cols, [None] * len(cols))
    return [c for c, _ in k] if k else []


def part_out(p, cols, rows, defs):
    tok = cache_put({"p": p, "cols": cols, "rows": rows, "defs": defs})
    return {"token": tok, "cols": cols, "heads": [M.App.head_name(defs, c) for c in cols], "rows": rows_text(rows),
            "defs": defs_out(defs, cols), "keys": key_cols(p, cols)}


def api_old_fetch(body, user):
    cols, rows, defs = APP.run_db(MES.fetch_old)
    out = part_out("mold", cols, rows, defs)
    out.update(code_col=MES.code_col(), link_col=MES.link_col(), ref=(CFG.get("ref_mold") or "").upper(),
               time=f"{dt.datetime.now():%Y-%m-%d %H:%M}", item_on=MES.enabled("item"),
               show_reg=bool(CFG.get("show_reg_tabs")))
    return out


def api_old_items(body, user):
    d = cache_get(body.get("token"))
    i = int(body.get("idx"))
    cc = MES.code_col()
    if cc not in d["cols"]:
        raise UserError(f"금형 표에 금형코드 컬럼({cc})이 없습니다.")
    code = d["rows"][i][d["cols"].index(cc)]
    cols, rows, defs = APP.run_db(lambda c: MES.fetch_items_of(c, code))
    out = part_out("item", cols, rows, defs)
    out.update(code=show(code).strip(), hint="")
    if not rows:
        try:
            cols2, rows2, _ = APP.run_db(lambda c: MES.find_items(c, show(code).strip()))
            if rows2:
                li = cols2.index(MES.link_col())
                found = sorted({repr(r[li]) for r in rows2})[:3]
                out["hint"] = f"※ 비슷한 금형코드로 {len(rows2)}건 있음: {', '.join(found)} (공백/대소문자 차이?)"
        except Exception:
            pass
    return out


def api_old_find_items(body, user):
    q = str(body.get("q") or "").strip()
    if not q:
        raise UserError("찾을 금형코드나 품번 일부를 입력하세요.")
    cols, rows, defs = APP.run_db(lambda c: MES.find_items(c, q))
    out = part_out("item", cols, rows, defs)
    out.update(search=q)
    return out


def old_edit_block(p, col, cols, row, defs):
    c = defs.get(col)
    if c is None:
        return "컬럼 정보를 찾을 수 없습니다."
    if M.App.row_key(APP, p, cols, row) is None:
        return f"{M.PART_NAME[p]} 매핑에 중복검사키가 없어서 어느 줄인지 특정할 수 없습니다."
    if (p == "mold" and col == MES.code_col()) or (p == "item" and col == MES.link_col()):
        return ("금형코드는 작업지시·샷수 이력 등 여러 곳에 연결되어 있어 여기서 수정할 수 없습니다.\n"
                "새 코드가 필요하면 [다른 이름으로 저장 모드]를 쓰세요.")
    if not c.get("editable", True):
        return f"{col} ({c.get('dbtype', '')}) 형식은 여기서 수정할 수 없습니다."
    return None


def api_old_save(body, user):
    """edits: [{token, idx, changes: {col: 새 글자}}] -> 기존 Mes.apply_updates 로 저장"""
    todo = []
    for e in body.get("edits") or []:
        d = cache_get(e.get("token"))
        p, cols, defs = d["p"], d["cols"], d["defs"]
        row = d["rows"][int(e["idx"])]
        key = M.App.row_key(APP, p, cols, row)
        if key is None:
            raise UserError(f"{M.PART_NAME[p]} 매핑에 중복검사키가 없어서 어느 줄인지 특정할 수 없습니다.")
        label = f"{M.PART_NAME[p]} " + "/".join(show(v) for _, v in key)
        ch = {}
        for col, text in (e.get("changes") or {}).items():
            why = old_edit_block(p, col, cols, row, defs)
            if why:
                raise UserError(f"{label} · {col}: {why}")
            new = check_value(defs[col], text, f"{label} · {defs[col].get('name', col)}")
            old = row[cols.index(col)]
            if cell_text(new) == cell_text(old) and (new is None) == (old is None):
                continue
            ch[col] = (old, new)
        if ch:
            todo.append({"p": p, "key": dict(key), "changes": ch, "label": label})
    if not todo:
        raise UserError("수정한 내용이 없습니다.")
    try:
        done = APP.run_db(lambda c: MES.apply_updates(c, todo))
    except ValueError as e:
        raise UserError(str(e))
    rows = []
    for e in done:
        by = {c["col"]: c for c in CFG[e["p"]]["columns"]}
        txt = "; ".join(f"{by[col]['name'] if col in by else col}: {cell_text(o)} -> {cell_text(n)}"
                        for col, (o, n) in e["changes"].items())
        rows.append([now_s(), "", f"{M.PART_NAME[e['p']]}수정", MES.table(e["p"]),
                     "/".join(show(v) for v in e["key"].values()), txt])
    APP.append_log(rows)
    return {"ok": True, "n": len(done)}


def api_old_delete_items(body, user):
    d = cache_get(body.get("token"))
    if d["p"] != "item":
        raise UserError("품번 표가 아닙니다.")
    cols = d["cols"]
    lc = MES.link_col()
    if lc not in cols:
        raise UserError(f"품번 표에 연결 컬럼({lc})이 없습니다.")
    li = cols.index(lc)
    picked = [d["rows"][int(i)] for i in body.get("idxs") or []]
    if not picked:
        raise UserError("아래 품번 표에서 지울 줄을 체크(☑)하거나 선택하세요.")

    def find(conn):
        out, used, cache = [], set(), {}
        for row in picked:
            v = row[li]
            if v not in cache:
                cache[v] = M.App.rid_rows(APP, conn, "item", lc, v)
            want = [cell_text(x) for x in row]
            hit = next((r for r in cache[v] if r["RID__"] not in used
                        and [cell_text(r.get(c)) for c in cols] == want), None)
            if not hit:
                raise UserError("품번 줄을 MES에서 찾지 못했습니다 (그 사이 바뀜?). 다시 조회하세요.")
            used.add(hit["RID__"])
            out.append(hit)
        return out
    hits = APP.run_db(find)
    label = lambda r: f"금형 {show(r.get(lc)).strip()} - 품번 {show(r.get('ITEM_CODE')).strip()}"  # noqa: E731
    if not body.get("confirm"):
        return {"preview": True, "lines": [label(r) for r in hits]}
    APP.run_db(lambda c: M.App.delete_plan(APP, c, [("item", r["RID__"], label(r)) for r in hits]))
    M.App.log_deleted(APP, [("item", r) for r in hits], dt.datetime.now(), "품번연결삭제")
    return {"ok": True, "n": len(hits)}


_plans = {}


def job_old_delete_mold_check(body, user, progress, cancelled):
    d = cache_get(body.get("token"))
    cc = MES.code_col()
    if cc not in d["cols"]:
        raise UserError(f"금형 표에 금형코드 컬럼({cc})이 없습니다.")
    ci = d["cols"].index(cc)
    codes = []
    for i in body.get("idxs") or []:
        c = d["rows"][int(i)][ci]
        if c not in codes:
            codes.append(c)
    if not codes:
        raise UserError("위 금형 표에서 지울 금형을 체크(☑)하거나 선택하세요.")
    ref = (CFG.get("ref_mold") or "").strip().upper()
    if ref and any(show(c).strip().upper() == ref for c in codes):
        raise UserError(f"기준 금형({ref})은 삭제할 수 없습니다. 체크를 풀거나 다른 금형을 기준으로 지정하세요.")
    lc = MES.link_col()

    def work(conn):
        h, out = {"c": conn}, []
        try:
            for c in codes:
                used, failed = MES.usage_of_code(h, c, progress, cancelled)
                if cancelled():
                    raise UserError("취소했습니다. (MES는 그대로)")
                out.append({"code": c, "used": used, "failed": failed,
                            "mold": M.App.rid_rows(APP, h["c"], "mold", cc, c),
                            "items": M.App.rid_rows(APP, h["c"], "item", lc, c) if MES.enabled("item") else []})
        finally:
            if h["c"] is not conn:
                try:
                    h["c"].close()
                except Exception:
                    pass
        return out
    info = APP.run_db(work)
    blocked = [x for x in info if x["used"]]
    ok = [x for x in info if not x["used"] and x["mold"]]
    pid = uuid.uuid4().hex[:16]
    _plans[pid] = ok
    return {"plan": pid,
            "blocked": [{"code": show(x["code"]).strip(), "used": [f"{t}.{c}" for t, c, _n in x["used"][:6]]} for x in blocked],
            "ok": [{"code": show(x["code"]).strip(), "items": len(x["items"])} for x in ok],
            "failed": sorted({t for x in ok for t in x["failed"]})}


def api_old_delete_mold(body, user):
    ok = _plans.pop(body.get("plan") or "", None)
    if not ok:
        raise UserError("삭제 확인 정보가 없습니다. [선택 금형 삭제]를 다시 누르세요.")
    one = show(ok[0]["code"]).strip() if len(ok) == 1 else None
    typed = (body.get("typed") or "").strip()
    if (typed.upper() != one.upper()) if one else (typed != "삭제"):
        raise UserError("입력한 글자가 달라서 삭제하지 않았습니다.")
    plan, rows = [], []
    for x in ok:
        cs = show(x["code"]).strip()
        for r in x["items"]:
            plan.append(("item", r["RID__"], f"금형 {cs} - 품번 {show(r.get('ITEM_CODE')).strip()}"))
            rows.append(("item", r))
        for r in x["mold"]:
            plan.append(("mold", r["RID__"], f"금형 {cs}"))
            rows.append(("mold", r))
    APP.run_db(lambda c: M.App.delete_plan(APP, c, plan))
    M.App.log_deleted(APP, rows, dt.datetime.now(), "금형삭제")
    return {"ok": True, "n": len(ok), "ni": sum(len(x["items"]) for x in ok)}


def api_old_items_of(body, user):
    """다른 이름으로 저장 창: 원본 금형의 품번 (ITEM_CODE·캐비티·ST)"""
    md = cache_get(body.get("token"))
    cc = MES.code_col()
    if cc not in md["cols"]:
        raise UserError(f"금형 표에 금형코드 컬럼({cc})이 없습니다.")
    row = md["rows"][int(body.get("idx"))]
    code = row[md["cols"].index(cc)]
    icols, irows, _d = APP.run_db(lambda c: MES.fetch_items_of(c, code))
    edit = [c for c in ("ITEM_CODE", "CAVITY_QTY", "ST_VALUE") if c in icols]
    return {"cols": edit, "all": icols, "rows": [{c: cell_text(r[icols.index(c)]) for c in icols} for r in irows],
            "has_drawing": "DRAWING_NO" in md["cols"],
            "drawing_same": "DRAWING_NO" in md["cols"] and show(row[md["cols"].index("DRAWING_NO")]).strip() == show(code).strip(),
            "has_name": "MOLD_NAME" in md["cols"], "code": show(code).strip()}


def api_old_save_as(body, user):
    """다른 이름으로 저장 (기존 on_save_as_mold 의 save): 원본 그대로, 새 금형 + 품번 추가"""
    md = cache_get(body.get("token"))
    cols, defs = md["cols"], md["defs"]
    cc, lc = MES.code_col(), MES.link_col()
    if cc not in cols:
        raise UserError(f"금형 표에 금형코드 컬럼({cc})이 없습니다.")
    mrow = md["rows"][int(body.get("idx"))]
    old_code = show(mrow[cols.index(cc)]).strip()
    mvals = dict(zip(cols, mrow))
    for col, text in (body.get("mold_changes") or {}).items():
        if col in defs and col != cc and defs[col].get("editable", True):
            mvals[col] = check_value(defs[col], text, defs[col].get("name", col))
    nc = "MOLD_NAME" if "MOLD_NAME" in cols else None
    code, name = str(body.get("new_code") or "").strip(), str(body.get("new_name") or "").strip()
    if not code or code.upper() == old_code.upper():
        raise UserError("새 금형코드를 넣으세요 (원본과 달라야 합니다).")
    if nc and not name:
        raise UserError("새 금형명을 넣으세요.")
    icols, irows, idefs = APP.run_db(lambda c: MES.fetch_items_of(c, mrow[cols.index(cc)]))
    by_orig = {show(r[icols.index("ITEM_CODE")]).strip(): r for r in irows} if "ITEM_CODE" in icols else {}
    use = []
    for it in body.get("items") or []:
        if not it.get("on"):
            continue
        base = by_orig.get(it.get("orig") or "") or (irows[0] if irows else [None] * len(icols))
        vals = dict(zip(icols, base))
        if not it.get("orig"):
            vals["ITEM_CODE"] = None
        for col, text in (it.get("changes") or {}).items():
            if col in idefs and col != lc and idefs[col].get("editable", True):
                vals[col] = None if str(text).strip() == "" else check_value(dict(idefs[col], req=False), text, col)
        use.append(vals)
    codes = [show(v.get("ITEM_CODE")).strip() for v in use]
    if any(not c for c in codes):
        raise UserError("포함(☑)한 품번 중 비어 있는 품번이 있습니다. 넣거나 체크를 푸세요.")
    if len(set(c.upper() for c in codes)) != len(codes):
        raise UserError("같은 품번이 두 번 들어 있습니다.")
    ok_missing = set(body.get("ok_missing") or [])
    reg = reg_user()
    now = dt.datetime.now()
    new_m = dict(mvals)
    new_m[cc] = code
    if nc:
        new_m[nc] = name
    if body.get("dw") and "DRAWING_NO" in new_m:
        new_m["DRAWING_NO"] = code

    def prep(row):
        for c in list(row):
            if any(w in c for w in M.RESET_WORDS):
                row[c] = None
            elif M.RE_AUTO_DATE.search(c) or M.RE_UPD_DATE.search(c):
                if isinstance(row[c], dt.datetime) or row[c] is None:
                    row[c] = now
            elif reg and (M.RE_AUTO_USER.search(c) or M.RE_UPD_USER.search(c)):
                row[c] = reg
        return row
    new_m = prep(new_m)
    new_items = []
    for v in use:
        r = prep(dict(v))
        r[lc] = code
        new_items.append(r)
    tm, ti = MES.table("mold"), MES.table("item")

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"SELECT COUNT(*) FROM {tm} WHERE UPPER(TRIM({ident(cc, '컬럼')})) = :c", c=code.upper())
        if cur.fetchone()[0]:
            raise UserError(f"금형코드 {code} 는 이미 있습니다. 다른 코드를 넣으세요.")
        bad = {}
        for i, ic in enumerate(codes):
            cur.execute(f"SELECT COUNT(*) FROM {mst_tbl()} WHERE ITEM_CODE = :i", i=ic)
            if not cur.fetchone()[0] and ic not in ok_missing:
                bad[i] = ["warn", "품목 마스터에 없음 - 한 번 더 누르면 그대로 저장"]
                continue
            cur.execute(f"SELECT MAX({ident(lc, '컬럼')}) FROM {ti} WHERE ITEM_CODE = :i", i=ic)
            m = cur.fetchone()[0]
            if m:
                bad[i] = ["err", f"이미 금형 {show(m).strip()}에 연결됨 → 새 품번으로"]
        if bad:
            return bad
        mm, mi = table_meta(conn, tm), table_meta(conn, ti)
        try:
            M.insert_dict(cur, tm, mm, new_m)
            for r in new_items:
                M.insert_dict(cur, ti, mi, r)
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        return None
    try:
        bad = APP.run_db(work)
    except UserError:
        raise
    except Exception as e:
        raise RuntimeError("등록 중 오류 -> 전부 취소 (MES는 그대로)\n\n" + ora_hint(e))
    if bad:
        return {"bad": {str(k): v for k, v in bad.items()}, "codes": codes}
    APP.append_log([[now_s(), "", "금형 다른이름저장", tm, f"{old_code} -> {code}",
                     f"품번 {len(new_items)}개: {', '.join(codes)}"]])
    return {"ok": True, "code": code, "old": old_code, "n": len(new_items)}


def api_old_template(body, user):
    """선택행으로 등록양식 만들기 -> 엑셀 대량등록 탭에 불러옴"""
    d = cache_get(body.get("token"))
    idxs = [int(i) for i in body.get("idxs") or []]
    if not idxs:
        raise UserError("복사할 금형을 선택하세요.")
    cc = MES.code_col()
    molds = [dict(zip(d["cols"], d["rows"][i])) for i in idxs]
    codes = [d["rows"][i][d["cols"].index(cc)] for i in idxs]
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    tok = uuid.uuid4().hex[:16]
    path = os.path.join(UPLOAD_DIR, tok + ".xlsx")
    conn = None
    items = []
    try:
        try:
            conn = APP.connect()
        except Exception:
            conn = None
        if conn is not None and MES.enabled("item"):
            for code in codes:
                cols, rows, _ = MES.fetch_items_of(conn, code)
                items += [dict(zip(cols, r)) for r in rows]
        APP.xl.make_template(path, prefill={"mold": molds, "item": items}, conn=conn)
    finally:
        if conn is not None:
            conn.close()
    out = bulk_view(tok)
    out["msg"] = (f"금형 {len(molds)}건, 품번 {len(items)}건을 양식에 복사했습니다.\n\n"
                  "[엑셀 받기]로 내려받아 금형 시트와 품번 시트의 금형코드(주황)를 새 코드로 입력하고\n"
                  "저장한 뒤 [엑셀 불러오기]로 다시 올리고 [① 검사]를 누르세요.")
    return out


# ==============================================================
# 엑셀 대량등록 / 엑셀칸 설정 (숨김 탭)
# ==============================================================
def bulk_load(tok):
    path = upload_path(tok)
    try:
        loaded = APP.xl.load(path)
    except ValueError as e:
        raise UserError(str(e))
    return path, {p: {"rows": loaded[p][0], "info": loaded[p][1], "results": []} for p in M.PARTS}


def bulk_view(tok, data=None):
    if data is None:
        _p, data = bulk_load(tok)
    parts = {}
    for p in M.PARTS:
        if not MES.enabled(p):
            parts[p] = {"enabled": False}
            continue
        used = MES.used(p)
        res_by = {r["rno"]: r for r in data[p]["results"]}
        rows = []
        for rno, row in data[p]["rows"]:
            r = res_by.get(rno)
            rows.append({"rno": rno, "vals": [cell_text(row.get(c["col"])) for c in used],
                         "msg": r["msg"] if r else "", "status": r["status"] if r else ""})
        parts[p] = {"enabled": True, "names": [c["name"] for c in used], "rows": rows,
                    "table": CFG[p].get("table", ""), "sheet": (data[p]["info"] or {}).get("sheet", "")}
    return {"token": tok, "parts": parts, "allow": bool(CFG.get("allow_item_to_existing"))}


def api_bulk_template(body, user):
    try:
        MES.validate_map()
    except ValueError as e:
        raise UserError(str(e))
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, f"tpl_{uuid.uuid4().hex[:8]}.xlsx")
    conn = None
    try:
        try:
            conn = APP.connect()
        except Exception:
            conn = None
        APP.xl.make_template(path, conn=conn)
        with open(path, "rb") as f:
            data = f.read()
    finally:
        if conn is not None:
            conn.close()
        try:
            os.remove(path)
        except Exception:
            pass
    return FileOut(data, f"금형등록_{dt.datetime.now():%Y%m%d}.xlsx", XLSX)


def api_bulk_upload(body, user):
    need_xl()
    tok, _path = save_upload(body)
    out = bulk_view(tok)
    out["name"] = body.get("name", "")
    if not any(out["parts"][p].get("rows") for p in M.PARTS):
        out["msg"] = "입력된 내용이 없습니다."
    return out


def api_bulk_check(body, user):
    tok = body.get("token")
    path, data = bulk_load(tok)
    if not any(data[p]["rows"] for p in M.PARTS):
        raise UserError("입력된 내용이 없습니다.")
    try:
        mres, ires = APP.run_db(lambda c: MES.check(c, data["mold"]["rows"], data["item"]["rows"]))
    except ValueError as e:
        raise UserError(str(e))
    data["mold"]["results"], data["item"]["results"] = mres, ires
    APP.xl.write_results(path, [(data[p]["info"], data[p]["results"]) for p in M.PARTS])
    out = bulk_view(tok, data)
    n = {p: {s: sum(r["status"] == s for r in data[p]["results"]) for s in ("OK", "경고", "오류")} for p in M.PARTS}
    out["msg"] = "검사 완료\n\n" + "\n".join(f"{M.PART_NAME[p]}: 정상 {n[p]['OK']} / 경고 {n[p]['경고']} / 오류 {n[p]['오류']}"
                                         for p in M.PARTS if MES.enabled(p))
    out["nerr"] = sum(n[p]["오류"] for p in M.PARTS)
    if out["nerr"]:
        out["msg"] += "\n\n[엑셀 받기]로 내려받으면 빨간 칸으로 표시되어 있습니다."
    return out


def api_bulk_register(body, user):
    tok = body.get("token")
    path, data = bulk_load(tok)
    if not any(data[p]["rows"] for p in M.PARTS):
        raise UserError("입력된 내용이 없습니다.")

    def work(conn):
        mres, ires = MES.check(conn, data["mold"]["rows"], data["item"]["rows"])
        data["mold"]["results"], data["item"]["results"] = mres, ires
        n = {p: {s: sum(r["status"] == s for r in data[p]["results"]) for s in ("OK", "경고", "오류")} for p in M.PARTS}
        nerr = sum(n[p]["오류"] for p in M.PARTS)
        nwarn = sum(n[p]["경고"] for p in M.PARTS)
        if nerr:
            APP.xl.write_results(path, [(data[p]["info"], data[p]["results"]) for p in M.PARTS])
            return {"nerr": nerr, "msg": f"오류 {nerr}건이 있어 등록하지 않았습니다.\n빨간 칸을 고친 뒤 다시 하세요."}
        if not body.get("confirm"):
            msg = f"금형 {len(mres)}건 -> {MES.table('mold')}"
            if MES.enabled("item"):
                msg += f"\n품번 {len(ires)}건 -> {MES.table('item')}"
            ref = CFG.get("ref_mold") or ""
            for p in M.PARTS:
                auto = MES.last_auto.get(p) or {}
                if auto and data[p]["results"]:
                    how = {"SYSDATE": "현재시각", "{USER}": "등록자", "{COPY}": f"기준금형 {ref} 값"}
                    msg += f"\n\n[{M.PART_NAME[p]}] 자동으로 채우는 필수 컬럼:\n  " + \
                           "\n  ".join(f"{c} = {how[h]}" for c, h in list(auto.items())[:10])
            trg = MES.triggers(conn)
            tl = [f"{M.PART_NAME[p]}: " + ", ".join(v) for p, v in trg.items() if v]
            if tl:
                msg += "\n\n※ 이 테이블에는 MES 자체 자동동작(트리거)이 있습니다:\n  " + "\n  ".join(tl)
            msg += "\n\n새 줄만 추가(INSERT)합니다. 기존 금형·품번 줄은 바꾸지 않습니다.\n위 내용을 MES에 등록합니다."
            if nwarn:
                msg += f"\n※ 경고 {nwarn}건 포함 (노란 줄 확인)"
            return {"preview": True, "msg": msg + "\n\n계속할까요?"}
        return {"done": MES.register(conn, mres, ires)}
    try:
        res = APP.run_db(work)
    except ValueError as e:
        raise UserError(str(e))
    if "done" not in res:
        res.update(bulk_view(tok, data))
        return res
    done = res["done"]
    now = dt.datetime.now()
    M.App.write_log(APP, done, now)
    try:
        APP.xl.write_results(path, [(data[p]["info"], data[p]["results"]) for p in M.PARTS],
                             f"등록완료 {now:%Y-%m-%d %H:%M}")
    except Exception:
        pass
    nm = sum(1 for d in done if d[0] == "mold")
    ok, miss = getattr(MES, "last_verify", (len(done), []))
    msg = f"등록 완료: 금형 {nm}건, 품번 {len(done) - nm}건\n[작업기록]에 기록했습니다.\n\nMES에서 다시 읽어 확인: {ok}/{len(done)}건 확인됨"
    if miss:
        msg += "\n※ 확인 안 됨: " + ", ".join(miss[:10])
    out = bulk_view(tok)
    out.update(ok=True, msg=msg)
    return out


def api_bulk_download(body, user):
    path = upload_path(body.get("token"))
    with open(path, "rb") as f:
        return FileOut(f.read(), body.get("name") or "금형등록_결과.xlsx", XLSX)


def api_map_get(body, user):
    return {"fields": [list(f) for f in M.MAP_FIELDS],
            "parts": {p: {"table": CFG[p].get("table", ""), "codecol": CFG[p].get("code_col" if p == "mold" else "link_col", ""),
                          "enabled": CFG[p].get("enabled", True), "columns": CFG[p]["columns"]} for p in M.PARTS},
            "names": M.PART_NAME}


def api_map_fetch_cols(body, user):
    p, table = body.get("p"), str(body.get("table") or "").strip().upper()
    if p not in M.PARTS or not table:
        raise UserError("테이블명을 넣으세요.")
    cols, code = APP.run_db(lambda c: MES.fetch_columns(c, table, p))
    if not cols:
        raise UserError("테이블을 찾지 못했습니다. 스키마/테이블명을 확인하세요.")
    auto = [f"{c['col']}={c['default']}" for c in cols if not c["use"] and c["default"]]
    msg = f"{table}: {len(cols)}개 컬럼을 불러왔습니다.\n\n금형코드 컬럼: {code or '못 찾음 - 직접 입력하세요'}\n"
    if auto:
        msg += "\n자동으로 기본값 처리한 컬럼 (입력칸에서 뺌):\n - " + "\n - ".join(auto[:12])
        if len(auto) > 12:
            msg += f"\n - 외 {len(auto) - 12}개"
    msg += "\n\n정리할 것:\n- 입력 안 할 컬럼은 사용 N (NOT NULL이면 기본값 필요)\n- {COPY}는 기준 금형 값을 복사합니다\n- 끝나면 [저장]"
    return {"columns": cols, "code": code, "msg": msg}


def api_map_save(body, user):
    parts = body.get("parts") or {}
    warns = []
    for p in M.PARTS:
        m = parts.get(p) or {}
        lst = m.get("columns") or []
        if p == "item" and not m.get("enabled"):
            continue
        try:
            ident(m.get("table"), f"{M.PART_NAME[p]} 테이블")
            for c in lst:
                ident(c["col"], "DB컬럼")
                if c.get("ref"):
                    ident(c["ref"], "존재확인")
        except ValueError as e:
            raise UserError(f"[{M.PART_NAME[p]}] {e}")
        cc = (m.get("codecol") or "").strip().upper()
        c_code = next((c for c in lst if c["col"] == cc), None)
        if not c_code or not c_code["use"]:
            raise UserError(f"[{M.PART_NAME[p]}] 금형코드 컬럼({cc or '미지정'})이 목록에 없거나 사용 N입니다.")
        miss = [c["col"] for c in lst if not c["use"] and c["req"] and not c["default"]]
        if miss:
            warns.append(f"[{M.PART_NAME[p]}] 필수인데 사용 N·기본값 없음: " + ", ".join(miss))
        if not any(c["use"] and c["key"] for c in lst):
            warns.append(f"[{M.PART_NAME[p]}] 중복검사키가 없습니다.")
    with LOCK:
        for p in M.PARTS:
            m = parts.get(p) or {}
            if m.get("columns") is not None:
                CFG[p]["columns"] = [M.mk(bool(c.get("use")), str(c.get("col", "")).upper(), c.get("name", ""),
                                          c.get("type", "TEXT"), bool(c.get("req")), bool(c.get("key")),
                                          int(c.get("maxlen") or 0), c.get("default", ""), str(c.get("ref", "")).upper(),
                                          bool(c.get("sim")), c.get("note", "")) for c in m["columns"]]
        CFG["item"]["enabled"] = bool((parts.get("item") or {}).get("enabled", True))
        APP.save_cfg()
        fresh = M.load_cfg()                    # 금형·품번 테이블 고정값 다시 적용
        CFG["mold"].update(fresh["mold"])
        CFG["item"].update(fresh["item"])
    msg = "저장했습니다.\n예전에 만든 엑셀 양식은 [엑셀 양식 만들기]로 새로 만드세요."
    if warns:
        msg += "\n\n※ 확인 필요 (DB에 기본값이 없으면 등록 실패):\n" + "\n".join(warns)
    return {"ok": True, "msg": msg}


def api_map_find_tables(body, user):
    found = APP.run_db(MES.find_mold_tables)
    return {"cols": ["추천", "테이블", "설명", "기본키(PK)", "참조수", "컬럼수", "행수"],
            "rows": [[("★" if f["score"] >= 50 else ""), f["table"], f["cmt"], f["pk"], f["refcnt"], f["colcnt"],
                      "" if f["rows"] is None else f"{f['rows']:,}" + ("" if f["exact"] else " (추정)")] for f in found]}


def api_table_preview(body, user):
    t = str(body.get("table") or "").strip().upper()
    try:
        hdr, rows = APP.run_db(lambda c: MES.preview(c, t))
    except ValueError as e:
        raise UserError(str(e))
    return {"cols": hdr, "rows": rows}


# ==============================================================
# 금형 간편등록 (숨김 탭)
# ==============================================================
QUICK_FIELDS = M.QUICK_FIELDS


def quick_T(key):
    return {"mold": MES.table("mold"), "link": MES.table("item"),
            "bom": MES.qualify(CFG.get("bom_table") or "ICOM_ITEM_CHILD"), "mst": mst_tbl()}[key]


def quick_base(conn, code):
    f1 = M.QuickTab.fetch_one
    mold, md = f1(conn, f"SELECT * FROM {quick_T('mold')} WHERE TRIM(MOLD_CODE) = :c AND ROWNUM = 1", {"c": code})
    link, ld = f1(conn, f"SELECT * FROM {quick_T('link')} WHERE TRIM(MOLD_CODE) = :c AND ROWNUM = 1", {"c": code})
    bom, bd = None, None
    if link:
        bom, bd = f1(conn, f"SELECT * FROM {quick_T('bom')} WHERE ITEM_CODE = :i AND ROWNUM = 1", {"i": link["ITEM_CODE"]})
    if not bd:
        _, bd = f1(conn, f"SELECT * FROM {quick_T('bom')} WHERE ROWNUM = 1", {})
        bom = None
    if not ld:
        _, ld = f1(conn, f"SELECT * FROM {quick_T('link')} WHERE ROWNUM = 1", {})
    if not mold:
        raise UserError(f"기준 금형 {code}를 찾지 못했습니다.")
    return {"mold": mold, "mold_desc": md, "link": link, "link_desc": ld, "bom": bom, "bom_desc": bd}


def api_quick_find(body, user):
    q = str(body.get("q") or "").strip().upper()

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"SELECT * FROM (SELECT MOLD_CODE, MOLD_NAME FROM {quick_T('mold')} WHERE UPPER(MOLD_CODE) LIKE :q "
                    f"OR UPPER(MOLD_NAME) LIKE :q ORDER BY MOLD_CODE) WHERE ROWNUM <= 500", q=f"%{q}%")
        return cur.fetchall()
    return {"rows": [[show(a).strip(), show(b).strip()] for a, b in APP.run_db(work)]}


def api_quick_base(body, user):
    code = str(body.get("code") or "").strip()
    base = APP.run_db(lambda c: quick_base(c, code))
    cols = {d[0] for d in base["mold_desc"]}
    vals = {}
    for col, _l in QUICK_FIELDS:
        if col == "MOLD_CODE":
            continue
        if col not in cols and col != "CAVITY_QTY":
            vals[col] = "(이 테이블에 없음)"
            continue
        src = base["mold"].get(col) if col in cols else (base["link"] or {}).get("CAVITY_QTY")
        vals[col] = show(src).strip() if src is not None else ""
    mold = base["mold"]
    return {"vals": vals, "name": show(mold.get("MOLD_NAME")).strip(), "has_link": bool(base["link"]),
            "drawing_follows": show(mold.get("DRAWING_NO")).strip() == show(mold.get("MOLD_CODE")).strip()}


def api_quick_item(body, user):
    code = str(body.get("item") or "").strip()
    bcode = str(body.get("base") or "").strip()
    if not code:
        raise UserError("품번을 넣으세요.")

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"SELECT MAX(ITEM_NAME) FROM {quick_T('mst')} WHERE ITEM_CODE = :i", i=code)
        name = cur.fetchone()[0]
        cur.execute(f"SELECT CHILD_ITEM_CODE, UNIT_PER_QTY FROM {quick_T('bom')} WHERE ITEM_CODE = :i "
                    f"ORDER BY NVL(UNIT_PER_QTY, 0) DESC", i=code)
        boms = cur.fetchall()
        base = quick_base(conn, bcode) if bcode else None
        return name, boms, base
    name, boms, base = APP.run_db(work)
    it = {"item": code, "name": show(name) if name else "(품목 마스터에 없음)", "raw": "", "qpu": "", "note": ""}
    if boms:
        it["raw"] = show(boms[0][0]).strip()
        it["qpu"] = show(boms[0][1])
        it["note"] = f"기존 단위중량 있음 ({len(boms)}건) - 그대로 둡니다"
    elif base and base.get("bom"):
        it["raw"] = show(base["bom"].get("CHILD_ITEM_CODE")).strip()
        it["note"] = "원소재는 기준 금형 값을 넣었습니다. 확인 후 단위중량을 입력하세요"
    return it


def quick_obj(body):
    """기존 QuickTab.build() 를 쓰기 위한 껍데기"""
    q = object.__new__(M.QuickTab)
    q.app, q.mes = APP, MES
    code = str(body.get("base") or "").strip()
    if not code:
        raise UserError("① 비슷한 기존 금형을 먼저 고르세요.")
    q.base = APP.run_db(lambda c: quick_base(c, code))
    f = body.get("fields") or {}
    q.fvars = {col: FV(str(f.get(col) or "")) for col, _l in QUICK_FIELDS}
    q.items = []
    for it in body.get("items") or []:
        x = {"item": str(it.get("item") or "").strip(), "name": it.get("name", ""), "raw": str(it.get("raw") or "").strip()}
        for k in ("qpu", "cav"):
            t = str(it.get(k) or "").strip()
            try:
                x[k] = conv({"type": "NUM"}, t) if t else None
                if x[k] is not None and x[k] < 0:
                    raise ValueError
            except ValueError:
                raise UserError(f"{x['item']}: {'단위중량' if k == 'qpu' else '캐비티'}은 숫자로 넣으세요.")
        q.items.append(x)
    return q


def quick_check(q):
    """기존 QuickTab.check 와 같은 검사. return 오류 문장들 (줄별 결과는 q.items 에)"""
    code = q.fvars["MOLD_CODE"].get().strip()
    name = q.fvars["MOLD_NAME"].get().strip()
    errs = []
    if not code:
        errs.append("금형코드를 입력하세요.")
    if not name:
        errs.append("금형명을 입력하세요.")
    mcols = {d[0]: d for d in q.base["mold_desc"]}
    if code and "MOLD_CODE" in mcols and mcols["MOLD_CODE"][3] and blen(code, int(CFG.get("kor_bytes") or 3)) > mcols["MOLD_CODE"][3]:
        errs.append(f"금형코드가 너무 깁니다 (최대 {mcols['MOLD_CODE'][3]}byte).")
    seen = set()
    for it in q.items:
        if it["item"].upper() in seen:
            errs.append(f"품번 {it['item']}가 두 번 있습니다.")
        seen.add(it["item"].upper())

    def work(conn):
        cur = conn.cursor()
        exists = False
        if code:
            cur.execute(f"SELECT COUNT(*) FROM {quick_T('mold')} WHERE TRIM(MOLD_CODE) = :c", c=code)
            exists = cur.fetchone()[0] > 0
        for it in q.items:
            it["level"], notes = "ok", []
            cur.execute(f"SELECT MAX(TRIM(MOLD_CODE)) FROM {quick_T('link')} WHERE ITEM_CODE = :i", i=it["item"])
            other = cur.fetchone()[0]
            if other:
                it["level"] = "err"
                notes.append(f"이미 금형 {other}에 연결된 품번 (MES 규칙: 품번 하나 = 금형 하나)")
            cur.execute(f"SELECT COUNT(*) FROM {quick_T('mst')} WHERE ITEM_CODE = :i", i=it["item"])
            if not cur.fetchone()[0]:
                it["level"] = "err" if it["level"] == "err" else "warn"
                notes.append("품목 마스터에 없는 품번 (오타 확인)")
            it["bom_action"] = None
            if it.get("raw"):
                cur.execute(f"SELECT UNIT_PER_QTY FROM {quick_T('bom')} WHERE ITEM_CODE = :i AND CHILD_ITEM_CODE = :r",
                            i=it["item"], r=it["raw"])
                ex = cur.fetchall()
                if ex:
                    it["bom_action"] = "keep"
                    notes.append(f"단위중량 기존값 {show(ex[0][0])} 유지 (바꾸려면 [단위중량 관리] 탭)")
                elif it.get("qpu"):
                    it["bom_action"] = "insert"
                    notes.append(f"단위중량 {show(it['qpu'])} 새로 등록")
                    cur.execute(f"SELECT COUNT(*) FROM {quick_T('mst')} WHERE ITEM_CODE = :r", r=it["raw"])
                    if not cur.fetchone()[0]:
                        it["level"] = "err" if it["level"] == "err" else "warn"
                        notes.append("원소재 코드가 품목 마스터에 없음 (코일 코드 확인)")
                else:
                    it["level"] = "err" if it["level"] == "err" else "warn"
                    notes.append("단위중량이 비어 있음 → 소재 사용량이 0으로 쌓입니다")
            else:
                it["level"] = "err" if it["level"] == "err" else "warn"
                notes.append("원소재·단위중량 미입력 → 소재 사용량이 0으로 쌓입니다")
            it["note"] = " / ".join(notes) if notes else "OK"
        return exists
    if APP.run_db(work):
        errs.append(f"금형코드 {code}는 이미 있습니다.")
    if not q.items:
        errs.append("품번이 없습니다. 품번이 없으면 작업지시를 낼 수 없습니다.")
    return errs


def quick_items_out(q):
    return [{"item": it["item"], "name": it.get("name", ""), "raw": it.get("raw", ""),
             "qpu": show(it.get("qpu")) if it.get("qpu") is not None else "",
             "cav": show(it.get("cav")) if it.get("cav") is not None else "",
             "note": it.get("note", ""), "level": it.get("level", "")} for it in q.items]


def api_quick_check(body, user):
    q = quick_obj(body)
    errs = quick_check(q)
    nerr = len(errs) + sum(1 for it in q.items if it["level"] == "err")
    nwarn = sum(1 for it in q.items if it["level"] == "warn")
    return {"errs": errs, "items": quick_items_out(q), "nerr": nerr, "nwarn": nwarn}


def api_quick_register(body, user):
    q = quick_obj(body)
    errs = quick_check(q)
    nerr = len(errs) + sum(1 for it in q.items if it["level"] == "err")
    if nerr:
        return {"errs": errs or ["④ 검사에서 오류가 있습니다. 빨간 줄과 메시지를 확인하세요."], "items": quick_items_out(q),
                "nerr": nerr, "nwarn": 0}
    steps = q.build()
    summary = "\n".join(f"  {kind}: " + " / ".join(show(v.get(k)) for k in ("MOLD_CODE", "ITEM_CODE", "CHILD_ITEM_CODE", "UNIT_PER_QTY") if k in v)
                        for kind, _, v, _ in steps)
    if not body.get("confirm"):
        return {"preview": True, "msg": f"아래를 한 번에 추가합니다 (기존 내용은 바꾸지 않음):\n\n{summary}\n\n"
                                        f"한 줄이라도 실패하면 전부 취소됩니다. 계속할까요?"}
    APP.run_db(lambda conn: MES.insert_many(conn, [(t, v, a) for _, t, v, a in steps]))
    APP.append_log([[now_s(), "", f"간편{kind}등록", t,
                     "/".join(show(v.get(k)) for k in ("MOLD_CODE", "ITEM_CODE", "CHILD_ITEM_CODE") if k in v),
                     "; ".join(f"{k}={cell_text(x)}" for k, x in v.items())] for kind, t, v, a in steps])
    code = q.fvars["MOLD_CODE"].get().strip()
    return {"ok": True, "msg": f"등록 완료: 금형 {code} (품번 {len(q.items)}개, 단위중량 {sum(1 for s in steps if s[0] == '단위중량')}건)\n"
                               f"[③ 작업지시 관리]에서 이 품번으로 작업지시를 낼 수 있습니다."}


# ==============================================================
# ④ 소재 관리 (기존 ReceiptTab / ReceiptForm 과 같은 규칙)
# ==============================================================
RCV = object.__new__(M.ReceiptTab)          # 단위 판단·재고 계산·삭제 함수만 씀
RCV.app, RCV.rows, RCV.unit, RCV.unit_note, RCV.unit_ok = APP, [], "kg", "아직 판단 전", False
T_RCV, T_LOT, T_STK = M.ReceiptTab.T_RCV, M.ReceiptTab.T_LOT, M.ReceiptTab.T_STK


def rcv_sql(f):
    d1, d2 = parse_day_u(f.get("d1")), parse_day_u(f.get("d2"))
    if d1 > d2:
        d1, d2 = d2, d1
    s = f"""
        SELECT R.RECEIPT_NO, R.SEQ, R.RECEIPT_DATE, R.ITEM_CODE,
               (SELECT MAX(M.ITEM_NAME) FROM ICOM_ITEM_MASTER M WHERE M.ITEM_CODE = R.ITEM_CODE) ITEM_NAME,
               R.HEAT_NO, R.VENDOR_SITE_ID, R.QTY, R.QTY_UNIT, R.NET_WEIGHT, R.TOTAL_WEIGHT,
               R.IO_FLAG, R.ENTER_BY, R.ENTER_DATE,
               (SELECT COUNT(*) FROM {M.ReceiptTab.T_USE} U WHERE U.RECEIPT_NO = R.RECEIPT_NO
                   AND U.RECEIPT_SEQ = R.SEQ AND NVL(U.HEAT_NO, '~') = NVL(R.HEAT_NO, '~')) USE_CNT,
               (SELECT COUNT(*) FROM {M.ReceiptTab.T_ISS} I WHERE I.RECEIPT_NO = R.RECEIPT_NO
                   AND I.RECEIPT_SEQ = R.SEQ AND NVL(I.HEAT_NO, '~') = NVL(R.HEAT_NO, '~')) ISS_CNT,
               (SELECT COUNT(*) FROM {T_RCV} X WHERE X.RECEIPT_NO = R.RECEIPT_NO
                   AND X.SEQ = R.SEQ AND X.IO_FLAG = 2) CANCEL_CNT
        FROM {T_RCV} R
        WHERE R.RECEIPT_DATE >= :sd AND R.RECEIPT_DATE < :ed"""
    p = {"sd": dt.datetime.combine(d1, dt.time()), "ed": dt.datetime.combine(d2 + dt.timedelta(days=1), dt.time())}
    if not f.get("cancel"):
        s += " AND R.IO_FLAG = 1"
    v = {k: str(f.get(k) or "").strip() for k in ("item", "heat", "rno", "vendor")}
    if v["item"]:
        s += (" AND (UPPER(R.ITEM_CODE) LIKE :item OR EXISTS (SELECT 1 FROM ICOM_ITEM_MASTER M "
              "WHERE M.ITEM_CODE = R.ITEM_CODE AND UPPER(M.ITEM_NAME) LIKE :item))")
        p["item"] = f"%{v['item'].upper()}%"
    if v["heat"]:
        s += " AND UPPER(R.HEAT_NO) LIKE :heat"
        p["heat"] = f"%{v['heat'].upper()}%"
    if v["rno"]:
        s += " AND TO_CHAR(R.RECEIPT_NO) LIKE :rno"
        p["rno"] = f"%{v['rno']}%"
    if v["vendor"]:
        s += " AND UPPER(R.VENDOR_SITE_ID) LIKE :vd"
        p["vd"] = f"%{v['vendor'].upper()}%"
    s += " ORDER BY R.RECEIPT_DATE DESC, R.RECEIPT_NO DESC, R.SEQ, R.IO_FLAG"
    return s, p


def rcv_disp(r):
    st, tag = M.ReceiptTab.state(r)
    io_ = {1: "입고", 2: "취소"}.get(r["IO_FLAG"], show(r["IO_FLAG"]))
    kg = lambda v: f"{RCV.to_kg(v):,.1f}" if v is not None else ""  # noqa: E731
    return [st, show(r["RECEIPT_NO"]), show(r["SEQ"]),
            r["RECEIPT_DATE"].strftime("%Y-%m-%d") if r["RECEIPT_DATE"] else "",
            show(r["ITEM_CODE"]), show(r["ITEM_NAME"]), show(r["HEAT_NO"]), show(r["VENDOR_SITE_ID"]),
            show(r["QTY"]), show(r["QTY_UNIT"]), kg(r["NET_WEIGHT"]), kg(r["TOTAL_WEIGHT"]),
            io_, f"{r['USE_CNT'] or 0}/{r['ISS_CNT'] or 0}", show(r["ENTER_BY"]),
            r["ENTER_DATE"].strftime("%Y-%m-%d %H:%M") if r["ENTER_DATE"] else ""], tag


def api_rcv_search(body, user):
    s, p = rcv_sql(body)

    def work(conn):
        RCV.unit_ok = False
        RCV.detect_unit(conn)
        cur = conn.cursor()
        cur.execute(s, p)
        names = [d[0] for d in cur.description]
        return [dict(zip(names, r)) for r in cur.fetchall()]
    rows = APP.run_db(work)
    disp = [rcv_disp(r) for r in rows]
    ins = [r for r in rows if r["IO_FLAG"] == 1]
    kg = sum((RCV.to_kg(r["NET_WEIGHT"]) or 0) for r in ins)
    ok = sum(1 for _d, t in disp if t == "ok")
    off = sum(1 for _d, t in disp if t == "off")
    return {"token": cache_put({"rcv": rows}), "cols": M.ReceiptTab.COLS, "rows": [d for d, _t in disp],
            "tags": [t for _d, t in disp], "io": [r["IO_FLAG"] for r in rows],
            "summary": f"{len(rows)}건  (입고 {len(ins)}건 / 순중량 합계 {kg:,.1f} kg)   수정·삭제 가능 {ok}건  ·  사용됨 {off}건"
                       f"      (MES 순중량 단위: {RCV.unit} - {RCV.unit_note})"}


def rcv_orig(conn, row):
    cur = conn.cursor()
    cur.execute(f"SELECT ROWIDTOCHAR(ROWID) RID__, R.* FROM {T_RCV} R "
                f"WHERE RECEIPT_NO = :rno AND SEQ = :seq AND IO_FLAG = 1", rno=row["RECEIPT_NO"], seq=row["SEQ"])
    rows = cur.fetchall()
    return dict(zip([d[0] for d in cur.description], rows[0])) if len(rows) == 1 else None


def api_rcv_form(body, user):
    """입력창 기초 자료 (기존 ReceiptForm.load)"""
    row = None
    if body.get("token") and body.get("idx") is not None:
        row = cache_get(body["token"])["rcv"][int(body["idx"])]
        if row["IO_FLAG"] != 1:
            raise UserError("취소(반품) 줄은 고칠 수 없습니다. 원래 입고 줄을 선택하세요.")

    def work(conn):
        RCV.detect_unit(conn)
        m_rcv = table_meta(conn, T_RCV)
        cur = conn.cursor()
        cur.execute(f"SELECT ITEM_CODE, COUNT(*) FROM {T_RCV} WHERE RECEIPT_DATE >= SYSDATE - 365 AND IO_FLAG = 1 "
                    f"GROUP BY ITEM_CODE ORDER BY 2 DESC")
        items = [show(a) for a, _ in cur.fetchall() if a]
        vendors = []
        try:
            cur.execute("SELECT VENDOR_SITE_ID, MAX(VENDOR_NAME) FROM ICOM_VENDORS GROUP BY VENDOR_SITE_ID "
                        "ORDER BY MAX(VENDOR_NAME)")
            vendors = [(show(a), show(b)) for a, b in cur.fetchall() if a is not None]
        except Exception:
            pass
        if not vendors:
            cur.execute(f"SELECT DISTINCT VENDOR_SITE_ID FROM {T_RCV} WHERE VENDOR_SITE_ID IS NOT NULL "
                        f"AND RECEIPT_DATE >= SYSDATE - 365")
            vendors = [(show(a), "") for (a,) in cur.fetchall()]
        order = "ENTER_DATE DESC NULLS LAST" if "ENTER_DATE" in m_rcv else "RECEIPT_NO DESC"
        cur.execute(f"SELECT RECEIPT_NO FROM (SELECT RECEIPT_NO FROM {T_RCV} ORDER BY {order}) WHERE ROWNUM = 1")
        r = cur.fetchone()
        orig, used = None, {}
        if row is not None:
            orig = rcv_orig(conn, row)
            if orig is None:
                raise UserError("선택한 입고(입고 줄)를 MES에서 찾지 못했습니다. 다시 조회하세요.")
            used = RCV.usage(conn, orig)
        return items, vendors, (show(r[0]) if r else "(없음)"), orig, used
    items, vendors, last_no, orig, used = APP.run_db(work)
    o = orig or {}
    kg0 = RCV.to_kg(o.get("NET_WEIGHT")) if o else None
    vl = lambda vid: next((f"{a} | {b}" if b else a for a, b in vendors if a == vid), vid)  # noqa: E731
    return {"items": items, "vendors": [f"{a} | {b}" if b else a for a, b in vendors], "last_no": last_no,
            "unit": RCV.unit, "unit_note": RCV.unit_note, "new": orig is None, "used": used,
            "vals": {"date": o["RECEIPT_DATE"].strftime("%Y-%m-%d") if o.get("RECEIPT_DATE") else "",
                     "item": show(o.get("ITEM_CODE")), "heat": show(o.get("HEAT_NO")),
                     "vendor": vl(show(o.get("VENDOR_SITE_ID"))) if o else "",
                     "kg": "" if kg0 is None else f"{kg0:g}", "qty": show(o.get("QTY")) if o else "1",
                     "charger": show(o.get("CHARGER")) if o else reg_user("ADMIN"),
                     "no": show(o.get("RECEIPT_NO"))}}


def api_rcv_next_no(body, user):
    d = parse_day_u(body.get("date"))
    no, same = APP.run_db(lambda c: RCV.next_no(c, d))
    return {"no": no, "same": show(same) if same else ""}


def api_rcv_item_lookup(body, user):
    it = str(body.get("item") or "").strip()
    if not it:
        return {"n": 0, "text": ""}

    def work(conn):
        cur = conn.cursor()
        cur.execute("SELECT MAX(ITEM_NAME), COUNT(*) FROM ICOM_ITEM_MASTER WHERE ITEM_CODE = :i", i=it)
        name, n = cur.fetchone()
        cur.execute(f"SELECT VENDOR_SITE_ID, NET_WEIGHT FROM (SELECT VENDOR_SITE_ID, NET_WEIGHT FROM {T_RCV} "
                    f"WHERE ITEM_CODE = :i AND IO_FLAG = 1 ORDER BY RECEIPT_DATE DESC, RECEIPT_NO DESC) "
                    f"WHERE ROWNUM = 1", i=it)
        last = cur.fetchone()
        cur.execute(f"SELECT NVL(SUM(QTY), 0), NVL(SUM(STOCK_WEIGHT_KG), 0) FROM {T_STK} WHERE ITEM_CODE = :i", i=it)
        return name, n, last, cur.fetchone()
    name, n, last, stk = APP.run_db(work)
    if not n:
        return {"n": 0, "text": f"품목 마스터에 없는 품번입니다: {it}  (ERP 품목 등록 확인)"}
    parts = [f"품명: {show(name)}", f"현재 재고: {show(stk[0])}롤 / {float(stk[1] or 0):,.1f}kg"]
    if last:
        parts.append(f"최근 업체: {show(last[0])}")
    return {"n": n, "text": "   ".join(parts), "last_vendor": show(last[0]) if last and last[0] is not None else ""}


def api_rcv_item_find(body, user):
    q = str(body.get("q") or "").strip().upper()

    def work(conn):
        cur = conn.cursor()
        cur.execute("SELECT * FROM (SELECT ITEM_CODE, MAX(ITEM_NAME) FROM ICOM_ITEM_MASTER "
                    "WHERE UPPER(ITEM_CODE) LIKE :q OR UPPER(ITEM_NAME) LIKE :q "
                    "GROUP BY ITEM_CODE ORDER BY ITEM_CODE) WHERE ROWNUM <= 300", q=f"%{q}%")
        return cur.fetchall()
    return {"rows": rows_text(APP.run_db(work))}


def rcv_inputs(vin, orig=None):
    d = parse_day_u(vin.get("date"))
    od = (orig or {}).get("RECEIPT_DATE")
    v = {"RECEIPT_DATE": od if od and od.date() == d else dt.datetime.combine(d, dt.time()),
         "ITEM_CODE": str(vin.get("item") or "").strip(), "HEAT_NO": str(vin.get("heat") or "").strip() or None,
         "VENDOR_SITE_ID": str(vin.get("vendor") or "").split(" | ")[0].strip() or None,
         "CHARGER": str(vin.get("charger") or "").strip() or None}
    if not v["ITEM_CODE"]:
        raise UserError("품번을 넣으세요.")
    if not v["HEAT_NO"]:
        raise UserError("HEAT_NO(소재로트)를 넣으세요.")
    try:
        kg = float(str(vin.get("kg") or "").replace(",", ""))
        qty = float(str(vin.get("qty") or "").replace(",", ""))
    except ValueError:
        raise UserError("중량(kg)과 수량은 숫자로 넣으세요.")
    if kg <= 0 or qty <= 0:
        raise UserError("중량과 수량은 0보다 커야 합니다.")
    if RCV.unit == "kg" and kg > 30000:
        raise UserError(f"중량 {kg:,} kg 이 너무 큽니다. kg 단위로 넣으세요.")
    qty = int(qty) if qty.is_integer() else qty
    net = RCV.from_kg(kg)
    if isinstance(net, float) and net.is_integer():
        net = int(net)
    v.update(QTY=qty, NET_WEIGHT=net)
    v["_kg"] = kg
    return v


def api_rcv_save_new(body, user):
    vin = body.get("values") or {}
    no = str(vin.get("no") or "").strip()

    def work(conn):
        RCV.detect_unit(conn)
        v = rcv_inputs(vin)
        if not no:
            raise UserError("입고번호를 넣으세요.")
        m_rcv, m_lot, m_stk = (table_meta(conn, x) for x in (T_RCV, T_LOT, T_STK))
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*) FROM ICOM_ITEM_MASTER WHERE ITEM_CODE = :i", i=v["ITEM_CODE"])
        if not cur.fetchone()[0]:
            raise UserError(f"품목 마스터에 없는 품번입니다: {v['ITEM_CODE']}")
        cur.execute(f"SELECT COUNT(*) FROM {T_RCV} WHERE TO_CHAR(RECEIPT_NO) = :n", n=no)
        if cur.fetchone()[0]:
            raise UserError(f"입고번호 {no}는 이미 있습니다.")
        warn = []
        cur.execute(f"SELECT RECEIPT_NO, RECEIPT_DATE FROM {T_RCV} WHERE ITEM_CODE = :i AND HEAT_NO = :h "
                    f"AND IO_FLAG = 1", i=v["ITEM_CODE"], h=v["HEAT_NO"])
        same = cur.fetchall()
        if same:
            warn.append("같은 품번·HEAT_NO 입고가 이미 있습니다: " + ", ".join(
                f"{show(a)}({b:%m/%d})" if b else show(a) for a, b in same[:5]))
        cur.execute(f"SELECT COUNT(*) FROM {T_STK} WHERE ITEM_CODE = :i", i=v["ITEM_CODE"])
        n = cur.fetchone()[0]
        if n > 1:
            raise UserError(f"소재 재고({T_STK})에 품번 {v['ITEM_CODE']} 줄이 {n}개라 재고를 더할 수 없습니다.")
        if not n:
            warn.append("이 소재는 재고 줄이 없어 새로 만듭니다.")
        if not body.get("confirm"):
            msg = (f"MES에 소재 입고 1건을 추가합니다.\n\n  입고번호: {no}\n  입고일: {v['RECEIPT_DATE']:%Y-%m-%d}\n"
                   f"  품번: {v['ITEM_CODE']}\n  HEAT_NO: {v['HEAT_NO']}\n  업체: {show(v['VENDOR_SITE_ID'])}\n"
                   f"  수량: {v['QTY']}롤   중량: {v['_kg']:,} kg\n\n입고 + 재고LOT 추가, 소재 재고에 더하기 (기존 입고는 안 바뀜)")
            if warn:
                msg += "\n\n※ 확인하세요:\n  " + "\n  ".join(warn)
            return {"preview": True, "msg": msg + "\n\n등록할까요?", "warn": bool(warn)}
        who = reg_user() or v["CHARGER"] or "ADMIN"
        now = dt.datetime.now()
        try:
            cur.execute(f"SELECT * FROM (SELECT * FROM {T_RCV} WHERE ITEM_CODE = :i AND IO_FLAG = 1 "
                        f"ORDER BY RECEIPT_DATE DESC) WHERE ROWNUM = 1", i=v["ITEM_CODE"])
            r = cur.fetchone()
            tpl = dict(zip([d[0] for d in cur.description], r)) if r else {}
            org = tpl.get("ORGANIZATION_ID") or 1
            unit = tpl.get("QTY_UNIT") or "RL"
            rcv = {c: tpl.get(c) for c, m in m_rcv.items() if m["notnull"] and not m["hasdef"]}
            rcv.update(RECEIPT_NO=no, SEQ=1, RECEIPT_DATE=v["RECEIPT_DATE"], ITEM_CODE=v["ITEM_CODE"],
                       VENDOR_SITE_ID=v["VENDOR_SITE_ID"], QTY=v["QTY"], QTY_UNIT=unit,
                       NET_WEIGHT=v["NET_WEIGHT"], TOTAL_WEIGHT=v["NET_WEIGHT"], HEAT_NO=v["HEAT_NO"],
                       CHARGER=v["CHARGER"], ORGANIZATION_ID=org, ENTER_DATE=now, ENTER_BY=who,
                       LAST_MODIFY_DATE=now, LAST_MODIFY_BY=who, IO_FLAG=1, DIVIDE_FLAG="N")
            M.insert_dict(cur, T_RCV, m_rcv, rcv)
            if "ENTER_DATE" in m_lot:
                cur.execute(f"SELECT * FROM (SELECT * FROM {T_LOT} ORDER BY ENTER_DATE DESC) WHERE ROWNUM = 1")
            else:
                cur.execute(f"SELECT * FROM {T_LOT} WHERE ROWNUM = 1")
            r = cur.fetchone()
            ltpl = dict(zip([d[0] for d in cur.description], r)) if r else {}
            lot = {c: ltpl.get(c) for c, m in m_lot.items() if m["notnull"] and not m["hasdef"]}
            lot.update(RECEIPT_NO=no, RECEIPT_SEQ=1, HEAT_NO=v["HEAT_NO"], ITEM_CODE=v["ITEM_CODE"],
                       VENDOR_SITE_ID=v["VENDOR_SITE_ID"], QTY=v["QTY"], QTY_UNIT=unit,
                       UNIT_WEIGHT=v["NET_WEIGHT"], TOTAL_WEIGHT=v["NET_WEIGHT"], ORGANIZATION_ID=org,
                       ENTER_DATE=now, ENTER_BY=who, LAST_MODIFY_DATE=now, LAST_MODIFY_BY=who, DIVIDE_FLAG="N")
            M.insert_dict(cur, T_LOT, m_lot, lot)
            smsg = RCV.stock_add(cur, v["ITEM_CODE"], v["QTY"], v["_kg"], who, org=org, unit=unit, meta=m_stk)
            conn.commit()
        except Exception as e:
            conn.rollback()
            raise RuntimeError("등록 중 오류 -> 전부 취소 (MES는 그대로)\n\n" + ora_hint(e))
        return {"ok": True, "no": no, "smsg": smsg, "date": f"{v['RECEIPT_DATE']:%Y-%m-%d}",
                "log": f"{v['ITEM_CODE']} HEAT {v['HEAT_NO']} {v['QTY']}롤 {v['_kg']}kg; {smsg}"}
    res = APP.run_db(work)
    if res.get("ok"):
        APP.append_log([[now_s(), "", "소재입고등록", T_RCV, no, res.pop("log")]])
    return res


def api_rcv_save_edit(body, user):
    row = cache_get(body.get("token"))["rcv"][int(body.get("idx"))]
    vin = body.get("values") or {}

    def work(conn):
        RCV.detect_unit(conn)
        o = rcv_orig(conn, row)
        if o is None:
            raise UserError("선택한 입고를 MES에서 찾지 못했습니다. 다시 조회하세요.")
        m_rcv, m_lot, m_stk = (table_meta(conn, x) for x in (T_RCV, T_LOT, T_STK))
        used = RCV.usage(conn, o)
        v = rcv_inputs(vin, o)
        kg_new = v.pop("_kg")
        ch = {c: (o.get(c), x) for c, x in v.items() if c in o and cell_text(o.get(c)) != cell_text(x)}
        if "NET_WEIGHT" in ch and "TOTAL_WEIGHT" in o and cell_text(o.get("TOTAL_WEIGHT")) == cell_text(o.get("NET_WEIGHT")):
            ch["TOTAL_WEIGHT"] = (o["TOTAL_WEIGHT"], v["NET_WEIGHT"])
        if not ch:
            raise UserError("바뀐 내용이 없습니다.")
        if used and set(ch) & set(M.ReceiptForm.FIELDS_LOCK_USED):
            raise UserError("사용·출고·취소 이력이 있는 입고는 업체·담당자만 고칠 수 있습니다.")
        names = {"RECEIPT_DATE": "입고일", "ITEM_CODE": "품번", "HEAT_NO": "HEAT_NO", "VENDOR_SITE_ID": "업체",
                 "QTY": "수량", "NET_WEIGHT": f"순중량({RCV.unit})", "TOTAL_WEIGHT": f"총중량({RCV.unit})", "CHARGER": "담당자"}
        stock_ch = bool(set(ch) & {"ITEM_CODE", "QTY", "NET_WEIGHT"})
        if not body.get("confirm"):
            lines = "\n".join(f"  {names.get(c, c)}: {cell_text(a) or '(빈칸)'} → {cell_text(b) or '(빈칸)'}" for c, (a, b) in ch.items())
            return {"preview": True, "msg": f"입고 {show(o['RECEIPT_NO'])}를 아래처럼 고칩니다.\n\n{lines}\n\n재고LOT도 같이 고칩니다."
                                            + ("\n소재 재고도 차이만큼 맞춥니다." if stock_ch else "") + "\n\n저장할까요?"}
        who = reg_user("PYEDIT")
        cur = conn.cursor()
        try:
            sets, where, b = [], ["ROWID = CHARTOROWID(:rid)"], {"rid": o["RID__"]}
            for i, (c, (a, x)) in enumerate(ch.items()):
                sets.append(f"{c} = :n{i}")
                b[f"n{i}"] = M.fit_type(m_rcv, c, x)
                if a is None:
                    where.append(f"{c} IS NULL")
                else:
                    where.append(f"{c} = :o{i}")
                    b[f"o{i}"] = a
            for c, how in (("LAST_MODIFY_DATE", "SYSDATE"), ("LAST_MODIFY_BY", who)):
                if c in m_rcv:
                    sets.append(f"{c} = SYSDATE" if how == "SYSDATE" else f"{c} = :mb")
                    if how != "SYSDATE":
                        b["mb"] = how
            cur.execute(f"UPDATE {T_RCV} SET {', '.join(sets)} WHERE {' AND '.join(where)}", b)
            if cur.rowcount != 1:
                raise UserError("그 사이 다른 곳에서 바뀌었거나 없어진 입고입니다. 다시 조회하세요.")
            lmap = {"HEAT_NO": "HEAT_NO", "ITEM_CODE": "ITEM_CODE", "VENDOR_SITE_ID": "VENDOR_SITE_ID",
                    "QTY": "QTY", "NET_WEIGHT": "UNIT_WEIGHT"}
            lset, lb = [], {"rno": o["RECEIPT_NO"], "seq": o["SEQ"], "heat": o.get("HEAT_NO")}
            for c, lc in lmap.items():
                if c in ch and lc in m_lot:
                    lset.append(f"{lc} = :l_{lc}")
                    lb[f"l_{lc}"] = M.fit_type(m_lot, lc, ch[c][1])
            if "NET_WEIGHT" in ch and "TOTAL_WEIGHT" in m_lot:
                lset.append("TOTAL_WEIGHT = :l_tw")
                lb["l_tw"] = M.fit_type(m_lot, "TOTAL_WEIGHT", ch["NET_WEIGHT"][1])
            lmsg = "LOT 변경 없음"
            if lset:
                if "LAST_MODIFY_DATE" in m_lot:
                    lset.append("LAST_MODIFY_DATE = SYSDATE")
                cur.execute(f"UPDATE {T_LOT} SET {', '.join(lset)} WHERE RECEIPT_NO = :rno AND RECEIPT_SEQ = :seq "
                            f"AND NVL(HEAT_NO, '~') = NVL(:heat, '~')", lb)
                if cur.rowcount != 1:
                    raise UserError(f"재고LOT 줄이 {cur.rowcount}개라 고칠 수 없습니다 (1개여야 함). 저장 안 함.")
                lmsg = "LOT 1줄 수정"
            smsg = ""
            if stock_ch:
                kg_old = RCV.to_kg(o.get("NET_WEIGHT")) or 0
                q_old = o.get("QTY") or 0
                s1 = RCV.stock_add(cur, o["ITEM_CODE"], -q_old, -kg_old, who, meta=m_stk)
                s2 = RCV.stock_add(cur, v["ITEM_CODE"], v["QTY"], kg_new, who, org=o.get("ORGANIZATION_ID") or 1,
                                   unit=o.get("QTY_UNIT") or "RL", meta=m_stk)
                smsg = f"; 재고: {s1} / {s2}" if o["ITEM_CODE"] != v["ITEM_CODE"] else f"; 재고: {s2}"
            conn.commit()
        except UserError:
            conn.rollback()
            raise
        except Exception as e:
            conn.rollback()
            raise RuntimeError("수정 중 오류 -> 전부 취소 (MES는 그대로)\n\n" + ora_hint(e))
        return {"ok": True, "no": show(o["RECEIPT_NO"]), "res": lmsg + smsg,
                "log": "; ".join(f"{c}: {cell_text(a)} -> {cell_text(x)}" for c, (a, x) in ch.items()) + f"; {lmsg + smsg}"}
    res = APP.run_db(work)
    if res.get("ok"):
        APP.append_log([[now_s(), "", "소재입고수정", T_RCV, res["no"], res.pop("log")]])
    return res


def api_rcv_delete(body, user):
    rows = cache_get(body.get("token"))["rcv"]
    sel = sorted({int(i) for i in body.get("idxs") or []})
    if not sel:
        raise UserError("삭제할 줄을 선택하세요.")
    targets, seen, blocked = [], set(), []
    for i in sel:
        r = rows[i]
        if M.ReceiptTab.state(r)[1] == "off":
            blocked.append(RCV.label(r))
            continue
        k = (r["RECEIPT_NO"], r["SEQ"])
        if k not in seen:
            seen.add(k)
            targets.append(r)
    if blocked:
        raise UserError("생산에 사용된 소재라 삭제할 수 없습니다:\n\n" + "\n".join(blocked[:15])
                        + ("\n\n나머지 줄만 선택해서 다시 하세요." if targets else ""))
    if not body.get("confirm"):
        lines = "\n".join(f"  {RCV.label(r)}" for r in targets[:20]) + \
            (f"\n  ... 외 {len(targets) - 20}건" if len(targets) > 20 else "")
        note = ("\n\n※ 같은 입고번호·SEQ의 입고줄과 취소줄은 함께 지워집니다."
                if any(r["CANCEL_CNT"] or r["IO_FLAG"] == 2 for r in targets) else "")
        return {"preview": True, "msg": f"MES에서 아래 {len(targets)}건의 소재 입고를 삭제합니다.\n"
                                        f"(입고 + 재고LOT 삭제, 소재 재고에서 수량·중량 차감){note}\n\n{lines}\n\n"
                                        f"되돌릴 수 없습니다. 진행하려면 '삭제'라고 입력하세요:"}
    if (body.get("typed") or "").strip() != "삭제":
        raise UserError("'삭제'라고 입력해야 지울 수 있습니다.")
    now = dt.datetime.now()
    results = APP.run_db(lambda c: [(r, *RCV.delete_one(c, r, now)) for r in targets])
    APP.append_log([[now_s(), "", "입고삭제", T_RCV, f"{show(r['RECEIPT_NO'])}-{show(r['SEQ'])}", msg]
                    for r, ok, msg in results if ok])
    okn = sum(1 for _, ok, _ in results if ok)
    fails = [f"  {RCV.label(r)}\n     → {msg}" for r, ok, msg in results if not ok]
    return {"ok": True, "msg": f"삭제 완료 {okn}건 / 실패 {len(fails)}건" + ("\n\n실패:\n" + "\n".join(fails[:10]) if fails else ""),
            "fail": bool(fails)}


# ==============================================================
# ⑤ SPM 관리 (spm.xlsx, 기존 프로그램과 같은 파일)
# ==============================================================
def spm_rows():
    try:
        return M.spm_load()
    except Exception as e:
        raise UserError(f"SPM 파일을 읽지 못했습니다: {e}")


def web_name():
    return getattr(TL, "name", "") or reg_user()


def api_spm_list(body, user):
    rows = sorted(spm_rows(), key=lambda x: x["item"].upper())
    return {"rows": [[r["item"], f"{r['spm']:g}", r.get("date", ""), r.get("user", ""), r.get("note", "")] for r in rows],
            "path": M.SPM_XLSX}


def api_spm_save(body, user):
    it = str(body.get("item") or "").strip().upper()
    try:
        s = float(str(body.get("spm") or "").replace(",", "").strip())
        if not 0 < s <= 1000:
            raise ValueError
    except ValueError:
        raise UserError("SPM은 0보다 큰 숫자로 넣으세요.")
    if not it:
        raise UserError("품번을 넣으세요.")
    with LOCK:
        rows = spm_rows()
        old = next((r for r in rows if r["item"].upper() == it), None)
        if body.get("new") and old and not body.get("force"):
            return {"ask": f"품번 {it}는 이미 SPM {old['spm']:g} 로 등록되어 있습니다.\n{s:g} 로 바꿀까요?"}
        if body.get("new") and not old and not body.get("force"):
            try:
                n = APP.run_db(lambda c: c.cursor().execute(
                    "SELECT COUNT(*) FROM ICOM_ITEM_MASTER WHERE ITEM_CODE = :i", i=it).fetchone()[0])
                if not n:
                    return {"ask": f"품목 마스터에 없는 품번입니다: {it}\n그래도 등록할까요?"}
            except Exception:
                pass
        row = old if old else {"item": it}
        if not old:
            rows.append(row)
        prev = dict(row)
        row.update(spm=s, note=str(body.get("note") or "").strip(), date=dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
                   user=web_name())
        try:
            M.spm_save(rows)
        except RuntimeError as e:
            raise UserError(str(e))
    APP.append_log([[now_s(), "", "SPM등록" if not old else "SPM수정", "spm.xlsx", it,
                     f"SPM {prev.get('spm', '') if old else '(없음)'} -> {s:g}"]])
    return {"ok": True, "item": it}


def api_spm_delete(body, user):
    items = {str(x).upper() for x in body.get("items") or []}
    if not items:
        raise UserError("삭제할 줄을 선택하세요.")
    with LOCK:
        rows = spm_rows()
        sel = [r for r in rows if r["item"].upper() in items]
        keep = [r for r in rows if r["item"].upper() not in items]
        try:
            M.spm_save(keep)
        except RuntimeError as e:
            raise UserError(str(e))
    APP.append_log([[now_s(), "", "SPM삭제", "spm.xlsx", r["item"], f"SPM {r['spm']:g}"] for r in sel])
    return {"ok": True, "n": len(sel)}


def api_spm_import(body, user):
    need_xl()
    _tok, path = save_upload(body)
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        got = []
        for r in wb.worksheets[0].iter_rows(values_only=True):
            if not r or len(r) < 2 or r[0] in (None, ""):
                continue
            try:
                got.append((str(r[0]).strip().upper(), float(str(r[1]).replace(",", ""))))
            except (TypeError, ValueError):
                continue
        wb.close()
    finally:
        try:
            os.remove(path)
        except Exception:
            pass
    if not got:
        raise UserError("가져올 줄이 없습니다. (A열 품번, B열 SPM 숫자)")
    with LOCK:
        rows = spm_rows()
        by = {r["item"].upper(): r for r in rows}
        add = [g for g in got if g[0] not in by]
        chg = [g for g in got if g[0] in by and by[g[0]]["spm"] != g[1]]
        if not add and not chg:
            return {"msg": f"{len(got)}줄 모두 지금 값과 같습니다."}
        if not body.get("confirm"):
            return {"ask": f"엑셀 {len(got)}줄 중\n  새 품번 {len(add)}개 추가\n  SPM 바뀜 {len(chg)}개\n\n"
                           f"반영할까요? (엑셀에 없는 기존 품번은 그대로 둡니다)"}
        now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
        for it, s in add + chg:
            r = by.get(it)
            if r is None:
                r = {"item": it, "note": ""}
                rows.append(r)
                by[it] = r
            r.update(spm=s, date=now, user=web_name())
        try:
            M.spm_save(rows)
        except RuntimeError as e:
            raise UserError(str(e))
    APP.append_log([[now_s(), "", "SPM가져오기", "spm.xlsx", body.get("name", ""), f"추가 {len(add)} / 변경 {len(chg)}"]])
    return {"ok": True, "msg": f"추가 {len(add)}개, 변경 {len(chg)}개 반영했습니다."}


def api_spm_missing(body, user):
    tb = wo_tbl()

    def work(conn):
        cur = conn.cursor()
        cur.execute(f"SELECT ITEM_CODE, MAX(ITEM_NAME), COUNT(*), MAX(WORK_ORDER_DATE), "
                    f"MAX(MACHINE_NAME) KEEP (DENSE_RANK LAST ORDER BY WORK_ORDER_DATE) FROM {tb} "
                    f"WHERE WORK_ORDER_DATE >= TRUNC(SYSDATE) - 30 AND ITEM_CODE IS NOT NULL "
                    f"GROUP BY ITEM_CODE ORDER BY ITEM_CODE")
        return cur.fetchall()
    rows = APP.run_db(work)
    have = {r["item"].upper() for r in spm_rows()}
    miss = [r for r in rows if show(r[0]).strip().upper() not in have]
    return {"total": len(rows), "rows": [[show(r[0]), show(r[1]), show(r[2]), r[3].strftime("%Y-%m-%d") if r[3] else "",
                                          show(r[4])] for r in miss]}


def api_spm_download(body, user):
    if not os.path.exists(M.SPM_XLSX):
        M.spm_load()
    with open(M.SPM_XLSX, "rb") as f:
        return FileOut(f.read(), "spm.xlsx", XLSX)


# ==============================================================
# MES 구조조사 / 전후 비교 (숨김 탭, 조회만)
# ==============================================================
INS = M.Inspector(APP)
_snap = {"data": {}, "time": None, "diff": []}


def api_find(body, user):
    kind = body.get("kind")
    db = INS
    if kind == "table":
        rows = APP.run_db(db.related_tables)
        data = [[n, c, k, ("" if r is None else f"{r:,}") + ("" if ex else " (추정)")] for n, c, k, r, ex in rows]
        return {"cols": ["테이블", "설명", "컬럼수", "행수"], "rows": rows_text(data), "tags": [""] * len(data),
                "msg": f"작업지시 관련 테이블 {len(data)}개 (이름에 키워드가 있거나 WORK_ORDER_NO 컬럼이 있는 것). 더블클릭 = 컬럼 구조"}
    if kind == "prog":
        rows = APP.run_db(db.programs)
        return {"cols": ["이름", "종류", "상태", "마지막 변경", "찾은 이유"], "rows": rows_text(rows),
                "tags": ["hot" if "내용" in r[4] else "" for r in rows],
                "msg": f"프로그램 {len(rows)}개.  노란 줄 = 안에서 작업지시 테이블에 INSERT/UPDATE 함 (가장 중요).  더블클릭 = 내용 보기"}
    if kind == "seq":
        rows = APP.run_db(db.sequences)
        return {"cols": ["번호생성기", "다음 번호 근처", "증가", "최소", "최대"], "rows": rows_text(rows), "tags": [""] * len(rows),
                "msg": f"번호 생성기(시퀀스) {len(rows)}개. 작업지시 번호가 여기서 나오는지 확인하는 용도입니다."}
    if kind == "trg":
        rows = APP.run_db(lambda c: db.triggers(c, [r[0] for r in db.related_tables(c)]))
        return {"cols": ["트리거", "테이블", "언제", "방식", "상태"], "rows": rows_text(rows), "tags": [""] * len(rows),
                "msg": f"관련 테이블에 걸린 자동동작(트리거) {len(rows)}개. 더블클릭 = 내용 보기"}
    if kind == "src":
        q = str(body.get("q") or "").strip()
        if len(q) < 3:
            raise UserError("3글자 이상 입력하세요. 예: IPRD_WORK_ORDER_MASTER")
        rows = APP.run_db(lambda c: db.search_source(c, q))
        return {"cols": ["프로그램", "종류", "줄", "내용"], "rows": rows_text(rows),
                "tags": ["hot" if re.search(r"\b(INSERT|UPDATE|DELETE|MERGE)\b", str(r[3]).upper()) else "" for r in rows],
                "msg": f"'{q}' 가 들어간 곳 {len(rows)}줄{' (1000줄까지)' if len(rows) >= 1000 else ''}.  노란 줄 = 쓰기(INSERT/UPDATE/DELETE).  더블클릭 = 프로그램 전체 보기"}
    raise UserError("없는 기능입니다.")


def api_find_columns(body, user):
    rows = APP.run_db(lambda c: INS.columns(c, str(body.get("table") or "")))
    return {"cols": ["컬럼", "형식", "길이", "빈값허용", "DB기본값", "설명"],
            "rows": [[a, b, show(c), d, str(e or "").strip(), f or ""] for a, b, c, d, e, f in rows]}


def api_find_source(body, user):
    name, typ = str(body.get("name") or ""), str(body.get("typ") or "")
    text, args = APP.run_db(lambda c: INS.source(c, name, typ))
    head = f"-- {typ} {name}" + (f"\n-- 입력값(파라미터):{args}" if args else "")
    return {"text": text or "(내용을 볼 권한이 없거나 암호화된 프로그램입니다)", "head": head}


def api_find_export_src(body, user):
    items = [(str(a), str(b)) for a, b in body.get("items") or []][:300]
    if not items:
        raise UserError("저장할 항목이 없습니다.")

    def work(conn):
        out = []
        for name, typ in items:
            text, args = INS.source(conn, name, typ)
            out.append(f"{'=' * 70}\n-- {typ} {name}" + (f"\n-- 입력값:{args}" if args else "") +
                       f"\n{'=' * 70}\n{text or '(내용을 볼 권한이 없거나 암호화됨)'}\n")
        return out
    parts = APP.run_db(work)
    return FileOut("\n".join(parts).encode("utf-8"), f"MES_프로그램내용_{dt.datetime.now():%Y%m%d_%H%M}.txt",
                   "text/plain; charset=utf-8")


def job_diff_snap(body, user, progress, cancelled):
    def work(conn):
        names = INS.all_tables(conn, body.get("scope", "kw") == "kw")
        lim = int(CFG.get("small_limit") or 30000)
        snap, skipped = {}, []
        for i, n in enumerate(names, 1):
            if cancelled():
                raise UserError("취소했습니다.")
            progress(f"상태 저장 중... {i}/{len(names)}  {n}")
            try:
                snap[n] = INS.snap_table(conn, n, lim)
            except Exception as e:
                skipped.append(f"{n}({str(e)[:40]})")
        return snap, skipped
    snap, skipped = APP.run_db(work)
    with LOCK:
        _snap.update(data=snap, time=dt.datetime.now(), diff=[])
    msg = f"{_snap['time']:%H:%M:%S} 상태 저장 완료 ({len(snap)}개 테이블)."
    if skipped:
        msg += f"  읽지 못한 테이블 {len(skipped)}개"
    return {"msg": msg}


def job_diff_compare(body, user, progress, cancelled):
    if not _snap["data"]:
        raise UserError("먼저 [① 찍기 전 저장]을 하세요.")
    tab = object.__new__(M.InspectTabs)
    tab.app, tab.db, tab.snap, tab.snap_time = APP, INS, _snap["data"], _snap["time"]

    def prog(m):
        if cancelled():
            raise UserError("취소했습니다.")
        progress(m)
    tab.progress = prog
    diff = APP.run_db(tab.compare)
    with LOCK:
        _snap["diff"] = diff
    return diff_view(True)


def diff_view(hide=True):
    rows, tags, view = [], [], []
    order = sorted(enumerate(_snap["diff"]),
                   key=lambda x: (-(abs(x[1]["after"] - x[1]["before"]) + len(x[1]["new"]) + len(x[1]["upd"])), x[1]["table"]))
    for i, d in order:
        changed = d["before"] != d["after"] or d["cols"]
        if hide and not changed:
            continue
        view.append(i)
        diffn = d["after"] - d["before"]
        rows.append([d["table"], f"{d['before']:,}", f"{d['after']:,}", f"{diffn:+,}" if diffn else "0",
                     len(d["new"]), len(d["upd"]), d["del"], ", ".join(d["cols"][:12]) + (" …" if len(d["cols"]) > 12 else ""), d["how"]])
        tags.append("del" if d["del"] and not d["new"] else ("new" if d["new"] or diffn > 0 else ("upd" if changed else "")))
    n = sum(1 for d in _snap["diff"] if d["before"] != d["after"] or d["cols"])
    return {"cols": ["테이블", "전", "후", "증감", "새 줄", "바뀐 줄", "없어진 줄", "바뀐 컬럼", "확인 방법"],
            "rows": rows_text(rows), "tags": tags, "view": view,
            "msg": f"저장 {_snap['time']:%H:%M:%S} 대비 바뀐 테이블 {n}개 / 전체 {len(_snap['diff'])}개.  더블클릭 = 새 줄·바뀐 줄 내용"}


def api_diff_view(body, user):
    if not _snap["diff"]:
        return {"cols": [], "rows": [], "tags": [], "view": [],
                "msg": (f"{_snap['time']:%H:%M:%S} 상태 저장됨 → MES에서 작업 후 [③ 찍은 후 비교]" if _snap["time"]
                        else "아직 저장한 상태가 없습니다.")}
    return diff_view(bool(body.get("hide", True)))


def api_diff_detail(body, user):
    d = _snap["diff"][int(body.get("i"))]
    if not d["hdr"]:
        return {"msg": f"{d['table']}: 줄 내용을 볼 수 없습니다.\n({d['how']})\n바뀐 컬럼: {', '.join(d['cols']) or '없음'}"}
    rows = [["새 줄"] + list(r) for r in d["new"]] + [["바뀐 줄"] + list(r) for r in d["upd"]]
    return {"title": f"{d['table']} - 새 줄 {len(d['new'])} / 바뀐 줄 {len(d['upd'])}", "cols": ["구분"] + list(d["hdr"]),
            "rows": [[M.cell(v) for v in r] for r in rows],
            "tags": ["new"] * len(d["new"]) + ["upd"] * len(d["upd"]), "note": f"바뀐 컬럼: {', '.join(d['cols']) or '-'}"}


def api_diff_export(body, user):
    if not _snap["diff"]:
        raise UserError("먼저 비교하세요.")
    sheets = [("요약", ["테이블", "전", "후", "증감", "새 줄", "바뀐 줄", "없어진 줄", "바뀐 컬럼", "확인 방법"],
               [[d["table"], d["before"], d["after"], d["after"] - d["before"], len(d["new"]), len(d["upd"]), d["del"],
                 ", ".join(d["cols"]), d["how"]] for d in _snap["diff"] if d["before"] != d["after"] or d["cols"]])]
    for d in _snap["diff"]:
        if d["hdr"] and (d["new"] or d["upd"]):
            sheets.append((d["table"][:28], ["구분"] + list(d["hdr"]),
                           [["새 줄"] + [M.cell(v, 2000) for v in r] for r in d["new"]] +
                           [["바뀐 줄"] + [M.cell(v, 2000) for v in r] for r in d["upd"]]))
    return FileOut(xlsx_of(sheets), f"전후비교_{dt.datetime.now():%Y%m%d_%H%M}.xlsx", XLSX)


# ==============================================================
# 환경설정 / 작업기록 / 사용법 / 공용
# ==============================================================
SET_FIELDS = M.App.SET_FIELDS


def get_cfg_path(path):
    d = CFG
    for k in path.split("."):
        d = d.get(k, "") if isinstance(d, dict) else ""
    return d


def show_flags():
    return {"dev": bool(CFG.get("show_dev_tabs")), "reg": bool(CFG.get("show_reg_tabs")), "item": bool(CFG.get("show_item_tab"))}


def api_meta(body, user):
    specs = {}
    for k, s in M.GRID_SPECS.items():
        specs[k] = {"title": s["title"], "hint": s.get("hint", ""), "filters": [list(f) for f in s["filters"]],
                    "hide_new": bool(s.get("hide_new")), "match": bool(s.get("match_col") or s.get("match_cols")),
                    "tools": [t[0] for t in s.get("tools", [])], "auto_fetch": bool(s.get("auto_fetch")),
                    "new_hint": s.get("new_hint", "")}
    return {"version": VERSION, "program": os.path.basename(PROGRAM), "app_title": M.APP_TITLE,
            "tabs": [[k, t] for k, t in M.App.TAB_TITLES], "page_info": M.App.PAGE_INFO, "help": M.App.HELP,
            "hidden": {"dev": list(M.App.DEV_TABS), "reg": list(M.App.REG_TABS), "item": list(M.App.ITEM_TABS)},
            "show": show_flags(), "grids": specs, "quick_fields": [list(f) for f in QUICK_FIELDS],
            "pin": bool(str(CFG.get("web_pin") or "").strip()), "today": f"{dt.date.today():%Y-%m-%d}",
            "spm_path": M.SPM_XLSX, "rcv_cols": M.ReceiptTab.COLS}


def api_set_get(body, user):
    return {"fields": [[p, label, hint, show(get_cfg_path(p))] for p, label, hint in SET_FIELDS],
            "db_user": M.DB_USER, "show": show_flags(),
            "allow": bool(CFG.get("allow_item_to_existing")), "web_pin": str(CFG.get("web_pin") or ""),
            "auto_quit": auto_quit_on(),
            "files": {"설정": M.CFG_PATH, "작업기록": M.LOG_PATH, "삭제 백업": M.BACKUP_PATH, "SPM": M.SPM_XLSX,
                      "프로그램": PROGRAM}}


def api_set_save(body, user):
    with LOCK:
        for path, val in (body.get("values") or {}).items():
            if path not in {p for p, _l, _h in SET_FIELDS}:
                continue
            val = str(val).strip()
            if path in ("kor_bytes", "limit", "small_limit"):
                val = int(val) if val.isdigit() else M.DEFAULT_CFG[path]
            keys = path.split(".")
            if len(keys) == 2:
                CFG[keys[0]][keys[1]] = val
            else:
                CFG[path] = val
        for k, ck in (("dev", "show_dev_tabs"), ("reg", "show_reg_tabs"), ("item", "show_item_tab")):
            if k in (body.get("show") or {}):
                CFG[ck] = bool(body["show"][k])
        if "allow" in body:
            CFG["allow_item_to_existing"] = bool(body["allow"])
        if "auto_quit" in body:
            CFG["web_auto_quit"] = bool(body["auto_quit"])
        if "web_pin" in body:
            CFG["web_pin"] = str(body.get("web_pin") or "").strip()
        APP.save_cfg()
    return {"ok": True, "show": show_flags()}


def api_set_test(body, user):
    def fn(conn):
        cur = conn.cursor()
        lines = ["접속 성공."]
        for p in M.PARTS:
            if not MES.enabled(p) or not CFG[p].get("table"):
                continue
            try:
                cur.execute(f"SELECT COUNT(*) FROM {MES.table(p)}")
                lines.append(f"{M.PART_NAME[p]}: {MES.table(p)} {cur.fetchone()[0]:,}건")
            except Exception as e:
                lines.append(f"{M.PART_NAME[p]}: {MES.table(p)} 조회 실패 - {e}")
        return "\n".join(lines)
    return {"msg": APP.run_db(fn)}


def api_log_list(body, user):
    if not os.path.exists(M.LOG_PATH):
        return {"cols": M.App.LOG_COLS, "rows": [], "total": 0}
    with open(M.LOG_PATH, encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))[1:]
    q = str(body.get("q") or "").strip().lower()
    rows = list(reversed(rows))
    if q:
        rows = [r for r in rows if any(q in str(v).lower() for v in r)]
    return {"cols": M.App.LOG_COLS, "rows": rows[:3000], "total": len(rows)}


def api_log_file(body, user):
    path = M.BACKUP_PATH if body.get("kind") == "backup" else M.LOG_PATH
    if not os.path.exists(path):
        raise UserError("아직 기록 파일이 없습니다.")
    with open(path, "rb") as f:
        return FileOut(f.read(), os.path.basename(path),
                       "text/csv; charset=utf-8" if path.endswith(".csv") else "application/octet-stream")


def api_xlsx(body, user):
    """화면 표를 그대로 엑셀로 (소재 입고·조사 결과 등)"""
    return FileOut(xlsx_of([(body.get("sheet") or "Sheet1", body.get("cols") or [], body.get("rows") or [])]),
                   (body.get("name") or "표") + f"_{dt.datetime.now():%Y%m%d_%H%M}.xlsx", XLSX)


def api_ping(body, user):
    def work(conn):
        cur = conn.cursor()
        cur.execute("SELECT SYSDATE FROM DUAL")
        return cur.fetchone()[0]
    return {"ok": True, "db_time": show(APP.run_db(work))}


# ---------- 오래 걸리는 작업 (진행 표시 + 취소) ----------
JOBS = {}


def api_job_start(body, user):
    fn = JOB_ROUTES.get(body.get("path"))
    if not fn:
        raise UserError("없는 기능입니다.")
    jid = uuid.uuid4().hex[:12]
    job = {"msg": "MES 접속 중...", "done": False, "stop": False, "result": None, "error": None, "kind": None,
           "t": dt.datetime.now()}
    JOBS[jid] = job
    name = getattr(TL, "name", "")

    def run():
        TL.user, TL.name = user, name
        try:
            job["result"] = jsonable(fn(body.get("body") or {}, user, lambda m: job.__setitem__("msg", m),
                                        lambda: job["stop"]))
        except UserError as e:
            job["error"], job["kind"] = str(e), "warn"
        except Exception as e:
            traceback.print_exc()
            job["error"] = ora_hint(e)
        finally:
            job["done"] = True
    threading.Thread(target=run, daemon=True).start()
    for k in [k for k, j in list(JOBS.items()) if j["done"] and (dt.datetime.now() - j["t"]).seconds > 3600]:
        JOBS.pop(k, None)
    return {"id": jid}


def api_job_poll(body, user):
    job = JOBS.get(body.get("id"))
    if not job:
        raise UserError("작업 정보를 찾지 못했습니다.")
    if not job["done"]:
        return {"done": False, "msg": job["msg"]}
    JOBS.pop(body.get("id"), None)
    return {"done": True, "result": job["result"], "error": job["error"], "kind": job["kind"]}


def api_job_cancel(body, user):
    job = JOBS.get(body.get("id"))
    if job:
        job["stop"] = True
    return {"ok": True}


ROUTES = {
    "/api/meta": api_meta, "/api/ping": api_ping,
    "/api/grid/search": api_grid_search, "/api/grid/save": api_grid_save, "/api/grid/delete": api_grid_delete,
    "/api/grid/new_form": api_grid_new_form, "/api/grid/insert": api_grid_insert,
    "/api/grid/export": api_grid_export, "/api/grid/import": api_grid_import, "/api/grid/tool": api_grid_tool,
    "/api/wo/form": api_wo_form, "/api/wo/form_save": api_wo_form_save,
    "/api/item/lookup": api_wo_item_lookup, "/api/item/find": api_wo_item_find,
    "/api/batch/lists": api_batch_lists, "/api/batch/spm": api_batch_spm, "/api/batch/recent": api_batch_recent,
    "/api/batch/check": api_batch_check, "/api/batch/make_no": api_batch_make_no, "/api/batch/apply": api_batch_apply,
    "/api/old/fetch": api_old_fetch, "/api/old/items": api_old_items, "/api/old/find_items": api_old_find_items,
    "/api/old/save": api_old_save, "/api/old/delete_items": api_old_delete_items,
    "/api/old/delete_mold": api_old_delete_mold, "/api/old/save_as": api_old_save_as,
    "/api/old/items_of": api_old_items_of, "/api/old/template": api_old_template,
    "/api/bulk/template": api_bulk_template, "/api/bulk/upload": api_bulk_upload, "/api/bulk/check": api_bulk_check,
    "/api/bulk/register": api_bulk_register, "/api/bulk/download": api_bulk_download,
    "/api/map/get": api_map_get, "/api/map/fetch_cols": api_map_fetch_cols, "/api/map/save": api_map_save,
    "/api/map/find_tables": api_map_find_tables, "/api/table/preview": api_table_preview,
    "/api/quick/find": api_quick_find, "/api/quick/base": api_quick_base, "/api/quick/item": api_quick_item,
    "/api/quick/check": api_quick_check, "/api/quick/register": api_quick_register,
    "/api/rcv/search": api_rcv_search, "/api/rcv/form": api_rcv_form, "/api/rcv/next_no": api_rcv_next_no,
    "/api/rcv/item_lookup": api_rcv_item_lookup, "/api/rcv/item_find": api_rcv_item_find,
    "/api/rcv/save_new": api_rcv_save_new, "/api/rcv/save_edit": api_rcv_save_edit, "/api/rcv/delete": api_rcv_delete,
    "/api/spm/list": api_spm_list, "/api/spm/save": api_spm_save, "/api/spm/delete": api_spm_delete,
    "/api/spm/import": api_spm_import, "/api/spm/missing": api_spm_missing, "/api/spm/download": api_spm_download,
    "/api/find": api_find, "/api/find/columns": api_find_columns, "/api/find/source": api_find_source,
    "/api/find/export_src": api_find_export_src,
    "/api/diff/view": api_diff_view, "/api/diff/detail": api_diff_detail, "/api/diff/export": api_diff_export,
    "/api/set/get": api_set_get, "/api/set/save": api_set_save, "/api/set/test": api_set_test,
    "/api/log/list": api_log_list, "/api/log/file": api_log_file, "/api/xlsx": api_xlsx,
    "/api/job/start": api_job_start, "/api/job/poll": api_job_poll, "/api/job/cancel": api_job_cancel,
}
JOB_ROUTES = {
    "/api/grid/delete_check": job_grid_delete_check,
    "/api/old/delete_mold_check": job_old_delete_mold_check,
    "/api/diff/snap": job_diff_snap, "/api/diff/compare": job_diff_compare,
}
# --------------------------------------------------------------
# 웹 화면을 모두 닫으면 서버도 끄기 (열린 화면이 20초마다 신호를 보냄)
# --------------------------------------------------------------
PAGES = {}            # 화면 id -> 마지막 신호 시각
PAGES_LOCK = threading.Lock()
PAGE_SEEN = [False]   # 화면이 한 번이라도 열렸는지 (서버만 켜 두고 화면 안 연 경우는 안 끔)
HB_DEAD = 150         # 이 시간(초) 동안 신호 없으면 닫힌 화면으로 봄 (숨은 탭은 신호가 1분에 1번까지 늦어질 수 있음)
QUIT_WAIT = 15        # 화면이 모두 닫힌 뒤 이만큼 기다렸다 끔 (새로고침 대비)


def auto_quit_on():
    return CFG.get("web_auto_quit", True) is not False


def api_hb(body, user):
    pid = str(body.get("id") or "")[:40]
    if pid:
        with PAGES_LOCK:
            PAGES[pid] = time.time()
            PAGE_SEEN[0] = True
    return {"ok": True, "auto_quit": auto_quit_on()}


def api_bye(body, user):
    with PAGES_LOCK:
        PAGES.pop(str(body.get("id") or "")[:40], None)
    return {"ok": True}


def quit_watch(srv):
    empty_since = None
    while True:
        time.sleep(3)
        now = time.time()
        with PAGES_LOCK:
            for k in [k for k, t in PAGES.items() if now - t > HB_DEAD]:
                PAGES.pop(k, None)
            alive = len(PAGES)
        if not auto_quit_on() or not PAGE_SEEN[0] or alive:
            empty_since = None
            continue
        if empty_since is None:
            empty_since = now
        elif now - empty_since >= QUIT_WAIT:
            print(f"\n[{dt.datetime.now():%H:%M:%S}] 열린 웹 화면이 없어 서버를 끕니다.")
            srv.shutdown()
            return


ROUTES["/api/hb"] = api_hb
ROUTES["/api/bye"] = api_bye
NO_PIN = {"/api/meta", "/api/hb", "/api/bye"}


def find_page():
    """웹 화면 html 찾기: index.html 이 있으면 그것, 없으면 이 웹 화면 html (이름이 달라도 됨)"""
    p = os.path.join(HERE, "index.html")
    if os.path.exists(p):
        return p
    hits = []
    for f in glob.glob(os.path.join(HERE, "*.htm*")):
        try:
            with open(f, "r", encoding="utf-8", errors="ignore") as fh:
                if "startHeartbeat" in fh.read():
                    hits.append(f)
        except Exception:
            pass
    return max(hits, key=os.path.getmtime) if hits else None


class Handler(BaseHTTPRequestHandler):
    server_version = "TJD-MES-Web"

    def log_message(self, fmt, *args):
        if "/api/" in (self.path or "") and not any(x in (self.path or "") for x in ("/api/job/poll", "/api/hb", "/api/bye")):
            sys.stdout.write(f"[{dt.datetime.now():%H:%M:%S}] {self.client_address[0]} {fmt % args}\n")

    def send(self, code, data, ctype="application/json; charset=utf-8", extra=None):
        if not isinstance(data, (bytes, bytearray)):
            data = json.dumps(jsonable(data), ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        p = urlparse(self.path).path
        if p in ("/", "/index.html"):
            path = find_page()
            if not path:
                return self.send(404, "웹 화면 html 파일이 서버 파일과 같은 폴더에 없습니다.".encode("utf-8"),
                                 "text/plain; charset=utf-8")
            with open(path, "rb") as f:
                return self.send(200, f.read(), "text/html; charset=utf-8")
        self.send(404, {"error": "없는 주소입니다."})

    def do_POST(self):
        p = urlparse(self.path).path
        try:
            n = int(self.headers.get("Content-Length") or 0)
            body = json.loads(self.rfile.read(n).decode("utf-8") or "{}") if n else {}
        except Exception:
            return self.send(400, {"error": "요청 형식 오류"})
        pin = str(CFG.get("web_pin") or "").strip()
        if pin and p not in NO_PIN and self.headers.get("X-Pin", "") != pin:
            return self.send(401, {"error": "PIN이 맞지 않습니다.", "pin": True})
        name = unquote(self.headers.get("X-User", "") or "").strip()[:20] or "이름없음"
        user = f"{name}(웹 {self.client_address[0]})"
        TL.user, TL.name = user, name
        try:
            fn = ROUTES.get(p)
            if not fn:
                return self.send(404, {"error": "없는 기능입니다."})
            res = fn(body, user)
            if isinstance(res, FileOut):
                return self.send(200, res.data, res.ctype,
                                 {"Content-Disposition": "attachment; filename*=UTF-8''" + quote(res.name)})
            self.send(200, res)
        except UserError as e:
            self.send(400, {"error": str(e), "kind": "warn"})
        except Exception as e:
            traceback.print_exc()
            self.send(500, {"error": ora_hint(e)})
        finally:
            TL.user = TL.name = ""


def local_ips():
    ips = set()
    try:
        for info in socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET):
            ips.add(info[4][0])
    except Exception:
        pass
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(((CFG["oracle"].get("dsn") or "192.168.1.250").split(":")[0], 1521))
        ips.add(s.getsockname()[0])
        s.close()
    except Exception:
        pass
    return sorted(i for i in ips if not i.startswith("127."))


def main():
    port = int(sys.argv[1]) if len(sys.argv) > 1 else int(CFG.get("web_port") or 8000)
    print("=" * 64)
    print(f" 태진다이텍 MES 웹 서버  {VERSION}")
    print(f" 프로그램: {os.path.basename(PROGRAM)}  (DB 규칙을 이 파일에서 그대로 씀)")
    pg = find_page()
    print(f" 웹 화면: {os.path.basename(pg) if pg else '※ html 파일 없음 (서버 파일과 같은 폴더에 두세요)'}")
    print(f" 설정: {M.CFG_PATH if os.path.exists(M.CFG_PATH) else '(설정 파일 없음 - 기본값)'}")
    print(f" DB  : {CFG['oracle']['dsn']}")
    if oracledb is None:
        print(" ※ oracledb 없음: python -m pip install oracledb")
    else:
        try:
            print(f" DB 접속 확인: OK (DB 시각 {api_ping({}, '')['db_time']})")
        except Exception as e:
            print(" ※ DB 접속 실패 (서버는 켜 둡니다. 화면에서 다시 시도하세요)\n   " + str(e).replace("\n", "\n   "))
    if openpyxl is None:
        print(" ※ openpyxl 없음 (엑셀·SPM 안 됨): python -m pip install openpyxl")
    print("-" * 64)
    print(" 다른 PC·폰 브라우저에서 접속:")
    for ip in local_ips() or ["이PC주소"]:
        print(f"   http://{ip}:{port}")
    print(f"   (이 PC에서는 http://localhost:{port})")
    print(" 끄려면 이 창을 닫거나 Ctrl+C")
    if auto_quit_on():
        print(" ※ 웹 화면을 모두 닫으면 서버도 자동으로 꺼집니다 (환경설정에서 끌 수 있음)")
    print("=" * 64)
    srv = ThreadingHTTPServer(("0.0.0.0", port), Handler)
    srv.daemon_threads = True
    threading.Thread(target=quit_watch, args=(srv,), daemon=True).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass
    srv.server_close()


if __name__ == "__main__":
    main()
